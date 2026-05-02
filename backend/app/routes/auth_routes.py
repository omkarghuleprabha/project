from flask import Blueprint, render_template, request, redirect, url_for, flash, session, jsonify, current_app, send_file
import mysql.connector
from werkzeug.security import generate_password_hash, check_password_hash
from app.utils.db import get_db
from app.utils.complaints import (
    complaint_original_photo_select,
    complaint_progress_percent,
    complaint_status_class,
    complaint_status_key,
    ensure_complaint_workflow_columns,
    get_complaint_columns,
    normalize_complaint_status,
    save_complaint_upload,
)
from app.utils.location import build_location_payload, ensure_request_location_columns
from app.utils.inquiries import fetch_inquiries
from app.utils.equipment import (
    create_equipment_request,
    ensure_equipment_schema,
    equipment_catalog_groups,
    fetch_inventory,
    fetch_request_rows,
    fetch_transaction_rows,
    inventory_summary,
    issue_inventory,
    mark_request_fulfilled,
    mark_request_rejected,
    normalize_equipment_item,
    restock_inventory,
)
from app.utils.staff_portal import (
    ANNOUNCEMENT_TARGET_ROLES,
    attach_verification_urls,
    encode_target_roles,
    ensure_announcements_table,
    ensure_announcement_views_table,
    ensure_staff_verification_columns,
    ensure_state_admin_table,
    fetch_visible_announcements,
    mark_announcements_seen,
    normalize_verification_status,
    role_display_name,
    save_verification_upload,
    staff_role_requires_verification,
)
from datetime import datetime, timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import re
import textwrap

auth_bp = Blueprint('auth_bp', __name__)

# Helper to map roles to table names
ROLE_MAP = {
    'state_admin': 'state_admins',
    'district_admin': 'district_admins',
    'admin': 'taluka_admins',
    'worker': 'village_workers',
    'user': 'users'
}

STAFF_PENDING_LOGIN_MESSAGE = "User verification is pending please try after time."
STAFF_REJECTED_LOGIN_MESSAGE = "Kindly Register properly after some time."

MAP_PRESET_CENTERS = {
    'akole': {'lat': 19.5318, 'lng': 73.9975},
}

REPORT_DATASETS = {
    'summary': {
        'label': 'Executive Summary',
        'description': 'Quick operational totals and progress metrics for the selected scope.',
        'icon': 'fa-chart-line',
    },
    'villages': {
        'label': 'Villages',
        'description': 'Village master list with citizens, workers, complaints, and door-pickup activity.',
        'icon': 'fa-map-location-dot',
    },
    'citizens': {
        'label': 'Citizens',
        'description': 'Citizen registrations across the selected area.',
        'icon': 'fa-users',
    },
    'workers': {
        'label': 'Workers',
        'description': 'Field staff, assignment load, and availability status.',
        'icon': 'fa-user-hard-hat',
    },
    'complaints': {
        'label': 'Complaints',
        'description': 'Complaint intake, assignment, and resolution activity.',
        'icon': 'fa-file-circle-exclamation',
    },
    'door_pickups': {
        'label': 'Door Pickup Services',
        'description': 'Door-to-door and bulk garbage collection service requests.',
        'icon': 'fa-truck-ramp-box',
    },
    'payments': {
        'label': 'Payments',
        'description': 'Collection payments and revenue sharing records.',
        'icon': 'fa-wallet',
    },
    'tasks': {
        'label': 'Operations Tasks',
        'description': 'Manual work assignments created for ground operations.',
        'icon': 'fa-list-check',
    },
    'all_operations': {
        'label': 'All Operations Bundle',
        'description': 'A multi-section export containing summary, citizens, workers, complaints, pickups, payments, and tasks.',
        'icon': 'fa-layer-group',
    },
}

REPORT_PREVIEW_LIMIT = 60
REPORT_SECTION_PREVIEW_LIMIT = 10
REPORT_QUICK_RANGES = {
    'all': {'label': 'All Time', 'days': None},
    '7d': {'label': 'Last 7 Days', 'days': 7},
    '30d': {'label': 'Last 30 Days', 'days': 30},
    '90d': {'label': 'Last 90 Days', 'days': 90},
}


def _post_redirect_target(default_endpoint):
    next_page = (request.form.get('next_page') or '').strip()
    if next_page.startswith('/'):
        return next_page

    referrer = (request.referrer or '').strip()
    if referrer.startswith(request.host_url):
        return referrer

    return url_for(default_endpoint)


def _empty_worker_counts():
    return {
        'total_tasks': 0,
        'completed_tasks': 0,
        'active_tasks': 0,
        'open_tasks': 0,
        'pending_tasks': 0,
        'in_progress_tasks': 0,
    }


def _merge_worker_counts(target, source):
    target['total_tasks'] += source.get('total_tasks') or 0
    target['completed_tasks'] += source.get('completed_tasks') or 0
    target['active_tasks'] += source.get('active_tasks') or 0
    target['open_tasks'] += source.get('active_tasks') or 0
    target['pending_tasks'] += source.get('pending_tasks') or 0
    target['in_progress_tasks'] += source.get('in_progress_tasks') or 0


def _format_timestamp(value, include_time=False):
    if not value:
        return 'N/A'
    return value.strftime('%d %b %Y %I:%M %p') if include_time else value.strftime('%d %b %Y')


def _slug_text(value):
    return (value or '').strip().lower()


def _build_map_query(*parts):
    return ', '.join([part.strip() for part in parts if part and str(part).strip()])


def _complaint_photo_url(filename):
    clean_name = (filename or '').strip()
    if not clean_name:
        return ''
    return url_for('static', filename=f'uploads/complaints/{clean_name}')


def _attach_complaint_photo_urls(complaint):
    original_photo_path = complaint.get('original_photo_path') or ''
    resolution_photo_path = complaint.get('resolution_photo_path') or ''
    complaint['original_photo_url'] = _complaint_photo_url(original_photo_path)
    complaint['resolution_photo_url'] = _complaint_photo_url(resolution_photo_path)
    complaint['has_original_photo'] = bool(complaint['original_photo_url'])
    complaint['has_resolution_photo'] = bool(complaint['resolution_photo_url'])
    return complaint


def _safe_int(value):
    try:
        return int(value) if value not in (None, '') else None
    except (TypeError, ValueError):
        return None


def _positive_int(value, default=None, minimum=1):
    parsed = _safe_int(value)
    if parsed is None:
        return default
    return max(parsed, minimum)


def _merge_equipment_note(prefix, notes=None):
    clean_prefix = (prefix or '').strip()
    clean_notes = (notes or '').strip()
    if clean_prefix and clean_notes:
        return f"{clean_prefix} {clean_notes}"
    return clean_prefix or clean_notes or None


def _equipment_scope_name(scope_role, scope_id, cursor, cache=None):
    cache = cache if cache is not None else {}
    cache_key = (scope_role, scope_id)
    if cache_key in cache:
        return cache[cache_key]

    scope_label = role_display_name(scope_role)
    if scope_role == 'supplier':
        cache[cache_key] = 'Supplier Intake'
        return cache[cache_key]

    if not scope_role or scope_id in (None, ''):
        cache[cache_key] = scope_label
        return cache[cache_key]

    display_name = f"{scope_label} #{scope_id}"

    if scope_role == 'worker':
        cursor.execute("""
            SELECT
                COALESCE(vw.name, 'Worker') AS name,
                COALESCE(v.name, 'Assigned Village') AS area_name
            FROM village_workers vw
            LEFT JOIN villages v ON vw.village_id = v.id
            WHERE vw.id = %s
        """, (scope_id,))
        record = cursor.fetchone() or {}
        name = record.get('name') or 'Worker'
        area_name = record.get('area_name')
        display_name = f"{name} - {area_name}" if area_name else name
    elif scope_role == 'taluka':
        cursor.execute("""
            SELECT
                COALESCE(t.name, 'Taluka') AS name,
                COALESCE(d.name, 'District') AS parent_name
            FROM talukas t
            LEFT JOIN districts d ON t.district_id = d.id
            WHERE t.id = %s
        """, (scope_id,))
        record = cursor.fetchone() or {}
        name = record.get('name') or 'Taluka'
        parent_name = record.get('parent_name')
        display_name = f"{name}, {parent_name}" if parent_name else name
    elif scope_role == 'district':
        cursor.execute("""
            SELECT
                COALESCE(d.name, 'District') AS name,
                COALESCE(s.name, 'State') AS parent_name
            FROM districts d
            LEFT JOIN states s ON d.state_id = s.id
            WHERE d.id = %s
        """, (scope_id,))
        record = cursor.fetchone() or {}
        name = record.get('name') or 'District'
        parent_name = record.get('parent_name')
        display_name = f"{name}, {parent_name}" if parent_name else name
    elif scope_role == 'state':
        cursor.execute("""
            SELECT COALESCE(name, 'State') AS name
            FROM states
            WHERE id = %s
        """, (scope_id,))
        record = cursor.fetchone() or {}
        display_name = record.get('name') or 'State'

    cache[cache_key] = display_name
    return display_name


def _decorate_equipment_request_rows(request_rows, cursor):
    cache = {}
    decorated_rows = []
    for row in request_rows or []:
        requester_scope_name = row.get('requester_scope_name') or _equipment_scope_name(
            row.get('requester_scope_role'),
            row.get('requester_scope_id'),
            cursor,
            cache,
        )
        target_scope_name = _equipment_scope_name(
            row.get('target_scope_role'),
            row.get('target_scope_id'),
            cursor,
            cache,
        )
        decorated_rows.append({
            **row,
            'requester_scope_name': requester_scope_name,
            'target_scope_name': target_scope_name,
        })
    return decorated_rows


def _decorate_equipment_transaction_rows(transaction_rows, cursor):
    cache = {}
    decorated_rows = []
    for row in transaction_rows or []:
        decorated_rows.append({
            **row,
            'from_scope_name': _equipment_scope_name(
                row.get('from_scope_role'),
                row.get('from_scope_id'),
                cursor,
                cache,
            ),
            'to_scope_name': _equipment_scope_name(
                row.get('to_scope_role'),
                row.get('to_scope_id'),
                cursor,
                cache,
            ),
        })
    return decorated_rows


def _get_equipment_request_record(request_id, cursor):
    ensure_equipment_schema(cursor)
    cursor.execute("""
        SELECT
            id,
            requester_role,
            requester_id,
            requester_scope_role,
            requester_scope_id,
            target_role,
            target_scope_role,
            target_scope_id,
            item_key,
            item_name,
            category_name,
            quantity_requested,
            quantity_fulfilled,
            priority,
            reason,
            notes,
            status,
            created_at,
            updated_at,
            fulfilled_at
        FROM equipment_requests
        WHERE id = %s
        LIMIT 1
    """, (request_id,))
    return cursor.fetchone()


def _fetch_equipment_scope_transactions(scope_role, scope_id, cursor, limit=40):
    rows = fetch_transaction_rows(
        cursor,
        "WHERE (from_scope_role = %s AND from_scope_id = %s) OR (to_scope_role = %s AND to_scope_id = %s)",
        (scope_role, scope_id, scope_role, scope_id),
        limit=limit,
    )
    return _decorate_equipment_transaction_rows(rows, cursor)


def _build_worker_equipment_context(worker_profile, cursor):
    worker_id = worker_profile.get('id')
    inventory_rows = fetch_inventory(cursor, 'worker', worker_id)
    request_history = _decorate_equipment_request_rows(
        fetch_request_rows(
            cursor,
            "WHERE requester_role = %s AND requester_id = %s",
            ('worker', worker_id),
            limit=24,
        ),
        cursor,
    )
    transaction_rows = _fetch_equipment_scope_transactions('worker', worker_id, cursor, limit=24)
    return {
        'equipment_role': 'worker',
        'equipment_catalog': equipment_catalog_groups(),
        'inventory_rows': inventory_rows,
        'inventory_totals': inventory_summary(inventory_rows),
        'request_history': request_history,
        'transaction_rows': transaction_rows,
        'pending_request_count': sum(1 for row in request_history if row.get('status') == 'pending'),
    }


def _validate_registration_password(password):
    password = password or ''
    rules = [
        (len(password) >= 8, "at least 8 characters"),
        (re.search(r"[A-Z]", password) is not None, "one uppercase letter"),
        (re.search(r"[a-z]", password) is not None, "one lowercase letter"),
        (re.search(r"\d", password) is not None, "one number"),
        (re.search(r"[^A-Za-z0-9]", password) is not None, "one special character"),
    ]

    missing_rules = [label for is_valid, label in rules if not is_valid]
    return len(missing_rules) == 0, missing_rules


def _build_taluka_equipment_context(admin_profile, admin_id, cursor):
    taluka_id = admin_profile.get('taluka_id')
    inventory_rows = fetch_inventory(cursor, 'taluka', taluka_id)
    incoming_requests = _decorate_equipment_request_rows(
        fetch_request_rows(
            cursor,
            "WHERE target_role = %s AND target_scope_role = %s AND target_scope_id = %s",
            ('admin', 'taluka', taluka_id),
            limit=40,
        ),
        cursor,
    )
    outgoing_requests = _decorate_equipment_request_rows(
        fetch_request_rows(
            cursor,
            "WHERE requester_role = %s AND requester_id = %s AND target_role = %s",
            ('admin', admin_id, 'district_admin'),
            limit=24,
        ),
        cursor,
    )
    transaction_rows = _fetch_equipment_scope_transactions('taluka', taluka_id, cursor, limit=36)
    return {
        'equipment_role': 'taluka',
        'equipment_catalog': equipment_catalog_groups(),
        'inventory_rows': inventory_rows,
        'inventory_totals': inventory_summary(inventory_rows),
        'incoming_requests': incoming_requests,
        'incoming_pending_requests': [row for row in incoming_requests if row.get('status') == 'pending'],
        'incoming_closed_requests': [row for row in incoming_requests if row.get('status') != 'pending'][:12],
        'outgoing_requests': outgoing_requests,
        'transaction_rows': transaction_rows,
        'worker_options': _get_taluka_worker_options(taluka_id, cursor),
        'upstream_scope_name': admin_profile.get('district_name') or 'District Office',
    }


def _build_district_equipment_context(district_profile, admin_id, cursor):
    district_id = district_profile.get('district_id')
    inventory_rows = fetch_inventory(cursor, 'district', district_id)
    incoming_requests = _decorate_equipment_request_rows(
        fetch_request_rows(
            cursor,
            "WHERE target_role = %s AND target_scope_role = %s AND target_scope_id = %s",
            ('district_admin', 'district', district_id),
            limit=40,
        ),
        cursor,
    )
    outgoing_requests = _decorate_equipment_request_rows(
        fetch_request_rows(
            cursor,
            "WHERE requester_role = %s AND requester_id = %s AND target_role = %s",
            ('district_admin', admin_id, 'state_admin'),
            limit=24,
        ),
        cursor,
    )
    transaction_rows = _fetch_equipment_scope_transactions('district', district_id, cursor, limit=36)
    return {
        'equipment_role': 'district',
        'equipment_catalog': equipment_catalog_groups(),
        'inventory_rows': inventory_rows,
        'inventory_totals': inventory_summary(inventory_rows),
        'incoming_requests': incoming_requests,
        'incoming_pending_requests': [row for row in incoming_requests if row.get('status') == 'pending'],
        'incoming_closed_requests': [row for row in incoming_requests if row.get('status') != 'pending'][:12],
        'outgoing_requests': outgoing_requests,
        'transaction_rows': transaction_rows,
        'taluka_admin_options': _get_district_taluka_admin_options(district_id, cursor),
        'upstream_scope_name': district_profile.get('state_name') or 'State Reserve',
    }


def _build_state_equipment_context(state_profile, cursor):
    state_id = state_profile.get('state_id')
    inventory_rows = fetch_inventory(cursor, 'state', state_id)
    incoming_requests = _decorate_equipment_request_rows(
        fetch_request_rows(
            cursor,
            "WHERE target_role = %s AND target_scope_role = %s AND target_scope_id = %s",
            ('state_admin', 'state', state_id),
            limit=40,
        ),
        cursor,
    )
    transaction_rows = _fetch_equipment_scope_transactions('state', state_id, cursor, limit=36)
    return {
        'equipment_role': 'state',
        'equipment_catalog': equipment_catalog_groups(),
        'inventory_rows': inventory_rows,
        'inventory_totals': inventory_summary(inventory_rows),
        'incoming_requests': incoming_requests,
        'incoming_pending_requests': [row for row in incoming_requests if row.get('status') == 'pending'],
        'incoming_closed_requests': [row for row in incoming_requests if row.get('status') != 'pending'][:12],
        'transaction_rows': transaction_rows,
        'district_admin_options': _get_state_district_admin_options(state_id, cursor),
    }


def _set_session_scope_data(role, user, cursor):
    scope_data = {
        'scope_state_id': None,
        'scope_district_id': None,
        'scope_taluka_id': None,
        'scope_village_id': None,
    }

    if role == 'state_admin':
        profile = _get_state_admin_profile(user.get('id'), cursor)
        if profile:
            scope_data['scope_state_id'] = profile.get('state_id')
    elif role == 'district_admin':
        profile = _get_district_admin_profile(user.get('id'), cursor)
        if profile:
            scope_data['scope_state_id'] = profile.get('state_id')
            scope_data['scope_district_id'] = profile.get('district_id')
    elif role == 'admin':
        profile = _get_taluka_admin_profile(user.get('id'), cursor)
        if profile:
            scope_data['scope_state_id'] = profile.get('state_id')
            scope_data['scope_district_id'] = profile.get('district_id')
            scope_data['scope_taluka_id'] = profile.get('taluka_id')
    elif role == 'worker':
        profile = _get_worker_profile(user.get('id'), cursor)
        if profile:
            scope_data['scope_state_id'] = profile.get('state_id')
            scope_data['scope_district_id'] = profile.get('district_id')
            scope_data['scope_taluka_id'] = profile.get('taluka_id')
            scope_data['scope_village_id'] = profile.get('village_id')
    elif role == 'user':
        scope_data['scope_state_id'] = user.get('state_id')
        scope_data['scope_district_id'] = user.get('district_id')
        scope_data['scope_taluka_id'] = user.get('taluka_id')
        scope_data['scope_village_id'] = user.get('village_id')

    for key, value in scope_data.items():
        session[key] = value


def _save_staff_verification_files(role):
    aadhaar_photo = request.files.get('aadhaar_photo')
    selfie_photo = request.files.get('selfie_photo')
    if not staff_role_requires_verification(role):
        return {'aadhaar_photo_path': None, 'selfie_photo_path': None}, None
    if not aadhaar_photo or not getattr(aadhaar_photo, 'filename', ''):
        return None, "Aadhaar card photo is required for staff verification."
    if not selfie_photo or not getattr(selfie_photo, 'filename', ''):
        return None, "A live selfie photo is required for staff verification."
    return {
        'aadhaar_photo_path': save_verification_upload(aadhaar_photo, current_app.root_path, prefix='aadhaar'),
        'selfie_photo_path': save_verification_upload(selfie_photo, current_app.root_path, prefix='selfie'),
    }, None


def _is_today(value):
    if not value or not hasattr(value, 'date'):
        return False
    return value.date() == datetime.now().date()


def _taluka_scope_center(taluka_name):
    return MAP_PRESET_CENTERS.get(_slug_text(taluka_name))


def _build_worker_map_config(worker_profile):
    if not worker_profile:
        return {
            'center': None,
            'center_query': '',
            'scope_label': 'Assigned Area',
        }

    return {
        'center': _taluka_scope_center(worker_profile.get('taluka_name')),
        'center_query': _build_map_query(
            worker_profile.get('village_name'),
            worker_profile.get('taluka_name'),
            worker_profile.get('district_name'),
            'Maharashtra',
            'India',
        ),
        'scope_label': worker_profile.get('taluka_name') or worker_profile.get('village_name') or 'Assigned Area',
    }


def _build_taluka_map_payload(admin_profile, complaints, requests, assigned_tasks):
    taluka_name = admin_profile.get('taluka_name')
    district_name = admin_profile.get('district_name')
    markers = []

    for complaint in complaints or []:
        status_label = complaint.get('status') or 'Pending'
        status_key = complaint_status_class(status_label)
        filed_at = complaint.get('created_at')
        updated_at = complaint.get('updated_at') or complaint.get('resolved_at') or complaint.get('assigned_at') or filed_at
        map_query = _build_map_query(
            complaint.get('village'),
            complaint.get('taluka') or taluka_name,
            district_name,
            'Maharashtra',
            'India',
        )
        location_payload = build_location_payload(
            latitude=complaint.get('latitude'),
            longitude=complaint.get('longitude'),
            accuracy=complaint.get('location_accuracy_meters'),
            query=map_query,
            directions=True,
        )
        markers.append({
            'id': f"CMP-{complaint.get('id')}",
            'category': 'complaint',
            'category_label': 'Complaint',
            'title': complaint.get('title') or 'Complaint',
            'description': complaint.get('description') or 'Citizen complaint',
            'status': status_key,
            'status_label': status_label,
            'location_label': complaint.get('village') or taluka_name or 'Complaint Area',
            'map_query': map_query,
            'secondary_label': complaint.get('citizen_name') or 'Citizen',
            'assigned_to': complaint.get('worker_name') or 'Unassigned',
            'time_label': _format_timestamp(updated_at, include_time=True),
            'sort_at': updated_at.isoformat() if updated_at else '',
            'is_today': _is_today(filed_at),
            'is_open': status_key != 'completed',
            **location_payload,
        })

    for pickup_request in requests or []:
        status_key = (pickup_request.get('status') or 'pending').lower()
        created_at = pickup_request.get('created_at')
        map_query = _build_map_query(
            pickup_request.get('village_name'),
            taluka_name,
            district_name,
            'Maharashtra',
            'India',
        )
        location_payload = build_location_payload(
            latitude=pickup_request.get('latitude'),
            longitude=pickup_request.get('longitude'),
            accuracy=pickup_request.get('location_accuracy_meters'),
            query=map_query,
            directions=True,
        )
        markers.append({
            'id': f"REQ-{pickup_request.get('id')}",
            'category': 'request',
            'category_label': 'Pickup Request',
            'title': pickup_request.get('garbage_type') or 'Garbage Collection',
            'description': f"Pickup request from {pickup_request.get('citizen_name') or 'Citizen'}",
            'status': status_key,
            'status_label': status_key.replace('_', ' ').title(),
            'location_label': pickup_request.get('village_name') or taluka_name or 'Taluka Area',
            'map_query': map_query,
            'secondary_label': pickup_request.get('citizen_name') or 'Citizen',
            'assigned_to': pickup_request.get('worker_name') or 'Unassigned',
            'time_label': _format_timestamp(created_at, include_time=True),
            'sort_at': created_at.isoformat() if created_at else '',
            'is_today': _is_today(created_at),
            'is_open': status_key != 'completed',
            **location_payload,
        })

    for assigned_task in assigned_tasks or []:
        status_key = (assigned_task.get('status') or 'pending').lower()
        assigned_at = assigned_task.get('assigned_at')
        markers.append({
            'id': f"TASK-{assigned_task.get('id')}",
            'category': 'assigned_task',
            'category_label': 'Assigned Task',
            'title': assigned_task.get('description') or assigned_task.get('location_name') or 'Assigned task',
            'description': assigned_task.get('description') or 'Manual task assigned by taluka office',
            'status': status_key,
            'status_label': status_key.replace('_', ' ').title(),
            'location_label': assigned_task.get('location_name') or assigned_task.get('village_name') or taluka_name or 'Assigned Area',
            'map_query': _build_map_query(
                assigned_task.get('location_name') or assigned_task.get('village_name'),
                taluka_name,
                district_name,
                'Maharashtra',
                'India',
            ),
            'secondary_label': assigned_task.get('worker_name') or 'Worker',
            'assigned_to': assigned_task.get('worker_name') or 'Worker',
            'time_label': _format_timestamp(assigned_at, include_time=True),
            'sort_at': assigned_at.isoformat() if assigned_at else '',
            'is_today': _is_today(assigned_at),
            'is_open': status_key != 'completed',
        })

    markers.sort(key=lambda item: item.get('sort_at') or '', reverse=True)

    today_counts = {
        'complaints': sum(1 for item in markers if item.get('category') == 'complaint' and item.get('is_today')),
        'requests': sum(1 for item in markers if item.get('category') == 'request' and item.get('is_today')),
        'assigned_tasks': sum(1 for item in markers if item.get('category') == 'assigned_task' and item.get('is_today')),
    }
    today_counts['total'] = sum(today_counts.values())

    return {
        'scope_label': f"{taluka_name or 'Taluka'} Scope",
        'center': _taluka_scope_center(taluka_name),
        'center_query': _build_map_query(taluka_name, district_name, 'Maharashtra', 'India'),
        'markers': markers,
        'today_counts': today_counts,
        'open_total': sum(1 for item in markers if item.get('is_open')),
    }


def _format_user_complaint(complaint):
    status_text = normalize_complaint_status(complaint.get('status'))
    normalized_status = complaint_status_class(status_text)
    last_update = complaint.get('updated_at') or complaint.get('resolved_at') or complaint.get('assigned_at') or complaint.get('created_at')
    map_query = _build_map_query(
        complaint.get('village'),
        complaint.get('taluka'),
        complaint.get('district'),
        'Maharashtra',
        'India',
    )
    formatted = {
        'ticket_id': complaint.get('id'),
        'title': complaint.get('title') or 'Complaint',
        'village': complaint.get('village') or 'Unknown Village',
        'taluka': complaint.get('taluka') or 'Unknown Taluka',
        'priority': complaint.get('priority') or 'Normal',
        'status_text': status_text,
        'status_class': normalized_status,
        'progress_percent': complaint_progress_percent(status_text),
        'assigned_worker': complaint.get('worker_name') or 'Pending Assignment',
        'admin_name': complaint.get('admin_name') or 'Taluka Office',
        'date': _format_timestamp(complaint.get('created_at')),
        'last_update': _format_timestamp(last_update, include_time=True),
        **build_location_payload(
            latitude=complaint.get('latitude'),
            longitude=complaint.get('longitude'),
            accuracy=complaint.get('location_accuracy_meters'),
            query=map_query,
        ),
    }
    return _attach_complaint_photo_urls(formatted)


def _get_user_complaints(user_id, cursor, limit=5):
    complaint_columns = ensure_complaint_workflow_columns(cursor)
    original_photo_select = complaint_original_photo_select(complaint_columns)
    query = """
        SELECT
            c.id,
            c.title,
            c.district,
            c.taluka,
            c.village,
            c.latitude,
            c.longitude,
            c.location_accuracy_meters,
            {original_photo_select},
            c.resolution_photo_path,
            c.priority,
            c.status,
            c.created_at,
            c.assigned_at,
            c.updated_at,
            c.resolved_at,
            COALESCE(vw.name, 'Pending Assignment') AS worker_name,
            COALESCE(ta.name, 'Taluka Office') AS admin_name
        FROM complaints c
        LEFT JOIN village_workers vw ON c.worker_id = vw.id
        LEFT JOIN taluka_admins ta ON c.admin_id = ta.id
        WHERE c.user_id = %s
        ORDER BY c.created_at DESC
    """.format(original_photo_select=original_photo_select)
    if limit is not None:
        query += f"\n        LIMIT {max(int(limit), 0)}"

    cursor.execute(query, (user_id,))
    return [_format_user_complaint(complaint) for complaint in (cursor.fetchall() or [])]


def _get_taluka_complaints(admin_profile, cursor, limit=None):
    complaint_columns = ensure_complaint_workflow_columns(cursor)
    original_photo_select = complaint_original_photo_select(complaint_columns)

    query = """
        SELECT
            c.id,
            c.title,
            c.description,
            c.village,
            c.taluka,
            c.latitude,
            c.longitude,
            c.location_accuracy_meters,
            {original_photo_select},
            c.resolution_photo_path,
            c.priority,
            c.status,
            c.created_at,
            c.assigned_at,
            c.updated_at,
            c.resolved_at,
            c.worker_id,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(vw.name, 'Unassigned') AS worker_name
        FROM complaints c
        LEFT JOIN users u ON c.user_id = u.id
        LEFT JOIN village_workers vw ON c.worker_id = vw.id
        WHERE c.taluka = %s
        ORDER BY COALESCE(c.updated_at, c.assigned_at, c.created_at) DESC
    """.format(original_photo_select=original_photo_select)
    if limit is not None:
        query += f"\n        LIMIT {max(int(limit), 0)}"

    cursor.execute(query, (admin_profile.get('taluka_name'),))
    complaints = cursor.fetchall() or []
    for complaint in complaints:
        complaint['status'] = normalize_complaint_status(complaint.get('status'))
        complaint['status_class'] = complaint_status_class(complaint.get('status'))
        complaint['assigned_at_text'] = _format_timestamp(complaint.get('assigned_at'), include_time=True)
        complaint['updated_at_text'] = _format_timestamp(
            complaint.get('updated_at') or complaint.get('resolved_at') or complaint.get('assigned_at') or complaint.get('created_at'),
            include_time=True,
        )
        complaint.update(build_location_payload(
            latitude=complaint.get('latitude'),
            longitude=complaint.get('longitude'),
            accuracy=complaint.get('location_accuracy_meters'),
            query=_build_map_query(
                complaint.get('village'),
                complaint.get('taluka') or admin_profile.get('taluka_name'),
                admin_profile.get('district_name'),
                'Maharashtra',
                'India',
            ),
            directions=True,
        ))
        _attach_complaint_photo_urls(complaint)
    return complaints


def _get_worker_profile(worker_id, cursor):
    ensure_staff_verification_columns(cursor, 'village_workers')
    cursor.execute("""
        SELECT
            vw.id,
            vw.name,
            vw.email,
            vw.phone,
            vw.village_id,
            vw.vehicle_no,
            vw.status,
            vw.aadhaar_no,
            vw.aadhaar_photo_path,
            vw.selfie_photo_path,
            COALESCE(vw.verification_status, 'pending') AS verification_status,
            vw.verification_notes,
            vw.verified_by_role,
            vw.verified_by_id,
            vw.verified_at,
            vw.created_at,
            v.name AS village_name,
            t.name AS taluka_name,
            t.id AS taluka_id,
            d.id AS district_id,
            d.name AS district_name,
            s.id AS state_id,
            s.name AS state_name
        FROM village_workers vw
        LEFT JOIN villages v ON vw.village_id = v.id
        LEFT JOIN talukas t ON v.taluka_id = t.id
        LEFT JOIN districts d ON t.district_id = d.id
        LEFT JOIN states s ON d.state_id = s.id
        WHERE vw.id = %s
    """, (worker_id,))
    return attach_verification_urls(cursor.fetchone())


def _get_taluka_worker_options(taluka_id, cursor):
    ensure_staff_verification_columns(cursor, 'village_workers')
    cursor.execute("""
        SELECT
            vw.id,
            vw.name,
            vw.email,
            vw.phone,
            vw.vehicle_no,
            vw.status,
            vw.aadhaar_no,
            vw.aadhaar_photo_path,
            vw.selfie_photo_path,
            COALESCE(vw.verification_status, 'pending') AS verification_status,
            vw.verification_notes,
            vw.verified_by_role,
            vw.verified_by_id,
            vw.verified_at,
            vw.created_at,
            vw.village_id,
            COALESCE(v.name, 'Not assigned') AS village_name
        FROM village_workers vw
        LEFT JOIN villages v ON vw.village_id = v.id
        WHERE vw.village_id IN (SELECT id FROM villages WHERE taluka_id = %s)
        ORDER BY vw.name ASC
    """, (taluka_id,))
    workers = cursor.fetchall() or []
    if not workers:
        return []

    counts_by_worker = {worker['id']: _empty_worker_counts() for worker in workers}

    cursor.execute("""
        SELECT
            r.worker_id,
            COUNT(*) AS total_tasks,
            SUM(CASE WHEN r.status = 'completed' THEN 1 ELSE 0 END) AS completed_tasks,
            SUM(CASE WHEN r.status IN ('pending', 'in_progress') THEN 1 ELSE 0 END) AS active_tasks,
            SUM(CASE WHEN r.status = 'pending' THEN 1 ELSE 0 END) AS pending_tasks,
            SUM(CASE WHEN r.status = 'in_progress' THEN 1 ELSE 0 END) AS in_progress_tasks
        FROM requests r
        INNER JOIN village_workers vw ON vw.id = r.worker_id
        WHERE vw.village_id IN (SELECT id FROM villages WHERE taluka_id = %s)
        GROUP BY r.worker_id
    """, (taluka_id,))
    for row in cursor.fetchall() or []:
        counts = counts_by_worker.setdefault(row.get('worker_id'), _empty_worker_counts())
        _merge_worker_counts(counts, row)

    cursor.execute("""
        SELECT
            t.worker_id,
            COUNT(*) AS total_tasks,
            SUM(CASE WHEN t.status = 'completed' THEN 1 ELSE 0 END) AS completed_tasks,
            SUM(CASE WHEN t.status = 'pending' THEN 1 ELSE 0 END) AS active_tasks,
            SUM(CASE WHEN t.status = 'pending' THEN 1 ELSE 0 END) AS pending_tasks,
            0 AS in_progress_tasks
        FROM tasks t
        INNER JOIN village_workers vw ON vw.id = t.worker_id
        WHERE vw.village_id IN (SELECT id FROM villages WHERE taluka_id = %s)
        GROUP BY t.worker_id
    """, (taluka_id,))
    for row in cursor.fetchall() or []:
        counts = counts_by_worker.setdefault(row.get('worker_id'), _empty_worker_counts())
        _merge_worker_counts(counts, row)

    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            c.worker_id,
            COUNT(*) AS total_tasks,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) = 'completed' THEN 1 ELSE 0 END) AS completed_tasks,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) IN ('pending', 'assigned', 'in progress', 'in_progress') THEN 1 ELSE 0 END) AS active_tasks,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) IN ('pending', 'assigned') THEN 1 ELSE 0 END) AS pending_tasks,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) IN ('in progress', 'in_progress') THEN 1 ELSE 0 END) AS in_progress_tasks
        FROM complaints c
        INNER JOIN village_workers vw ON vw.id = c.worker_id
        WHERE vw.village_id IN (SELECT id FROM villages WHERE taluka_id = %s)
        GROUP BY c.worker_id
    """, (taluka_id,))
    for row in cursor.fetchall() or []:
        counts = counts_by_worker.setdefault(row.get('worker_id'), _empty_worker_counts())
        _merge_worker_counts(counts, row)

    for worker in workers:
        worker.update(counts_by_worker.get(worker['id'], _empty_worker_counts()))
        attach_verification_urls(worker)

    return workers


def _get_taluka_village_rows(taluka_id, cursor):
    cursor.execute("""
        SELECT
            v.id,
            v.name,
            COUNT(DISTINCT u.id) AS citizens,
            COUNT(DISTINCT vw.id) AS workers,
            COUNT(r.id) AS total_requests,
            SUM(CASE WHEN r.status = 'completed' THEN 1 ELSE 0 END) AS completed_requests
        FROM villages v
        LEFT JOIN users u ON u.village_id = v.id
        LEFT JOIN village_workers vw ON vw.village_id = v.id
        LEFT JOIN requests r ON r.user_id = u.id
        WHERE v.taluka_id = %s
        GROUP BY v.id, v.name
        ORDER BY v.name ASC
    """, (taluka_id,))
    villages = cursor.fetchall() or []

    for village in villages:
        total_requests = village.get('total_requests') or 0
        completed_requests = village.get('completed_requests') or 0
        village['progress'] = round((completed_requests / total_requests) * 100) if total_requests else 0

    return villages


def _get_taluka_worker_record(taluka_id, worker_id, cursor):
    cursor.execute("""
        SELECT
            vw.id,
            vw.name,
            vw.village_id,
            COALESCE(v.name, 'Assigned Area') AS village_name
        FROM village_workers vw
        LEFT JOIN villages v ON vw.village_id = v.id
        WHERE vw.id = %s AND v.taluka_id = %s
    """, (worker_id, taluka_id))
    return cursor.fetchone()


def _get_taluka_recent_manual_tasks(taluka_id, cursor, limit=8):
    limit_clause = ""
    if limit is not None:
        limit_clause = f"\n        LIMIT {max(int(limit), 0)}"

    cursor.execute(f"""
        SELECT
            t.id,
            t.worker_id,
            COALESCE(t.location_name, v.name, 'Assigned Area') AS location_name,
            COALESCE(t.description, 'Assigned task') AS description,
            COALESCE(t.priority, 'medium') AS priority,
            COALESCE(t.status, 'pending') AS status,
            t.assigned_at,
            COALESCE(vw.name, 'Unknown Worker') AS worker_name,
            COALESCE(v.name, 'Unknown Village') AS village_name
        FROM tasks t
        LEFT JOIN village_workers vw ON t.worker_id = vw.id
        LEFT JOIN villages v ON vw.village_id = v.id
        WHERE vw.village_id IN (SELECT id FROM villages WHERE taluka_id = %s)
        ORDER BY t.assigned_at DESC
        {limit_clause}
    """, (taluka_id,))
    return cursor.fetchall() or []


def _get_taluka_request_items(taluka_id, cursor, limit=None):
    limit_clause = ""
    if limit is not None:
        limit_clause = f"\n        LIMIT {max(int(limit), 0)}"

    ensure_request_location_columns(cursor)
    cursor.execute(f"""
        SELECT
            r.id,
            COALESCE(r.garbage_type, 'Garbage Collection') AS garbage_type,
            COALESCE(r.status, 'pending') AS status,
            COALESCE(r.amount, 0) AS amount,
            r.latitude,
            r.longitude,
            r.location_accuracy_meters,
            r.created_at,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(v.name, 'Unknown Village') AS village_name,
            COALESCE(vw.name, 'Unassigned') AS worker_name,
            r.worker_id
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages v ON u.village_id = v.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        WHERE COALESCE(u.taluka_id, uv.taluka_id) = %s
        ORDER BY r.created_at DESC
        {limit_clause}
    """, (taluka_id,))

    request_items = cursor.fetchall() or []
    for request_item in request_items:
        request_item['status'] = (request_item.get('status') or 'pending').lower()
        request_item['status_label'] = request_item['status'].replace('_', ' ').title()
        request_item.update(build_location_payload(
            latitude=request_item.get('latitude'),
            longitude=request_item.get('longitude'),
            accuracy=request_item.get('location_accuracy_meters'),
            query=_build_map_query(
                request_item.get('village_name'),
                request_item.get('taluka_name'),
                'Maharashtra',
                'India',
            ),
            directions=True,
        ))
    return request_items


def _get_worker_assignment_stats(worker_id, cursor):
    cursor.execute("""
        SELECT
            COUNT(*) AS assigned,
            SUM(CASE WHEN status = 'pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN status = 'in_progress' THEN 1 ELSE 0 END) AS in_progress,
            SUM(CASE WHEN status = 'completed' THEN 1 ELSE 0 END) AS completed,
            SUM(CASE WHEN status = 'completed' AND DATE(created_at) = CURDATE() THEN 1 ELSE 0 END) AS completed_today,
            COALESCE(SUM(CASE WHEN MONTH(created_at) = MONTH(CURDATE()) AND YEAR(created_at) = YEAR(CURDATE()) THEN amount ELSE 0 END), 0) AS monthly_earnings,
            COALESCE(SUM(CASE WHEN DATE(created_at) = CURDATE() THEN amount ELSE 0 END), 0) AS today_earnings
        FROM requests
        WHERE worker_id = %s
    """, (worker_id,))
    request_stats = cursor.fetchone() or {}

    cursor.execute("""
        SELECT
            COUNT(*) AS assigned,
            SUM(CASE WHEN status = 'pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN status = 'completed' THEN 1 ELSE 0 END) AS completed
        FROM tasks
        WHERE worker_id = %s
    """, (worker_id,))
    manual_task_stats = cursor.fetchone() or {}

    assigned = (request_stats.get('assigned') or 0) + (manual_task_stats.get('assigned') or 0)
    pending = (request_stats.get('pending') or 0) + (manual_task_stats.get('pending') or 0)
    in_progress = request_stats.get('in_progress') or 0
    completed = (request_stats.get('completed') or 0) + (manual_task_stats.get('completed') or 0)
    completed_today = request_stats.get('completed_today') or 0

    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            COUNT(*) AS assigned,
            SUM(CASE WHEN LOWER(COALESCE(status, 'Pending')) IN ('pending', 'assigned') THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN LOWER(COALESCE(status, 'Pending')) IN ('in progress', 'in_progress') THEN 1 ELSE 0 END) AS in_progress,
            SUM(CASE WHEN LOWER(COALESCE(status, 'Pending')) = 'completed' THEN 1 ELSE 0 END) AS completed,
            SUM(
                CASE
                    WHEN LOWER(COALESCE(status, 'Pending')) = 'completed'
                     AND DATE(COALESCE(resolved_at, updated_at, created_at)) = CURDATE()
                    THEN 1
                    ELSE 0
                END
            ) AS completed_today
        FROM complaints
        WHERE worker_id = %s
    """, (worker_id,))
    complaint_stats = cursor.fetchone() or {}

    assigned += complaint_stats.get('assigned') or 0
    pending += complaint_stats.get('pending') or 0
    in_progress += complaint_stats.get('in_progress') or 0
    completed += complaint_stats.get('completed') or 0
    completed_today += complaint_stats.get('completed_today') or 0

    return {
        'assigned': assigned,
        'pending': pending,
        'in_progress': in_progress,
        'completed': completed,
        'completed_today': completed_today,
        'monthly_earnings': request_stats.get('monthly_earnings') or 0,
        'today_earnings': request_stats.get('today_earnings') or 0,
        'total_tasks': assigned,
        'pending_tasks': pending,
        'in_progress_tasks': in_progress,
        'completed_tasks': completed,
    }


def _get_worker_work_items(worker_id, cursor, worker_profile=None, status_filter='all', limit=None):
    worker_village_id = worker_profile.get('village_id') if worker_profile else None
    worker_taluka_name = worker_profile.get('taluka_name') if worker_profile else None
    worker_district_name = worker_profile.get('district_name') if worker_profile else None

    request_status_clause = ""
    task_status_clause = ""
    complaint_status_clause = ""
    if status_filter == 'active':
        request_status_clause = "AND COALESCE(r.status, 'pending') IN ('pending', 'in_progress')"
        task_status_clause = "AND COALESCE(t.status, 'pending') = 'pending'"
        complaint_status_clause = "AND LOWER(COALESCE(c.status, 'Pending')) IN ('pending', 'assigned', 'in progress', 'in_progress')"
    elif status_filter == 'completed':
        request_status_clause = "AND COALESCE(r.status, 'pending') = 'completed'"
        task_status_clause = "AND COALESCE(t.status, 'pending') = 'completed'"
        complaint_status_clause = "AND LOWER(COALESCE(c.status, 'Pending')) = 'completed'"

    ensure_request_location_columns(cursor)
    cursor.execute(f"""
        SELECT
            r.id,
            COALESCE(r.garbage_type, 'Garbage Collection') AS garbage_type,
            COALESCE(r.status, 'pending') AS status,
            COALESCE(r.amount, 0) AS amount,
            r.latitude,
            r.longitude,
            r.location_accuracy_meters,
            r.created_at,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(v.name, worker_v.name, 'Assigned Area') AS area_name
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages v ON u.village_id = v.id
        LEFT JOIN villages worker_v ON worker_v.id = %s
        WHERE r.worker_id = %s
        {request_status_clause}
    """, (worker_village_id, worker_id))
    request_items = []
    for row in cursor.fetchall() or []:
        status = (row.get('status') or 'pending').lower()
        map_query = _build_map_query(
            row.get('area_name'),
            worker_taluka_name,
            worker_district_name,
            'Maharashtra',
            'India',
        )
        request_items.append({
            'id': row.get('id'),
            'ticket_id': f"REQ-{row.get('id')}",
            'source': 'request',
            'source_label': 'Citizen Request',
            'citizen_name': row.get('citizen_name') or 'Citizen',
            'area_name': row.get('area_name') or 'Assigned Area',
            'location_name': row.get('area_name') or 'Assigned Area',
            'description': row.get('garbage_type') or 'Garbage Collection',
            'garbage_type': row.get('garbage_type') or 'Garbage Collection',
            'priority': 'medium',
            'priority_label': 'Medium',
            'status': status,
            'status_class': status.replace('_', '-'),
            'status_label': status.replace('_', ' ').title(),
            'amount': row.get('amount') or 0,
            'created_at': row.get('created_at'),
            'map_query': map_query,
            'can_start': status == 'pending',
            'can_complete': status in ('pending', 'in_progress'),
            'is_started': status == 'in_progress',
            **build_location_payload(
                latitude=row.get('latitude'),
                longitude=row.get('longitude'),
                accuracy=row.get('location_accuracy_meters'),
                query=map_query,
                directions=True,
            ),
        })

    cursor.execute(f"""
        SELECT
            t.id,
            COALESCE(t.location_name, v.name, 'Assigned Area') AS location_name,
            COALESCE(t.description, 'Assigned task') AS description,
            COALESCE(t.status, 'pending') AS status,
            COALESCE(t.priority, 'medium') AS priority,
            t.assigned_at,
            COALESCE(v.name, 'Assigned Area') AS village_name
        FROM tasks t
        LEFT JOIN village_workers vw ON t.worker_id = vw.id
        LEFT JOIN villages v ON vw.village_id = v.id
        WHERE t.worker_id = %s
        {task_status_clause}
    """, (worker_id,))
    manual_task_items = []
    for row in cursor.fetchall() or []:
        status = (row.get('status') or 'pending').lower()
        priority = (row.get('priority') or 'medium').lower()
        if priority not in ('low', 'medium', 'high'):
            priority = 'medium'
        manual_task_items.append({
            'id': row.get('id'),
            'ticket_id': f"TASK-{row.get('id')}",
            'source': 'manual_task',
            'source_label': 'Admin Task',
            'citizen_name': 'Taluka Admin',
            'area_name': row.get('location_name') or row.get('village_name') or 'Assigned Area',
            'location_name': row.get('location_name') or row.get('village_name') or 'Assigned Area',
            'description': row.get('description') or 'Assigned task',
            'garbage_type': row.get('description') or 'Assigned task',
            'priority': priority,
            'priority_label': priority.title(),
            'status': status,
            'status_class': status.replace('_', '-'),
            'status_label': status.replace('_', ' ').title(),
            'amount': None,
            'created_at': row.get('assigned_at'),
            'map_query': _build_map_query(
                row.get('location_name') or row.get('village_name'),
                worker_taluka_name,
                worker_district_name,
                'Maharashtra',
                'India',
            ),
            'can_start': False,
            'can_complete': status == 'pending',
            'is_started': True,
        })

    complaint_columns = ensure_complaint_workflow_columns(cursor)
    original_photo_select = complaint_original_photo_select(complaint_columns)
    cursor.execute(f"""
        SELECT
            c.id,
            c.title,
            c.description,
            c.village,
            c.taluka,
            c.latitude,
            c.longitude,
            c.location_accuracy_meters,
            {original_photo_select},
            c.resolution_photo_path,
            c.status,
            c.priority,
            c.created_at,
            c.assigned_at,
            c.updated_at,
            c.resolved_at,
            COALESCE(u.name, 'Citizen') AS citizen_name
        FROM complaints c
        LEFT JOIN users u ON c.user_id = u.id
        WHERE c.worker_id = %s
        {complaint_status_clause}
    """, (worker_id,))
    complaint_items = []
    for row in cursor.fetchall() or []:
        status_label = normalize_complaint_status(row.get('status'))
        status_key = complaint_status_key(row.get('status'))
        priority_label = (row.get('priority') or 'Normal').strip().title()
        priority_key = priority_label.lower()
        if priority_key in ('urgent', 'high'):
            task_priority = 'high'
        elif priority_key == 'low':
            task_priority = 'low'
        else:
            task_priority = 'medium'
        complaint_area = ', '.join([part for part in [row.get('village'), row.get('taluka')] if part]) or 'Complaint Area'
        map_query = _build_map_query(
            row.get('village'),
            row.get('taluka') or worker_taluka_name,
            worker_district_name,
            'Maharashtra',
            'India',
        )
        complaint_item = {
            'id': row.get('id'),
            'ticket_id': f"CMP-{row.get('id')}",
            'source': 'complaint',
            'source_label': 'Complaint',
            'citizen_name': row.get('citizen_name') or 'Citizen',
            'area_name': complaint_area,
            'location_name': row.get('village') or row.get('taluka') or 'Complaint Area',
            'description': row.get('description') or row.get('title') or 'Complaint',
            'garbage_type': row.get('title') or 'Complaint',
            'priority': task_priority,
            'priority_label': priority_label,
            'status': status_key,
            'status_class': complaint_status_class(status_label),
            'status_label': status_label,
            'amount': None,
            'created_at': row.get('resolved_at') or row.get('updated_at') or row.get('assigned_at') or row.get('created_at'),
            'map_query': map_query,
            'can_start': status_key in ('pending', 'assigned'),
            'can_complete': status_key in ('pending', 'assigned', 'in_progress'),
            'is_started': status_key == 'in_progress',
            **build_location_payload(
                latitude=row.get('latitude'),
                longitude=row.get('longitude'),
                accuracy=row.get('location_accuracy_meters'),
                query=map_query,
                directions=True,
            ),
        }
        complaint_items.append(_attach_complaint_photo_urls(complaint_item))

    work_items = request_items + manual_task_items + complaint_items
    work_items.sort(key=lambda item: item.get('created_at') or datetime.min, reverse=True)

    if status_filter == 'all':
        status_order = {'pending': 0, 'assigned': 1, 'in_progress': 2, 'completed': 3}
        work_items.sort(key=lambda item: status_order.get(item.get('status'), 3))

    if limit is not None:
        work_items = work_items[:max(int(limit), 0)]

    return work_items


def _get_taluka_admin_profile(admin_id, cursor):
    ensure_staff_verification_columns(cursor, 'taluka_admins')
    cursor.execute("""
        SELECT
            ta.id,
            ta.name,
            ta.email,
            ta.phone,
            ta.taluka_id,
            ta.aadhaar_no,
            ta.aadhaar_photo_path,
            ta.selfie_photo_path,
            COALESCE(ta.verification_status, 'pending') AS verification_status,
            ta.verification_notes,
            ta.verified_by_role,
            ta.verified_by_id,
            ta.verified_at,
            ta.created_at,
            t.name AS taluka_name,
            d.id AS district_id,
            d.name AS district_name,
            s.id AS state_id,
            s.name AS state_name
        FROM taluka_admins ta
        LEFT JOIN talukas t ON ta.taluka_id = t.id
        LEFT JOIN districts d ON t.district_id = d.id
        LEFT JOIN states s ON d.state_id = s.id
        WHERE ta.id = %s
    """, (admin_id,))
    return attach_verification_urls(cursor.fetchone())


def _get_taluka_overview_data(admin_profile, cursor):
    taluka_id = admin_profile.get('taluka_id')
    taluka_name = admin_profile.get('taluka_name')

    cursor.execute("SELECT COUNT(*) AS total FROM villages WHERE taluka_id = %s", (taluka_id,))
    villages_count = (cursor.fetchone() or {}).get('total', 0)

    cursor.execute("""
        SELECT COUNT(*) AS total
        FROM users
        WHERE COALESCE(taluka_id, 0) = %s
           OR village_id IN (SELECT id FROM villages WHERE taluka_id = %s)
    """, (taluka_id, taluka_id))
    citizens_count = (cursor.fetchone() or {}).get('total', 0)

    cursor.execute("""
        SELECT COUNT(*) AS total
        FROM village_workers
        WHERE village_id IN (SELECT id FROM villages WHERE taluka_id = %s)
    """, (taluka_id,))
    workers_count = (cursor.fetchone() or {}).get('total', 0)

    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN LOWER(status) = 'pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN LOWER(status) IN ('assigned', 'in progress', 'in_progress') THEN 1 ELSE 0 END) AS in_progress,
            SUM(CASE WHEN LOWER(status) = 'completed' THEN 1 ELSE 0 END) AS completed
        FROM complaints
        WHERE taluka = %s
    """, (taluka_name,))
    complaints_stats = cursor.fetchone() or {}

    cursor.execute("""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN status = 'pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN status = 'in_progress' THEN 1 ELSE 0 END) AS in_progress,
            SUM(CASE WHEN status = 'completed' THEN 1 ELSE 0 END) AS completed,
            COALESCE(SUM(CASE WHEN status = 'completed' THEN amount ELSE 0 END), 0) AS completed_value
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        WHERE COALESCE(u.taluka_id, uv.taluka_id) = %s
    """, (taluka_id,))
    requests_stats = cursor.fetchone() or {}

    dashboard_stats = {
        'villages': villages_count or 0,
        'citizens': citizens_count or 0,
        'workers': workers_count or 0,
        'complaints_total': complaints_stats.get('total') or 0,
        'complaints_pending': complaints_stats.get('pending') or 0,
        'complaints_in_progress': complaints_stats.get('in_progress') or 0,
        'complaints_completed': complaints_stats.get('completed') or 0,
        'requests_total': requests_stats.get('total') or 0,
        'requests_pending': requests_stats.get('pending') or 0,
        'requests_in_progress': requests_stats.get('in_progress') or 0,
        'requests_completed': requests_stats.get('completed') or 0,
        'completed_value': requests_stats.get('completed_value') or 0,
    }

    recent_complaints = _get_taluka_complaints(admin_profile, cursor, limit=6)
    map_complaints = _get_taluka_complaints(admin_profile, cursor)

    cursor.execute("""
        SELECT
            v.id,
            v.name,
            COUNT(DISTINCT u.id) AS citizens,
            COUNT(DISTINCT vw.id) AS workers,
            COUNT(r.id) AS total_requests,
            SUM(CASE WHEN r.status = 'completed' THEN 1 ELSE 0 END) AS completed_requests
        FROM villages v
        LEFT JOIN users u ON u.village_id = v.id
        LEFT JOIN village_workers vw ON vw.village_id = v.id
        LEFT JOIN requests r ON r.user_id = u.id
        WHERE v.taluka_id = %s
        GROUP BY v.id, v.name
        ORDER BY total_requests DESC, v.name ASC
        LIMIT 8
    """, (taluka_id,))
    village_rows = cursor.fetchall() or []
    village_progress = []
    for row in village_rows:
        total_requests = row.get('total_requests') or 0
        completed_requests = row.get('completed_requests') or 0
        progress = round((completed_requests / total_requests) * 100) if total_requests else 0
        village_progress.append({
            'id': row.get('id'),
            'name': row.get('name'),
            'citizens': row.get('citizens') or 0,
            'workers': row.get('workers') or 0,
            'total_requests': total_requests,
            'completed_requests': completed_requests,
            'progress': progress,
        })

    taluka_workers = _get_taluka_worker_options(taluka_id, cursor)
    worker_summary = sorted(
        taluka_workers,
        key=lambda item: (
            -(item.get('completed_tasks') or 0),
            -(item.get('total_tasks') or 0),
            (item.get('name') or '').lower(),
        )
    )[:6]

    recent_requests = _get_taluka_request_items(taluka_id, cursor, limit=8)
    map_requests = _get_taluka_request_items(taluka_id, cursor)
    map_assigned_tasks = _get_taluka_recent_manual_tasks(taluka_id, cursor, limit=None)

    available_workers = [
        {
            'id': worker.get('id'),
            'name': worker.get('name'),
            'village_id': worker.get('village_id'),
            'village_name': worker.get('village_name'),
            'status': worker.get('status'),
        }
        for worker in taluka_workers
    ]

    chart_data = {
        'complaints': [
            dashboard_stats['complaints_pending'],
            dashboard_stats['complaints_in_progress'],
            dashboard_stats['complaints_completed'],
        ],
        'requests': [
            dashboard_stats['requests_pending'],
            dashboard_stats['requests_in_progress'],
            dashboard_stats['requests_completed'],
        ],
        'villages': [item['name'] for item in village_progress[:5]],
        'village_completion': [item['progress'] for item in village_progress[:5]],
    }

    return {
        'dashboard_stats': dashboard_stats,
        'recent_complaints': recent_complaints,
        'village_progress': village_progress,
        'worker_summary': worker_summary,
        'recent_requests': recent_requests,
        'available_workers': available_workers,
        'chart_data': chart_data,
        'taluka_map': _build_taluka_map_payload(admin_profile, map_complaints, map_requests, map_assigned_tasks),
    }


def _district_performance_band(completion_rate, open_work):
    if completion_rate >= 85 and open_work <= 5:
        return 'excellent', 'Excellent'
    if completion_rate >= 70 and open_work <= 14:
        return 'good', 'Stable'
    if completion_rate >= 50:
        return 'attention', 'Needs Attention'
    return 'critical', 'Critical'


def _district_hotspot_band(open_complaints, high_priority):
    pressure_score = (open_complaints or 0) * 2 + (high_priority or 0)
    if pressure_score >= 10:
        return 'high', 'High Pressure'
    if pressure_score >= 5:
        return 'medium', 'Watch Closely'
    return 'low', 'Stable'


def _get_district_admin_profile(admin_id, cursor):
    ensure_staff_verification_columns(cursor, 'district_admins')
    cursor.execute("""
        SELECT
            da.id,
            da.name,
            da.email,
            da.phone,
            da.district_id,
            da.aadhaar_no,
            da.aadhaar_photo_path,
            da.selfie_photo_path,
            COALESCE(da.verification_status, 'pending') AS verification_status,
            da.verification_notes,
            da.verified_by_role,
            da.verified_by_id,
            da.verified_at,
            da.created_at,
            d.name AS district_name,
            s.id AS state_id,
            s.name AS state_name
        FROM district_admins da
        LEFT JOIN districts d ON da.district_id = d.id
        LEFT JOIN states s ON d.state_id = s.id
        WHERE da.id = %s
    """, (admin_id,))
    return attach_verification_urls(cursor.fetchone())


def _get_state_admin_profile(admin_id, cursor):
    ensure_state_admin_table(cursor)
    cursor.execute("""
        SELECT
            sa.id,
            sa.name,
            sa.email,
            sa.phone,
            sa.state_id,
            sa.created_at,
            s.name AS state_name
        FROM state_admins sa
        LEFT JOIN states s ON sa.state_id = s.id
        WHERE sa.id = %s
    """, (admin_id,))
    return cursor.fetchone()


def _get_district_complaints(admin_profile, cursor, limit=None):
    complaint_columns = ensure_complaint_workflow_columns(cursor)
    original_photo_select = complaint_original_photo_select(complaint_columns)

    query = """
        SELECT
            c.id,
            c.title,
            c.description,
            c.village,
            c.taluka,
            c.latitude,
            c.longitude,
            c.location_accuracy_meters,
            {original_photo_select},
            c.resolution_photo_path,
            c.priority,
            c.status,
            c.created_at,
            c.assigned_at,
            c.updated_at,
            c.resolved_at,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(vw.name, 'Unassigned') AS worker_name,
            COALESCE(ta.name, 'Taluka Office') AS admin_name,
            COALESCE(TIMESTAMPDIFF(HOUR, COALESCE(c.updated_at, c.assigned_at, c.created_at), NOW()), 0) AS hours_waiting
        FROM complaints c
        LEFT JOIN users u ON c.user_id = u.id
        LEFT JOIN village_workers vw ON c.worker_id = vw.id
        LEFT JOIN taluka_admins ta ON c.admin_id = ta.id
        WHERE c.district = %s
        ORDER BY COALESCE(c.updated_at, c.assigned_at, c.created_at) DESC
    """.format(original_photo_select=original_photo_select)
    if limit is not None:
        query += f"\n        LIMIT {max(int(limit), 0)}"

    cursor.execute(query, (admin_profile.get('district_name'),))
    complaints = cursor.fetchall() or []

    for complaint in complaints:
        status_text = normalize_complaint_status(complaint.get('status'))
        status_key = complaint_status_key(status_text)
        last_update = complaint.get('updated_at') or complaint.get('resolved_at') or complaint.get('assigned_at') or complaint.get('created_at')
        hours_waiting = complaint.get('hours_waiting') or 0
        complaint['status'] = status_text
        complaint['status_key'] = status_key
        complaint['status_class'] = status_key
        complaint['priority_key'] = _slug_text(complaint.get('priority') or 'normal')
        complaint['created_at_text'] = _format_timestamp(complaint.get('created_at'), include_time=True)
        complaint['updated_at_text'] = _format_timestamp(last_update, include_time=True)
        complaint['hours_waiting'] = hours_waiting
        complaint['age_label'] = f"{hours_waiting} hrs open"
        complaint['is_escalated'] = status_key != 'completed' and hours_waiting >= 48
        complaint['severity_class'] = 'critical' if hours_waiting >= 72 or complaint['priority_key'] == 'high' else 'attention'
        complaint.update(build_location_payload(
            latitude=complaint.get('latitude'),
            longitude=complaint.get('longitude'),
            accuracy=complaint.get('location_accuracy_meters'),
            query=_build_map_query(
                complaint.get('village'),
                complaint.get('taluka'),
                admin_profile.get('district_name'),
                'Maharashtra',
                'India',
            ),
            directions=True,
        ))
        _attach_complaint_photo_urls(complaint)

    return complaints


def _get_district_request_items(district_id, cursor, limit=None):
    limit_clause = ""
    if limit is not None:
        limit_clause = f"\n        LIMIT {max(int(limit), 0)}"

    ensure_request_location_columns(cursor)
    cursor.execute(f"""
        SELECT
            r.id,
            COALESCE(r.garbage_type, 'Garbage Collection') AS garbage_type,
            COALESCE(r.status, 'pending') AS status,
            COALESCE(r.amount, 0) AS amount,
            COALESCE(r.weight, 0) AS weight,
            r.latitude,
            r.longitude,
            r.location_accuracy_meters,
            r.created_at,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(uv.name, wv.name, 'Unknown Village') AS village_name,
            COALESCE(ut.name, uvt.name, wt.name, 'Unknown Taluka') AS taluka_name,
            COALESCE(vw.name, 'Unassigned') AS worker_name,
            r.worker_id
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas uvt ON uv.taluka_id = uvt.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        LEFT JOIN villages wv ON vw.village_id = wv.id
        LEFT JOIN talukas wt ON wv.taluka_id = wt.id
        WHERE COALESCE(u.district_id, ut.district_id, uvt.district_id, wt.district_id) = %s
        ORDER BY r.created_at DESC
        {limit_clause}
    """, (district_id,))

    request_items = cursor.fetchall() or []
    for request_item in request_items:
        request_item['status'] = (request_item.get('status') or 'pending').lower()
        request_item['status_label'] = request_item['status'].replace('_', ' ').title()
        request_item['created_at_text'] = _format_timestamp(request_item.get('created_at'), include_time=True)
        request_item.update(build_location_payload(
            latitude=request_item.get('latitude'),
            longitude=request_item.get('longitude'),
            accuracy=request_item.get('location_accuracy_meters'),
            query=_build_map_query(
                request_item.get('village_name'),
                request_item.get('taluka_name'),
                'Maharashtra',
                'India',
            ),
            directions=True,
        ))
    return request_items


def _get_district_recent_manual_tasks(district_id, cursor, limit=10):
    limit_clause = ""
    if limit is not None:
        limit_clause = f"\n        LIMIT {max(int(limit), 0)}"

    cursor.execute(f"""
        SELECT
            t.id,
            t.worker_id,
            COALESCE(t.location_name, v.name, 'Assigned Area') AS location_name,
            COALESCE(t.description, 'Assigned task') AS description,
            COALESCE(t.priority, 'medium') AS priority,
            COALESCE(t.status, 'pending') AS status,
            t.assigned_at,
            COALESCE(vw.name, 'Unknown Worker') AS worker_name,
            COALESCE(v.name, 'Unknown Village') AS village_name,
            COALESCE(tl.name, 'Unknown Taluka') AS taluka_name
        FROM tasks t
        LEFT JOIN village_workers vw ON t.worker_id = vw.id
        LEFT JOIN villages v ON vw.village_id = v.id
        LEFT JOIN talukas tl ON v.taluka_id = tl.id
        WHERE tl.district_id = %s
        ORDER BY t.assigned_at DESC
        {limit_clause}
    """, (district_id,))

    tasks = cursor.fetchall() or []
    for task in tasks:
        task['status'] = (task.get('status') or 'pending').lower()
        task['status_label'] = task['status'].replace('_', ' ').title()
        task['priority_key'] = _slug_text(task.get('priority') or 'medium')
        task['assigned_at_text'] = _format_timestamp(task.get('assigned_at'), include_time=True)
    return tasks


def _get_district_worker_summary(district_id, cursor, limit=None):
    cursor.execute("""
        SELECT
            vw.id,
            vw.name,
            vw.email,
            vw.phone,
            vw.vehicle_no,
            vw.status,
            vw.created_at,
            COALESCE(v.name, 'Not assigned') AS village_name,
            COALESCE(t.name, 'Unknown Taluka') AS taluka_name
        FROM village_workers vw
        LEFT JOIN villages v ON vw.village_id = v.id
        LEFT JOIN talukas t ON v.taluka_id = t.id
        WHERE t.district_id = %s
        ORDER BY vw.name ASC
    """, (district_id,))
    workers = cursor.fetchall() or []
    if not workers:
        return []

    counts_by_worker = {worker['id']: _empty_worker_counts() for worker in workers}

    cursor.execute("""
        SELECT
            r.worker_id,
            COUNT(*) AS total_tasks,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'completed' THEN 1 ELSE 0 END) AS completed_tasks,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) IN ('pending', 'in_progress') THEN 1 ELSE 0 END) AS active_tasks,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'pending' THEN 1 ELSE 0 END) AS pending_tasks,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'in_progress' THEN 1 ELSE 0 END) AS in_progress_tasks
        FROM requests r
        INNER JOIN village_workers vw ON vw.id = r.worker_id
        INNER JOIN villages v ON vw.village_id = v.id
        INNER JOIN talukas t ON v.taluka_id = t.id
        WHERE t.district_id = %s
        GROUP BY r.worker_id
    """, (district_id,))
    for row in cursor.fetchall() or []:
        counts = counts_by_worker.setdefault(row.get('worker_id'), _empty_worker_counts())
        _merge_worker_counts(counts, row)

    cursor.execute("""
        SELECT
            t.worker_id,
            COUNT(*) AS total_tasks,
            SUM(CASE WHEN t.status = 'completed' THEN 1 ELSE 0 END) AS completed_tasks,
            SUM(CASE WHEN t.status = 'pending' THEN 1 ELSE 0 END) AS active_tasks,
            SUM(CASE WHEN t.status = 'pending' THEN 1 ELSE 0 END) AS pending_tasks,
            0 AS in_progress_tasks
        FROM tasks t
        INNER JOIN village_workers vw ON vw.id = t.worker_id
        INNER JOIN villages v ON vw.village_id = v.id
        INNER JOIN talukas tl ON v.taluka_id = tl.id
        WHERE tl.district_id = %s
        GROUP BY t.worker_id
    """, (district_id,))
    for row in cursor.fetchall() or []:
        counts = counts_by_worker.setdefault(row.get('worker_id'), _empty_worker_counts())
        _merge_worker_counts(counts, row)

    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            c.worker_id,
            COUNT(*) AS total_tasks,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) = 'completed' THEN 1 ELSE 0 END) AS completed_tasks,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) IN ('pending', 'assigned', 'in progress', 'in_progress') THEN 1 ELSE 0 END) AS active_tasks,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) IN ('pending', 'assigned') THEN 1 ELSE 0 END) AS pending_tasks,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) IN ('in progress', 'in_progress') THEN 1 ELSE 0 END) AS in_progress_tasks
        FROM complaints c
        INNER JOIN village_workers vw ON vw.id = c.worker_id
        INNER JOIN villages v ON vw.village_id = v.id
        INNER JOIN talukas t ON v.taluka_id = t.id
        WHERE t.district_id = %s
        GROUP BY c.worker_id
    """, (district_id,))
    for row in cursor.fetchall() or []:
        counts = counts_by_worker.setdefault(row.get('worker_id'), _empty_worker_counts())
        _merge_worker_counts(counts, row)

    for worker in workers:
        worker.update(counts_by_worker.get(worker['id'], _empty_worker_counts()))
        total_tasks = worker.get('total_tasks') or 0
        completed_tasks = worker.get('completed_tasks') or 0
        worker['completion_rate'] = round((completed_tasks / total_tasks) * 100) if total_tasks else 0
        worker['status_key'] = _slug_text(worker.get('status') or 'active').replace(' ', '_')

    workers.sort(
        key=lambda item: (
            -(item.get('completed_tasks') or 0),
            -(item.get('completion_rate') or 0),
            (item.get('name') or '').lower(),
        )
    )

    if limit is not None:
        return workers[:max(int(limit), 0)]
    return workers


def _get_district_taluka_performance(district_id, district_name, cursor):
    cursor.execute("""
        SELECT
            t.id,
            t.name,
            ta.id AS admin_id,
            COALESCE(ta.name, 'Not Assigned') AS admin_name
        FROM talukas t
        LEFT JOIN taluka_admins ta ON ta.taluka_id = t.id
        WHERE t.district_id = %s
        ORDER BY t.name ASC
    """, (district_id,))
    taluka_rows = cursor.fetchall() or []
    if not taluka_rows:
        return []

    taluka_map = {}
    name_to_id = {}
    for row in taluka_rows:
        taluka_map[row['id']] = {
            'id': row.get('id'),
            'name': row.get('name'),
            'admin_id': row.get('admin_id'),
            'admin_name': row.get('admin_name'),
            'villages': 0,
            'workers': 0,
            'citizens': 0,
            'total_requests': 0,
            'pending_requests': 0,
            'in_progress_requests': 0,
            'completed_requests': 0,
            'completed_value': 0,
            'total_complaints': 0,
            'complaints_pending': 0,
            'complaints_in_progress': 0,
            'complaints_completed': 0,
        }
        name_to_id[_slug_text(row.get('name'))] = row.get('id')

    cursor.execute("""
        SELECT taluka_id, COUNT(*) AS total
        FROM villages
        WHERE taluka_id IN (SELECT id FROM talukas WHERE district_id = %s)
        GROUP BY taluka_id
    """, (district_id,))
    for row in cursor.fetchall() or []:
        taluka = taluka_map.get(row.get('taluka_id'))
        if taluka:
            taluka['villages'] = row.get('total') or 0

    cursor.execute("""
        SELECT v.taluka_id, COUNT(*) AS total
        FROM village_workers vw
        INNER JOIN villages v ON vw.village_id = v.id
        WHERE v.taluka_id IN (SELECT id FROM talukas WHERE district_id = %s)
        GROUP BY v.taluka_id
    """, (district_id,))
    for row in cursor.fetchall() or []:
        taluka = taluka_map.get(row.get('taluka_id'))
        if taluka:
            taluka['workers'] = row.get('total') or 0

    cursor.execute("""
        SELECT
            COALESCE(u.taluka_id, v.taluka_id) AS taluka_id,
            COUNT(*) AS total
        FROM users u
        LEFT JOIN villages v ON u.village_id = v.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas vt ON v.taluka_id = vt.id
        WHERE COALESCE(u.district_id, ut.district_id, vt.district_id) = %s
        GROUP BY COALESCE(u.taluka_id, v.taluka_id)
    """, (district_id,))
    for row in cursor.fetchall() or []:
        taluka = taluka_map.get(row.get('taluka_id'))
        if taluka:
            taluka['citizens'] = row.get('total') or 0

    cursor.execute("""
        SELECT
            COALESCE(u.taluka_id, uv.taluka_id, wv.taluka_id) AS taluka_id,
            COUNT(*) AS total_requests,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'pending' THEN 1 ELSE 0 END) AS pending_requests,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'in_progress' THEN 1 ELSE 0 END) AS in_progress_requests,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'completed' THEN 1 ELSE 0 END) AS completed_requests,
            COALESCE(SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'completed' THEN r.amount ELSE 0 END), 0) AS completed_value
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas uvt ON uv.taluka_id = uvt.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        LEFT JOIN villages wv ON vw.village_id = wv.id
        LEFT JOIN talukas wt ON wv.taluka_id = wt.id
        WHERE COALESCE(u.district_id, ut.district_id, uvt.district_id, wt.district_id) = %s
        GROUP BY COALESCE(u.taluka_id, uv.taluka_id, wv.taluka_id)
    """, (district_id,))
    for row in cursor.fetchall() or []:
        taluka = taluka_map.get(row.get('taluka_id'))
        if taluka:
            taluka['total_requests'] = row.get('total_requests') or 0
            taluka['pending_requests'] = row.get('pending_requests') or 0
            taluka['in_progress_requests'] = row.get('in_progress_requests') or 0
            taluka['completed_requests'] = row.get('completed_requests') or 0
            taluka['completed_value'] = row.get('completed_value') or 0

    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            COALESCE(c.taluka, 'Unknown Taluka') AS taluka_name,
            COUNT(*) AS total_complaints,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) = 'pending' THEN 1 ELSE 0 END) AS complaints_pending,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) IN ('assigned', 'in progress', 'in_progress') THEN 1 ELSE 0 END) AS complaints_in_progress,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) = 'completed' THEN 1 ELSE 0 END) AS complaints_completed
        FROM complaints c
        WHERE c.district = %s
        GROUP BY COALESCE(c.taluka, 'Unknown Taluka')
    """, (district_name,))
    for row in cursor.fetchall() or []:
        taluka_id = name_to_id.get(_slug_text(row.get('taluka_name')))
        taluka = taluka_map.get(taluka_id)
        if taluka:
            taluka['total_complaints'] = row.get('total_complaints') or 0
            taluka['complaints_pending'] = row.get('complaints_pending') or 0
            taluka['complaints_in_progress'] = row.get('complaints_in_progress') or 0
            taluka['complaints_completed'] = row.get('complaints_completed') or 0

    performance_rows = []
    for taluka in taluka_map.values():
        total_services = (taluka.get('total_requests') or 0) + (taluka.get('total_complaints') or 0)
        completed_services = (taluka.get('completed_requests') or 0) + (taluka.get('complaints_completed') or 0)
        open_work = (
            (taluka.get('pending_requests') or 0)
            + (taluka.get('in_progress_requests') or 0)
            + (taluka.get('complaints_pending') or 0)
            + (taluka.get('complaints_in_progress') or 0)
        )
        completion_rate = round((completed_services / total_services) * 100) if total_services else 0
        band_class, band_label = _district_performance_band(completion_rate, open_work)
        taluka['total_services'] = total_services
        taluka['completed_services'] = completed_services
        taluka['open_work'] = open_work
        taluka['completion_rate'] = completion_rate
        taluka['band_class'] = band_class
        taluka['band_label'] = band_label
        taluka['admin_assigned'] = bool(taluka.get('admin_id'))
        performance_rows.append(taluka)

    performance_rows.sort(
        key=lambda item: (
            -(item.get('completion_rate') or 0),
            item.get('open_work') or 0,
            (item.get('name') or '').lower(),
        )
    )
    return performance_rows


def _get_district_hotspots(district_name, cursor, limit=6):
    cursor.execute(f"""
        SELECT
            COALESCE(c.village, 'Unknown Village') AS village_name,
            COALESCE(c.taluka, 'Unknown Taluka') AS taluka_name,
            COUNT(*) AS total_complaints,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) != 'completed' THEN 1 ELSE 0 END) AS open_complaints,
            SUM(CASE WHEN LOWER(COALESCE(c.priority, 'Normal')) = 'high' THEN 1 ELSE 0 END) AS high_priority_count
        FROM complaints c
        WHERE c.district = %s
        GROUP BY COALESCE(c.taluka, 'Unknown Taluka'), COALESCE(c.village, 'Unknown Village')
        ORDER BY open_complaints DESC, high_priority_count DESC, total_complaints DESC
        LIMIT {max(int(limit), 0)}
    """, (district_name,))

    hotspots = cursor.fetchall() or []
    for hotspot in hotspots:
        level_class, level_label = _district_hotspot_band(hotspot.get('open_complaints'), hotspot.get('high_priority_count'))
        hotspot['level_class'] = level_class
        hotspot['level_label'] = level_label
        hotspot['location_label'] = f"{hotspot.get('village_name')} • {hotspot.get('taluka_name')}"
    return hotspots


def _get_village_names_by_taluka(district_id, cursor):
    cursor.execute("""
        SELECT
            v.taluka_id,
            v.name
        FROM villages v
        INNER JOIN talukas t ON v.taluka_id = t.id
        WHERE t.district_id = %s
        ORDER BY t.name ASC, v.name ASC
    """, (district_id,))

    village_names_by_taluka = {}
    for row in cursor.fetchall() or []:
        taluka_id = row.get('taluka_id')
        if taluka_id is None:
            continue
        village_names_by_taluka.setdefault(taluka_id, []).append(row.get('name'))

    return village_names_by_taluka


def ensure_district_admin_task_table(cursor):
    try:
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS district_admin_tasks (
                id INT NOT NULL AUTO_INCREMENT,
                district_admin_id INT NOT NULL,
                taluka_admin_id INT NOT NULL,
                location_name VARCHAR(255) DEFAULT NULL,
                description TEXT NOT NULL,
                priority ENUM('low', 'medium', 'high') DEFAULT 'medium',
                status ENUM('pending', 'completed') DEFAULT 'pending',
                assigned_at TIMESTAMP NULL DEFAULT CURRENT_TIMESTAMP,
                completed_at TIMESTAMP NULL DEFAULT NULL,
                PRIMARY KEY (id),
                KEY idx_district_admin_tasks_district_admin_id (district_admin_id),
                KEY idx_district_admin_tasks_taluka_admin_id (taluka_admin_id),
                CONSTRAINT district_admin_tasks_district_admin_fk
                    FOREIGN KEY (district_admin_id) REFERENCES district_admins (id) ON DELETE CASCADE,
                CONSTRAINT district_admin_tasks_taluka_admin_fk
                    FOREIGN KEY (taluka_admin_id) REFERENCES taluka_admins (id) ON DELETE CASCADE
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci
        """)
    except mysql.connector.Error as err:
        if getattr(err, 'errno', None) != 1050:
            raise


def _get_district_taluka_rows(district_id, cursor):
    village_names_by_taluka = _get_village_names_by_taluka(district_id, cursor)

    cursor.execute("""
        SELECT
            t.id,
            t.name,
            ta.id AS admin_id,
            COALESCE(ta.name, 'Not Assigned') AS admin_name,
            COALESCE(ta.email, '') AS admin_email,
            COALESCE(ta.phone, '') AS admin_phone,
            COALESCE(
                (SELECT COUNT(*) FROM villages v WHERE v.taluka_id = t.id),
                0
            ) AS villages,
            COALESCE(
                (SELECT COUNT(*)
                 FROM users u
                 LEFT JOIN villages uv ON u.village_id = uv.id
                 WHERE COALESCE(u.taluka_id, uv.taluka_id) = t.id),
                0
            ) AS citizens,
            COALESCE(
                (SELECT COUNT(*)
                 FROM village_workers vw
                 LEFT JOIN villages vv ON vw.village_id = vv.id
                 WHERE vv.taluka_id = t.id),
                0
            ) AS workers,
            COALESCE(
                (SELECT COUNT(*)
                 FROM requests r
                 LEFT JOIN users u ON r.user_id = u.id
                 LEFT JOIN villages uv ON u.village_id = uv.id
                 WHERE COALESCE(u.taluka_id, uv.taluka_id) = t.id),
                0
            ) AS total_requests,
            COALESCE(
                (SELECT SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'completed' THEN 1 ELSE 0 END)
                 FROM requests r
                 LEFT JOIN users u ON r.user_id = u.id
                 LEFT JOIN villages uv ON u.village_id = uv.id
                 WHERE COALESCE(u.taluka_id, uv.taluka_id) = t.id),
                0
            ) AS completed_requests,
            COALESCE(
                (SELECT COUNT(*)
                 FROM complaints c
                 LEFT JOIN districts d ON d.id = t.district_id
                 WHERE c.taluka = t.name AND c.district = d.name),
                0
            ) AS total_complaints,
            COALESCE(
                (SELECT SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) != 'completed' THEN 1 ELSE 0 END)
                 FROM complaints c
                 LEFT JOIN districts d ON d.id = t.district_id
                 WHERE c.taluka = t.name AND c.district = d.name),
                0
            ) AS open_complaints
        FROM talukas t
        LEFT JOIN taluka_admins ta ON ta.taluka_id = t.id
        WHERE t.district_id = %s
        ORDER BY t.name ASC
    """, (district_id,))
    talukas = cursor.fetchall() or []

    for taluka in talukas:
        total_requests = taluka.get('total_requests') or 0
        completed_requests = taluka.get('completed_requests') or 0
        village_names = village_names_by_taluka.get(taluka.get('id'), [])
        taluka['progress'] = round((completed_requests / total_requests) * 100) if total_requests else 0
        taluka['villages_list'] = village_names[:8]
        taluka['villages_more_count'] = max(len(village_names) - 8, 0)
        taluka['admin_assigned'] = bool(taluka.get('admin_id'))

    return talukas


def _get_district_taluka_record(district_id, taluka_id, cursor):
    cursor.execute("""
        SELECT
            t.id,
            t.name,
            t.district_id,
            d.name AS district_name,
            ta.id AS admin_id,
            COALESCE(ta.name, 'Not Assigned') AS admin_name,
            COALESCE(ta.email, '') AS admin_email,
            COALESCE(ta.phone, '') AS admin_phone
        FROM talukas t
        LEFT JOIN districts d ON t.district_id = d.id
        LEFT JOIN taluka_admins ta ON ta.taluka_id = t.id
        WHERE t.id = %s AND t.district_id = %s
    """, (taluka_id, district_id))
    return cursor.fetchone()


def _get_district_unassigned_talukas(district_id, cursor):
    cursor.execute("""
        SELECT
            t.id,
            t.name
        FROM talukas t
        LEFT JOIN taluka_admins ta ON ta.taluka_id = t.id
        WHERE t.district_id = %s AND ta.id IS NULL
        ORDER BY t.name ASC
    """, (district_id,))
    return cursor.fetchall() or []


def _get_district_taluka_admin_record(district_id, taluka_admin_id, cursor):
    ensure_staff_verification_columns(cursor, 'taluka_admins')
    cursor.execute("""
        SELECT
            ta.id,
            ta.name,
            ta.email,
            ta.phone,
            ta.aadhaar_no,
            ta.aadhaar_photo_path,
            ta.selfie_photo_path,
            COALESCE(ta.verification_status, 'pending') AS verification_status,
            ta.verification_notes,
            ta.verified_by_role,
            ta.verified_by_id,
            ta.verified_at,
            ta.created_at,
            ta.taluka_id,
            t.name AS taluka_name,
            d.id AS district_id,
            d.name AS district_name,
            s.id AS state_id,
            s.name AS state_name
        FROM taluka_admins ta
        INNER JOIN talukas t ON ta.taluka_id = t.id
        LEFT JOIN districts d ON t.district_id = d.id
        LEFT JOIN states s ON d.state_id = s.id
        WHERE ta.id = %s AND t.district_id = %s
    """, (taluka_admin_id, district_id))
    return attach_verification_urls(cursor.fetchone())


def _get_district_taluka_admin_options(district_id, cursor):
    ensure_district_admin_task_table(cursor)
    ensure_staff_verification_columns(cursor, 'taluka_admins')
    village_names_by_taluka = _get_village_names_by_taluka(district_id, cursor)

    cursor.execute("""
        SELECT
            ta.id,
            ta.name,
            ta.email,
            ta.phone,
            ta.aadhaar_no,
            ta.aadhaar_photo_path,
            ta.selfie_photo_path,
            COALESCE(ta.verification_status, 'pending') AS verification_status,
            ta.verification_notes,
            ta.verified_by_role,
            ta.verified_by_id,
            ta.verified_at,
            ta.created_at,
            ta.taluka_id,
            t.name AS taluka_name,
            COALESCE(
                (SELECT COUNT(*) FROM villages v WHERE v.taluka_id = t.id),
                0
            ) AS villages,
            COALESCE(
                (SELECT COUNT(*)
                 FROM village_workers vw
                 LEFT JOIN villages v ON vw.village_id = v.id
                 WHERE v.taluka_id = t.id),
                0
            ) AS workers,
            COALESCE(
                (SELECT COUNT(*) FROM district_admin_tasks dat WHERE dat.taluka_admin_id = ta.id),
                0
            ) AS total_district_tasks,
            COALESCE(
                (SELECT COUNT(*) FROM district_admin_tasks dat WHERE dat.taluka_admin_id = ta.id AND dat.status = 'pending'),
                0
            ) AS pending_district_tasks,
            COALESCE(
                (SELECT COUNT(*) FROM district_admin_tasks dat WHERE dat.taluka_admin_id = ta.id AND dat.status = 'completed'),
                0
            ) AS completed_district_tasks
        FROM taluka_admins ta
        INNER JOIN talukas t ON ta.taluka_id = t.id
        WHERE t.district_id = %s
        ORDER BY t.name ASC, ta.name ASC
    """, (district_id,))
    taluka_admins = cursor.fetchall() or []

    for taluka_admin in taluka_admins:
        village_names = village_names_by_taluka.get(taluka_admin.get('taluka_id'), [])
        taluka_admin['villages_list'] = village_names[:8]
        taluka_admin['villages_more_count'] = max(len(village_names) - 8, 0)
        attach_verification_urls(taluka_admin)

    return taluka_admins


def _get_state_district_admin_options(state_id, cursor):
    ensure_staff_verification_columns(cursor, 'district_admins')
    cursor.execute("""
        SELECT
            da.id,
            da.name,
            da.email,
            da.phone,
            da.district_id,
            da.aadhaar_no,
            da.aadhaar_photo_path,
            da.selfie_photo_path,
            COALESCE(da.verification_status, 'pending') AS verification_status,
            da.verification_notes,
            da.verified_by_role,
            da.verified_by_id,
            da.verified_at,
            da.created_at,
            d.name AS district_name,
            COALESCE(
                (SELECT COUNT(*) FROM talukas t WHERE t.district_id = d.id),
                0
            ) AS talukas,
            COALESCE(
                (SELECT COUNT(*)
                 FROM village_workers vw
                 LEFT JOIN villages v ON vw.village_id = v.id
                 LEFT JOIN talukas t ON v.taluka_id = t.id
                 WHERE t.district_id = d.id),
                0
            ) AS workers
        FROM district_admins da
        INNER JOIN districts d ON da.district_id = d.id
        WHERE d.state_id = %s
        ORDER BY d.name ASC, da.name ASC
    """, (state_id,))
    district_admins = cursor.fetchall() or []
    for district_admin in district_admins:
        attach_verification_urls(district_admin)
    return district_admins


def _get_state_district_rows(state_id, cursor):
    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            d.id,
            d.name,
            COALESCE((SELECT COUNT(*) FROM talukas t WHERE t.district_id = d.id), 0) AS talukas,
            COALESCE((
                SELECT COUNT(*)
                FROM villages v
                LEFT JOIN talukas t ON v.taluka_id = t.id
                WHERE t.district_id = d.id
            ), 0) AS villages,
            COALESCE((
                SELECT COUNT(*)
                FROM users u
                LEFT JOIN villages v ON u.village_id = v.id
                LEFT JOIN talukas t ON v.taluka_id = t.id
                WHERE t.district_id = d.id
            ), 0) AS citizens,
            COALESCE((
                SELECT COUNT(*)
                FROM village_workers vw
                LEFT JOIN villages v ON vw.village_id = v.id
                LEFT JOIN talukas t ON v.taluka_id = t.id
                WHERE t.district_id = d.id
            ), 0) AS workers,
            COALESCE((
                SELECT COUNT(*)
                FROM complaints c
                WHERE c.district = d.name AND LOWER(COALESCE(c.status, 'Pending')) = 'completed'
            ), 0) AS completed_complaints,
            COALESCE((
                SELECT COUNT(*)
                FROM complaints c
                WHERE c.district = d.name AND LOWER(COALESCE(c.status, 'Pending')) != 'completed'
            ), 0) AS open_complaints
        FROM districts d
        WHERE d.state_id = %s
        ORDER BY d.name ASC
    """, (state_id,))
    district_rows = cursor.fetchall() or []
    for district_row in district_rows:
        total_complaints = (district_row.get('completed_complaints') or 0) + (district_row.get('open_complaints') or 0)
        district_row['completion_rate'] = round(((district_row.get('completed_complaints') or 0) / total_complaints) * 100) if total_complaints else 0
    return district_rows


def _get_state_escalations(state_id, cursor, limit=8):
    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            c.id,
            c.title,
            c.priority,
            c.status,
            c.village,
            c.taluka,
            c.district,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(vw.name, 'Unassigned') AS worker_name,
            COALESCE(TIMESTAMPDIFF(HOUR, COALESCE(c.updated_at, c.assigned_at, c.created_at), NOW()), 0) AS hours_waiting,
            COALESCE(c.updated_at, c.assigned_at, c.created_at) AS last_action_at
        FROM complaints c
        LEFT JOIN users u ON c.user_id = u.id
        LEFT JOIN village_workers vw ON c.worker_id = vw.id
        LEFT JOIN districts d ON d.name = c.district
        WHERE d.state_id = %s
          AND LOWER(COALESCE(c.status, 'Pending')) != 'completed'
        ORDER BY hours_waiting DESC, last_action_at DESC
        LIMIT %s
    """, (state_id, max(int(limit), 1)))
    return cursor.fetchall() or []


def _get_state_overview_data(state_profile, cursor):
    state_id = state_profile.get('state_id')

    cursor.execute("SELECT COUNT(*) AS total FROM districts WHERE state_id = %s", (state_id,))
    districts_count = (cursor.fetchone() or {}).get('total', 0)

    cursor.execute("""
        SELECT
            COUNT(DISTINCT t.id) AS talukas,
            COUNT(DISTINCT v.id) AS villages,
            COUNT(DISTINCT u.id) AS citizens,
            COUNT(DISTINCT vw.id) AS workers
        FROM districts d
        LEFT JOIN talukas t ON t.district_id = d.id
        LEFT JOIN villages v ON v.taluka_id = t.id
        LEFT JOIN users u ON u.village_id = v.id
        LEFT JOIN village_workers vw ON vw.village_id = v.id
        WHERE d.state_id = %s
    """, (state_id,))
    scope_counts = cursor.fetchone() or {}

    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) = 'completed' THEN 1 ELSE 0 END) AS completed,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) != 'completed' THEN 1 ELSE 0 END) AS open_items
        FROM complaints c
        LEFT JOIN districts d ON d.name = c.district
        WHERE d.state_id = %s
    """, (state_id,))
    complaint_counts = cursor.fetchone() or {}

    ensure_staff_verification_columns(cursor, 'district_admins')
    ensure_staff_verification_columns(cursor, 'taluka_admins')
    ensure_staff_verification_columns(cursor, 'village_workers')

    cursor.execute("""
        SELECT
            SUM(CASE WHEN COALESCE(da.verification_status, 'pending') = 'pending' THEN 1 ELSE 0 END) AS district_pending,
            SUM(CASE WHEN COALESCE(da.verification_status, 'pending') = 'approved' THEN 1 ELSE 0 END) AS district_verified
        FROM district_admins da
        INNER JOIN districts d ON da.district_id = d.id
        WHERE d.state_id = %s
    """, (state_id,))
    district_verification = cursor.fetchone() or {}

    cursor.execute("""
        SELECT
            SUM(CASE WHEN COALESCE(ta.verification_status, 'pending') = 'pending' THEN 1 ELSE 0 END) AS taluka_pending
        FROM taluka_admins ta
        INNER JOIN talukas t ON ta.taluka_id = t.id
        INNER JOIN districts d ON t.district_id = d.id
        WHERE d.state_id = %s
    """, (state_id,))
    taluka_verification = cursor.fetchone() or {}

    cursor.execute("""
        SELECT
            SUM(CASE WHEN COALESCE(vw.verification_status, 'pending') = 'pending' THEN 1 ELSE 0 END) AS worker_pending
        FROM village_workers vw
        LEFT JOIN villages v ON vw.village_id = v.id
        LEFT JOIN talukas t ON v.taluka_id = t.id
        LEFT JOIN districts d ON t.district_id = d.id
        WHERE d.state_id = %s
    """, (state_id,))
    worker_verification = cursor.fetchone() or {}

    total_complaints = complaint_counts.get('total') or 0
    completed_complaints = complaint_counts.get('completed') or 0

    return {
        'state_stats': {
            'districts': districts_count or 0,
            'talukas': scope_counts.get('talukas') or 0,
            'villages': scope_counts.get('villages') or 0,
            'citizens': scope_counts.get('citizens') or 0,
            'workers': scope_counts.get('workers') or 0,
            'complaints_total': total_complaints,
            'complaints_completed': completed_complaints,
            'complaints_open': complaint_counts.get('open_items') or 0,
            'completion_rate': round((completed_complaints / total_complaints) * 100) if total_complaints else 0,
            'district_pending': district_verification.get('district_pending') or 0,
            'district_verified': district_verification.get('district_verified') or 0,
            'taluka_pending': taluka_verification.get('taluka_pending') or 0,
            'worker_pending': worker_verification.get('worker_pending') or 0,
        },
        'district_admins': _get_state_district_admin_options(state_id, cursor),
        'district_rows': _get_state_district_rows(state_id, cursor),
        'state_escalations': _get_state_escalations(state_id, cursor),
        'announcement_feed': fetch_visible_announcements(cursor, 'state_admin', state_id=state_id, limit=6),
    }


def _parse_report_date(value):
    clean_value = (value or '').strip()
    if not clean_value:
        return None
    try:
        return datetime.strptime(clean_value, '%Y-%m-%d').date()
    except ValueError:
        return None


def _resolve_report_date_filters():
    quick_range = (request.args.get('range') or 'all').strip().lower()
    if quick_range not in REPORT_QUICK_RANGES:
        quick_range = 'all'

    date_from = _parse_report_date(request.args.get('date_from'))
    date_to = _parse_report_date(request.args.get('date_to'))

    if not date_from and not date_to and REPORT_QUICK_RANGES[quick_range]['days']:
        today = datetime.now().date()
        span_days = REPORT_QUICK_RANGES[quick_range]['days']
        date_to = today
        date_from = today - timedelta(days=span_days - 1)

    if date_from and date_to and date_from > date_to:
        raise ValueError("The start date must be before the end date.")

    if date_from and date_to:
        label = f"{date_from.strftime('%d %b %Y')} to {date_to.strftime('%d %b %Y')}"
    elif date_from:
        label = f"From {date_from.strftime('%d %b %Y')}"
    elif date_to:
        label = f"Up to {date_to.strftime('%d %b %Y')}"
    else:
        label = REPORT_QUICK_RANGES[quick_range]['label']

    return {
        'range_key': quick_range,
        'date_from': date_from,
        'date_to': date_to,
        'date_from_value': date_from.isoformat() if date_from else '',
        'date_to_value': date_to.isoformat() if date_to else '',
        'label': label,
    }


def _fetch_report_district_options(state_id, cursor):
    cursor.execute("""
        SELECT id, name
        FROM districts
        WHERE state_id = %s
        ORDER BY name ASC
    """, (state_id,))
    return cursor.fetchall() or []


def _fetch_report_taluka_options(state_id, cursor, district_id=None):
    params = [state_id]
    filters = ["d.state_id = %s"]
    if district_id:
        filters.append("d.id = %s")
        params.append(district_id)

    cursor.execute(f"""
        SELECT
            t.id,
            t.name,
            d.id AS district_id,
            d.name AS district_name
        FROM talukas t
        INNER JOIN districts d ON t.district_id = d.id
        WHERE {' AND '.join(filters)}
        ORDER BY t.name ASC
    """, tuple(params))
    return cursor.fetchall() or []


def _fetch_report_village_options(cursor, taluka_id):
    if not taluka_id:
        return []

    cursor.execute("""
        SELECT id, name
        FROM villages
        WHERE taluka_id = %s
        ORDER BY name ASC
    """, (taluka_id,))
    return cursor.fetchall() or []


def _report_geo_scope_clauses(scope, state_expr, district_expr, taluka_expr, village_expr):
    clauses = [f"{state_expr} = %s"]
    params = [scope.get('state_id')]

    if scope.get('district_id'):
        clauses.append(f"{district_expr} = %s")
        params.append(scope.get('district_id'))
    if scope.get('taluka_id'):
        clauses.append(f"{taluka_expr} = %s")
        params.append(scope.get('taluka_id'))
    if scope.get('village_id'):
        clauses.append(f"{village_expr} = %s")
        params.append(scope.get('village_id'))

    return clauses, params


def _apply_report_date_filter(clauses, params, column_name, date_filters):
    if date_filters.get('date_from'):
        clauses.append(f"DATE({column_name}) >= %s")
        params.append(date_filters.get('date_from'))
    if date_filters.get('date_to'):
        clauses.append(f"DATE({column_name}) <= %s")
        params.append(date_filters.get('date_to'))


def _resolve_report_scope(role, profile, cursor):
    selected_district_id = _safe_int(request.args.get('district_id'))
    selected_taluka_id = _safe_int(request.args.get('taluka_id'))
    selected_village_id = _safe_int(request.args.get('village_id'))

    scope = {
        'state_id': profile.get('state_id'),
        'state_name': profile.get('state_name'),
        'district_id': None,
        'district_name': None,
        'taluka_id': None,
        'taluka_name': None,
        'village_id': None,
        'village_name': None,
    }

    if role == 'admin':
        scope['district_id'] = profile.get('district_id')
        scope['district_name'] = profile.get('district_name')
        scope['taluka_id'] = profile.get('taluka_id')
        scope['taluka_name'] = profile.get('taluka_name')

        if selected_district_id and selected_district_id != scope['district_id']:
            raise PermissionError("Taluka admins can access reports only for their own district and taluka.")
        if selected_taluka_id and selected_taluka_id != scope['taluka_id']:
            raise PermissionError("Taluka admins can access reports only for their assigned taluka.")

        if selected_village_id:
            cursor.execute("""
                SELECT id, name
                FROM villages
                WHERE id = %s AND taluka_id = %s
            """, (selected_village_id, scope['taluka_id']))
            village_row = cursor.fetchone()
            if not village_row:
                raise PermissionError("That village is outside your taluka report scope.")
            scope['village_id'] = village_row.get('id')
            scope['village_name'] = village_row.get('name')

        return scope

    if role == 'district_admin':
        scope['district_id'] = profile.get('district_id')
        scope['district_name'] = profile.get('district_name')

        if selected_district_id and selected_district_id != scope['district_id']:
            raise PermissionError("District admins can access reports only inside their own district.")

        if selected_taluka_id:
            cursor.execute("""
                SELECT t.id, t.name
                FROM talukas t
                WHERE t.id = %s AND t.district_id = %s
            """, (selected_taluka_id, scope['district_id']))
            taluka_row = cursor.fetchone()
            if not taluka_row:
                raise PermissionError("That taluka is outside your district report scope.")
            scope['taluka_id'] = taluka_row.get('id')
            scope['taluka_name'] = taluka_row.get('name')

        if selected_village_id:
            cursor.execute("""
                SELECT
                    v.id,
                    v.name,
                    t.id AS taluka_id,
                    t.name AS taluka_name
                FROM villages v
                INNER JOIN talukas t ON v.taluka_id = t.id
                WHERE v.id = %s
                  AND t.district_id = %s
                  AND (%s IS NULL OR t.id = %s)
            """, (selected_village_id, scope['district_id'], scope.get('taluka_id'), scope.get('taluka_id')))
            village_row = cursor.fetchone()
            if not village_row:
                raise PermissionError("That village is outside your district report scope.")
            scope['village_id'] = village_row.get('id')
            scope['village_name'] = village_row.get('name')
            scope['taluka_id'] = village_row.get('taluka_id')
            scope['taluka_name'] = village_row.get('taluka_name')

        return scope

    if selected_district_id:
        cursor.execute("""
            SELECT id, name
            FROM districts
            WHERE id = %s AND state_id = %s
        """, (selected_district_id, scope['state_id']))
        district_row = cursor.fetchone()
        if not district_row:
            raise PermissionError("That district is outside your state report scope.")
        scope['district_id'] = district_row.get('id')
        scope['district_name'] = district_row.get('name')

    if selected_taluka_id:
        cursor.execute("""
            SELECT
                t.id,
                t.name,
                d.id AS district_id,
                d.name AS district_name
            FROM talukas t
            INNER JOIN districts d ON t.district_id = d.id
            WHERE t.id = %s
              AND d.state_id = %s
              AND (%s IS NULL OR d.id = %s)
        """, (selected_taluka_id, scope['state_id'], scope.get('district_id'), scope.get('district_id')))
        taluka_row = cursor.fetchone()
        if not taluka_row:
            raise PermissionError("That taluka is outside your state report scope.")
        scope['taluka_id'] = taluka_row.get('id')
        scope['taluka_name'] = taluka_row.get('name')
        scope['district_id'] = taluka_row.get('district_id')
        scope['district_name'] = taluka_row.get('district_name')

    if selected_village_id:
        cursor.execute("""
            SELECT
                v.id,
                v.name,
                t.id AS taluka_id,
                t.name AS taluka_name,
                d.id AS district_id,
                d.name AS district_name
            FROM villages v
            INNER JOIN talukas t ON v.taluka_id = t.id
            INNER JOIN districts d ON t.district_id = d.id
            WHERE v.id = %s
              AND d.state_id = %s
              AND (%s IS NULL OR d.id = %s)
              AND (%s IS NULL OR t.id = %s)
        """, (
            selected_village_id,
            scope['state_id'],
            scope.get('district_id'),
            scope.get('district_id'),
            scope.get('taluka_id'),
            scope.get('taluka_id'),
        ))
        village_row = cursor.fetchone()
        if not village_row:
            raise PermissionError("That village is outside your state report scope.")
        scope['village_id'] = village_row.get('id')
        scope['village_name'] = village_row.get('name')
        scope['taluka_id'] = village_row.get('taluka_id')
        scope['taluka_name'] = village_row.get('taluka_name')
        scope['district_id'] = village_row.get('district_id')
        scope['district_name'] = village_row.get('district_name')

    return scope


def _build_report_filter_options(role, profile, scope, cursor):
    options = {
        'districts': [],
        'talukas': [],
        'villages': [],
    }

    if role == 'state_admin':
        options['districts'] = _fetch_report_district_options(profile.get('state_id'), cursor)
        options['talukas'] = _fetch_report_taluka_options(
            profile.get('state_id'),
            cursor,
            district_id=scope.get('district_id'),
        )
        if scope.get('taluka_id'):
            options['villages'] = _fetch_report_village_options(cursor, scope.get('taluka_id'))
        return options

    if role == 'district_admin':
        options['talukas'] = _fetch_report_taluka_options(
            profile.get('state_id'),
            cursor,
            district_id=profile.get('district_id'),
        )
        if scope.get('taluka_id'):
            options['villages'] = _fetch_report_village_options(cursor, scope.get('taluka_id'))
        return options

    options['villages'] = _fetch_report_village_options(cursor, profile.get('taluka_id'))
    return options


def _report_scope_label(scope):
    if scope.get('village_name'):
        return f"{scope.get('village_name')} village"
    if scope.get('taluka_name'):
        return f"{scope.get('taluka_name')} taluka"
    if scope.get('district_name'):
        return f"{scope.get('district_name')} district"
    return f"{scope.get('state_name') or 'State'} scope"


def _report_export_query(scope, date_filters, dataset_key):
    payload = {
        'dataset': dataset_key,
        'range': date_filters.get('range_key'),
    }
    if scope.get('district_id'):
        payload['district_id'] = scope.get('district_id')
    if scope.get('taluka_id'):
        payload['taluka_id'] = scope.get('taluka_id')
    if scope.get('village_id'):
        payload['village_id'] = scope.get('village_id')
    if date_filters.get('date_from_value'):
        payload['date_from'] = date_filters.get('date_from_value')
    if date_filters.get('date_to_value'):
        payload['date_to'] = date_filters.get('date_to_value')
    return payload


def _report_card(label, value, tone, icon, detail=''):
    return {
        'label': label,
        'value': value,
        'tone': tone,
        'icon': icon,
        'detail': detail,
    }


def _fetch_report_scalar(cursor, query, params):
    cursor.execute(query, tuple(params))
    row = cursor.fetchone() or {}
    return next(iter(row.values()), 0)


def _build_summary_dataset(scope, date_filters, cursor):
    geo_clauses, geo_params = _report_geo_scope_clauses(scope, "d.state_id", "d.id", "t.id", "v.id")

    villages_total = _fetch_report_scalar(cursor, f"""
        SELECT COUNT(*) AS total
        FROM villages v
        INNER JOIN talukas t ON v.taluka_id = t.id
        INNER JOIN districts d ON t.district_id = d.id
        WHERE {' AND '.join(geo_clauses)}
    """, geo_params)

    citizen_clauses, citizen_params = _report_geo_scope_clauses(scope, "d.state_id", "d.id", "t.id", "v.id")
    _apply_report_date_filter(citizen_clauses, citizen_params, "u.created_at", date_filters)
    citizens_total = _fetch_report_scalar(cursor, f"""
        SELECT COUNT(*) AS total
        FROM users u
        LEFT JOIN villages v ON u.village_id = v.id
        LEFT JOIN talukas t ON COALESCE(u.taluka_id, v.taluka_id) = t.id
        LEFT JOIN districts d ON COALESCE(u.district_id, t.district_id) = d.id
        WHERE {' AND '.join(citizen_clauses)}
    """, citizen_params)

    worker_clauses, worker_params = _report_geo_scope_clauses(scope, "d.state_id", "d.id", "t.id", "v.id")
    _apply_report_date_filter(worker_clauses, worker_params, "vw.created_at", date_filters)
    workers_total = _fetch_report_scalar(cursor, f"""
        SELECT COUNT(*) AS total
        FROM village_workers vw
        INNER JOIN villages v ON vw.village_id = v.id
        INNER JOIN talukas t ON v.taluka_id = t.id
        INNER JOIN districts d ON t.district_id = d.id
        WHERE {' AND '.join(worker_clauses)}
    """, worker_params)

    complaint_clauses = ["d.state_id = %s"]
    complaint_params = [scope.get('state_id')]
    if scope.get('district_id'):
        complaint_clauses.append("d.id = %s")
        complaint_params.append(scope.get('district_id'))
    if scope.get('taluka_id'):
        complaint_clauses.append("t.id = %s")
        complaint_params.append(scope.get('taluka_id'))
    if scope.get('village_id'):
        complaint_clauses.append("v.id = %s")
        complaint_params.append(scope.get('village_id'))
    _apply_report_date_filter(complaint_clauses, complaint_params, "c.created_at", date_filters)
    ensure_complaint_workflow_columns(cursor)
    cursor.execute(f"""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) = 'completed' THEN 1 ELSE 0 END) AS completed,
            SUM(CASE WHEN LOWER(COALESCE(c.status, 'Pending')) != 'completed' THEN 1 ELSE 0 END) AS open_items
        FROM complaints c
        INNER JOIN districts d ON d.name = c.district
        LEFT JOIN talukas t ON t.name = c.taluka AND t.district_id = d.id
        LEFT JOIN villages v ON v.name = c.village AND v.taluka_id = t.id
        WHERE {' AND '.join(complaint_clauses)}
    """, tuple(complaint_params))
    complaint_totals = cursor.fetchone() or {}

    request_clauses, request_params = _report_geo_scope_clauses(
        scope,
        "d.state_id",
        "d.id",
        "COALESCE(ut.id, uvt.id, wt.id)",
        "COALESCE(uv.id, wv.id)",
    )
    _apply_report_date_filter(request_clauses, request_params, "r.created_at", date_filters)
    cursor.execute(f"""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'completed' THEN 1 ELSE 0 END) AS completed,
            COALESCE(SUM(COALESCE(r.amount, 0)), 0) AS total_amount
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas uvt ON uv.taluka_id = uvt.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        LEFT JOIN villages wv ON vw.village_id = wv.id
        LEFT JOIN talukas wt ON wv.taluka_id = wt.id
        INNER JOIN districts d ON COALESCE(u.district_id, ut.district_id, uvt.district_id, wt.district_id) = d.id
        WHERE {' AND '.join(request_clauses)}
    """, tuple(request_params))
    request_totals = cursor.fetchone() or {}

    payment_clauses, payment_params = _report_geo_scope_clauses(
        scope,
        "d.state_id",
        "d.id",
        "COALESCE(ut.id, uvt.id, wt.id)",
        "COALESCE(uv.id, wv.id)",
    )
    _apply_report_date_filter(payment_clauses, payment_params, "p.created_at", date_filters)
    cursor.execute(f"""
        SELECT
            COUNT(*) AS total,
            COALESCE(SUM(COALESCE(p.total, 0)), 0) AS total_amount,
            COALESCE(SUM(COALESCE(p.admin_share, 0)), 0) AS admin_share,
            COALESCE(SUM(COALESCE(p.worker_share, 0)), 0) AS worker_share
        FROM payments p
        LEFT JOIN requests r ON p.request_id = r.id
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas uvt ON uv.taluka_id = uvt.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        LEFT JOIN villages wv ON vw.village_id = wv.id
        LEFT JOIN talukas wt ON wv.taluka_id = wt.id
        INNER JOIN districts d ON COALESCE(u.district_id, ut.district_id, uvt.district_id, wt.district_id) = d.id
        WHERE {' AND '.join(payment_clauses)}
    """, tuple(payment_params))
    payment_totals = cursor.fetchone() or {}

    task_clauses, task_params = _report_geo_scope_clauses(scope, "d.state_id", "d.id", "t.id", "v.id")
    _apply_report_date_filter(task_clauses, task_params, "ts.assigned_at", date_filters)
    cursor.execute(f"""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN COALESCE(ts.status, 'pending') = 'completed' THEN 1 ELSE 0 END) AS completed
        FROM tasks ts
        INNER JOIN village_workers vw ON ts.worker_id = vw.id
        INNER JOIN villages v ON vw.village_id = v.id
        INNER JOIN talukas t ON v.taluka_id = t.id
        INNER JOIN districts d ON t.district_id = d.id
        WHERE {' AND '.join(task_clauses)}
    """, tuple(task_params))
    task_totals = cursor.fetchone() or {}

    rows = [
        {'metric': 'Scope', 'value': _report_scope_label(scope)},
        {'metric': 'Date Window', 'value': date_filters.get('label')},
        {'metric': 'Villages', 'value': villages_total or 0},
        {'metric': 'Citizens', 'value': citizens_total or 0},
        {'metric': 'Workers', 'value': workers_total or 0},
        {'metric': 'Complaints Total', 'value': complaint_totals.get('total') or 0},
        {'metric': 'Complaints Open', 'value': complaint_totals.get('open_items') or 0},
        {'metric': 'Door Pickup Requests', 'value': request_totals.get('total') or 0},
        {'metric': 'Completed Pickups', 'value': request_totals.get('completed') or 0},
        {'metric': 'Payments Total', 'value': payment_totals.get('total_amount') or 0},
        {'metric': 'Operational Tasks', 'value': task_totals.get('total') or 0},
    ]

    cards = [
        _report_card('Villages', villages_total or 0, 'navy', 'fa-map'),
        _report_card('Citizens', citizens_total or 0, 'sky', 'fa-users', date_filters.get('label')),
        _report_card('Workers', workers_total or 0, 'mint', 'fa-user-hard-hat'),
        _report_card('Open Complaints', complaint_totals.get('open_items') or 0, 'rose', 'fa-file-circle-exclamation'),
        _report_card('Door Pickups', request_totals.get('total') or 0, 'amber', 'fa-truck-ramp-box'),
        _report_card('Payments', f"Rs. {payment_totals.get('total_amount') or 0}", 'navy', 'fa-wallet'),
    ]

    return {
        'summary_cards': cards,
        'preview_sections': [{
            'title': 'Executive Summary',
            'description': 'Key metrics for the currently selected report scope and time window.',
            'columns': [
                {'key': 'metric', 'label': 'Metric'},
                {'key': 'value', 'label': 'Value'},
            ],
            'rows': rows,
            'total_rows': len(rows),
        }],
        'export_sections': [{
            'title': 'Executive Summary',
            'columns': [
                {'key': 'metric', 'label': 'Metric'},
                {'key': 'value', 'label': 'Value'},
            ],
            'rows': rows,
        }],
        'notes': [],
    }


def _build_village_report_dataset(scope, date_filters, cursor):
    del date_filters
    geo_clauses, geo_params = _report_geo_scope_clauses(scope, "d.state_id", "d.id", "t.id", "v.id")
    ensure_complaint_workflow_columns(cursor)
    cursor.execute(f"""
        SELECT
            v.id,
            v.name AS village_name,
            t.name AS taluka_name,
            d.name AS district_name,
            COUNT(DISTINCT u.id) AS citizens,
            COUNT(DISTINCT vw.id) AS workers,
            COUNT(DISTINCT CASE WHEN c.id IS NOT NULL THEN c.id END) AS complaints_total,
            COUNT(DISTINCT CASE WHEN LOWER(COALESCE(c.status, 'Pending')) = 'completed' THEN c.id END) AS complaints_completed,
            COUNT(DISTINCT CASE WHEN r.id IS NOT NULL THEN r.id END) AS pickup_requests,
            COUNT(DISTINCT CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'completed' THEN r.id END) AS pickup_completed
        FROM villages v
        INNER JOIN talukas t ON v.taluka_id = t.id
        INNER JOIN districts d ON t.district_id = d.id
        LEFT JOIN users u ON u.village_id = v.id
        LEFT JOIN village_workers vw ON vw.village_id = v.id
        LEFT JOIN requests r ON r.user_id = u.id
        LEFT JOIN complaints c ON c.village = v.name AND c.taluka = t.name AND c.district = d.name
        WHERE {' AND '.join(geo_clauses)}
        GROUP BY v.id, v.name, t.name, d.name
        ORDER BY d.name ASC, t.name ASC, v.name ASC
    """, tuple(geo_params))
    rows = cursor.fetchall() or []
    for row in rows:
        total_pickups = row.get('pickup_requests') or 0
        completed_pickups = row.get('pickup_completed') or 0
        row['pickup_progress'] = round((completed_pickups / total_pickups) * 100) if total_pickups else 0

    return {
        'summary_cards': [
            _report_card('Villages', len(rows), 'navy', 'fa-map-location-dot'),
            _report_card('Citizens Linked', sum((row.get('citizens') or 0) for row in rows), 'sky', 'fa-users'),
            _report_card('Workers Linked', sum((row.get('workers') or 0) for row in rows), 'mint', 'fa-user-hard-hat'),
            _report_card('Pickup Completion', f"{round((sum((row.get('pickup_completed') or 0) for row in rows) / max(sum((row.get('pickup_requests') or 0) for row in rows), 1)) * 100)}%", 'amber', 'fa-chart-column'),
        ],
        'preview_sections': [{
            'title': 'Village Performance',
            'description': 'Master village records with service coverage inside the selected scope.',
            'columns': [
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'district_name', 'label': 'District'},
                {'key': 'citizens', 'label': 'Citizens'},
                {'key': 'workers', 'label': 'Workers'},
                {'key': 'complaints_total', 'label': 'Complaints'},
                {'key': 'pickup_requests', 'label': 'Door Pickups'},
                {'key': 'pickup_progress', 'label': 'Pickup %'},
            ],
            'rows': rows[:REPORT_PREVIEW_LIMIT],
            'total_rows': len(rows),
        }],
        'export_sections': [{
            'title': 'Villages',
            'columns': [
                {'key': 'id', 'label': 'Village ID'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'district_name', 'label': 'District'},
                {'key': 'citizens', 'label': 'Citizens'},
                {'key': 'workers', 'label': 'Workers'},
                {'key': 'complaints_total', 'label': 'Complaints Total'},
                {'key': 'complaints_completed', 'label': 'Complaints Completed'},
                {'key': 'pickup_requests', 'label': 'Door Pickups'},
                {'key': 'pickup_completed', 'label': 'Door Pickups Completed'},
                {'key': 'pickup_progress', 'label': 'Pickup Completion %'},
            ],
            'rows': rows,
        }],
        'notes': ['Village reports show master coverage data. Date filters are not applied to static village records.'],
    }


def _build_citizen_report_dataset(scope, date_filters, cursor):
    geo_clauses, geo_params = _report_geo_scope_clauses(scope, "d.state_id", "d.id", "t.id", "v.id")
    _apply_report_date_filter(geo_clauses, geo_params, "u.created_at", date_filters)
    cursor.execute(f"""
        SELECT
            u.id,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(u.email, '') AS email,
            COALESCE(u.phone, '') AS phone,
            d.name AS district_name,
            t.name AS taluka_name,
            v.name AS village_name,
            u.created_at
        FROM users u
        LEFT JOIN villages v ON u.village_id = v.id
        LEFT JOIN talukas t ON COALESCE(u.taluka_id, v.taluka_id) = t.id
        LEFT JOIN districts d ON COALESCE(u.district_id, t.district_id) = d.id
        WHERE {' AND '.join(geo_clauses)}
        ORDER BY u.created_at DESC, u.id DESC
    """, tuple(geo_params))
    rows = cursor.fetchall() or []
    for row in rows:
        row['created_at'] = _format_timestamp(row.get('created_at'), include_time=True)

    return {
        'summary_cards': [
            _report_card('Citizen Records', len(rows), 'navy', 'fa-users'),
            _report_card('Districts Covered', len({row.get('district_name') for row in rows if row.get('district_name')}), 'sky', 'fa-building'),
            _report_card('Talukas Covered', len({row.get('taluka_name') for row in rows if row.get('taluka_name')}), 'mint', 'fa-map'),
            _report_card('Date Window', date_filters.get('label'), 'amber', 'fa-calendar-days'),
        ],
        'preview_sections': [{
            'title': 'Citizen Registrations',
            'description': 'Citizens registered in the selected location and time window.',
            'columns': [
                {'key': 'citizen_name', 'label': 'Citizen'},
                {'key': 'phone', 'label': 'Phone'},
                {'key': 'district_name', 'label': 'District'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'created_at', 'label': 'Registered At'},
            ],
            'rows': rows[:REPORT_PREVIEW_LIMIT],
            'total_rows': len(rows),
        }],
        'export_sections': [{
            'title': 'Citizens',
            'columns': [
                {'key': 'id', 'label': 'Citizen ID'},
                {'key': 'citizen_name', 'label': 'Citizen'},
                {'key': 'email', 'label': 'Email'},
                {'key': 'phone', 'label': 'Phone'},
                {'key': 'district_name', 'label': 'District'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'created_at', 'label': 'Registered At'},
            ],
            'rows': rows,
        }],
        'notes': [],
    }


def _build_worker_report_dataset(scope, date_filters, cursor):
    geo_clauses, geo_params = _report_geo_scope_clauses(scope, "d.state_id", "d.id", "t.id", "v.id")
    _apply_report_date_filter(geo_clauses, geo_params, "vw.created_at", date_filters)
    cursor.execute(f"""
        SELECT
            vw.id,
            vw.name AS worker_name,
            COALESCE(vw.email, '') AS email,
            COALESCE(vw.phone, '') AS phone,
            COALESCE(vw.vehicle_no, '') AS vehicle_no,
            COALESCE(vw.status, 'Active') AS worker_status,
            v.name AS village_name,
            t.name AS taluka_name,
            d.name AS district_name,
            vw.created_at,
            (
                COALESCE((SELECT COUNT(*) FROM requests r WHERE r.worker_id = vw.id), 0)
                + COALESCE((SELECT COUNT(*) FROM tasks ts WHERE ts.worker_id = vw.id), 0)
                + COALESCE((SELECT COUNT(*) FROM complaints c WHERE c.worker_id = vw.id), 0)
            ) AS total_assignments,
            (
                COALESCE((SELECT COUNT(*) FROM requests r WHERE r.worker_id = vw.id AND LOWER(COALESCE(r.status, 'pending')) = 'completed'), 0)
                + COALESCE((SELECT COUNT(*) FROM tasks ts WHERE ts.worker_id = vw.id AND COALESCE(ts.status, 'pending') = 'completed'), 0)
                + COALESCE((SELECT COUNT(*) FROM complaints c WHERE c.worker_id = vw.id AND LOWER(COALESCE(c.status, 'Pending')) = 'completed'), 0)
            ) AS completed_assignments
        FROM village_workers vw
        INNER JOIN villages v ON vw.village_id = v.id
        INNER JOIN talukas t ON v.taluka_id = t.id
        INNER JOIN districts d ON t.district_id = d.id
        WHERE {' AND '.join(geo_clauses)}
        ORDER BY vw.name ASC
    """, tuple(geo_params))
    rows = cursor.fetchall() or []
    for row in rows:
        row['created_at'] = _format_timestamp(row.get('created_at'), include_time=True)
        total_assignments = row.get('total_assignments') or 0
        completed_assignments = row.get('completed_assignments') or 0
        row['completion_rate'] = round((completed_assignments / total_assignments) * 100) if total_assignments else 0

    return {
        'summary_cards': [
            _report_card('Workers', len(rows), 'navy', 'fa-user-hard-hat'),
            _report_card('Active', sum(1 for row in rows if _slug_text(row.get('worker_status')) == 'active'), 'mint', 'fa-circle-check'),
            _report_card('Completed Assignments', sum((row.get('completed_assignments') or 0) for row in rows), 'sky', 'fa-list-check'),
            _report_card('Avg Completion', f"{round(sum((row.get('completion_rate') or 0) for row in rows) / len(rows)) if rows else 0}%", 'amber', 'fa-chart-simple'),
        ],
        'preview_sections': [{
            'title': 'Field Workers',
            'description': 'Registered field team members with assignment performance.',
            'columns': [
                {'key': 'worker_name', 'label': 'Worker'},
                {'key': 'worker_status', 'label': 'Status'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'total_assignments', 'label': 'Assignments'},
                {'key': 'completion_rate', 'label': 'Completion %'},
            ],
            'rows': rows[:REPORT_PREVIEW_LIMIT],
            'total_rows': len(rows),
        }],
        'export_sections': [{
            'title': 'Workers',
            'columns': [
                {'key': 'id', 'label': 'Worker ID'},
                {'key': 'worker_name', 'label': 'Worker'},
                {'key': 'email', 'label': 'Email'},
                {'key': 'phone', 'label': 'Phone'},
                {'key': 'vehicle_no', 'label': 'Vehicle No'},
                {'key': 'worker_status', 'label': 'Status'},
                {'key': 'district_name', 'label': 'District'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'total_assignments', 'label': 'Total Assignments'},
                {'key': 'completed_assignments', 'label': 'Completed Assignments'},
                {'key': 'completion_rate', 'label': 'Completion %'},
                {'key': 'created_at', 'label': 'Created At'},
            ],
            'rows': rows,
        }],
        'notes': [],
    }


def _build_complaint_report_dataset(scope, date_filters, cursor):
    complaint_clauses = ["d.state_id = %s"]
    complaint_params = [scope.get('state_id')]
    if scope.get('district_id'):
        complaint_clauses.append("d.id = %s")
        complaint_params.append(scope.get('district_id'))
    if scope.get('taluka_id'):
        complaint_clauses.append("t.id = %s")
        complaint_params.append(scope.get('taluka_id'))
    if scope.get('village_id'):
        complaint_clauses.append("v.id = %s")
        complaint_params.append(scope.get('village_id'))
    _apply_report_date_filter(complaint_clauses, complaint_params, "c.created_at", date_filters)

    complaint_columns = ensure_complaint_workflow_columns(cursor)
    original_photo_select = complaint_original_photo_select(complaint_columns)

    cursor.execute(f"""
        SELECT
            c.id,
            c.title,
            c.description,
            COALESCE(c.priority, 'Normal') AS priority,
            COALESCE(c.status, 'Pending') AS complaint_status,
            c.district,
            c.taluka,
            c.village,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(vw.name, 'Unassigned') AS worker_name,
            c.created_at,
            COALESCE(c.updated_at, c.resolved_at, c.assigned_at, c.created_at) AS last_action_at,
            c.resolved_at,
            {original_photo_select}
        FROM complaints c
        INNER JOIN districts d ON d.name = c.district
        LEFT JOIN talukas t ON t.name = c.taluka AND t.district_id = d.id
        LEFT JOIN villages v ON v.name = c.village AND v.taluka_id = t.id
        LEFT JOIN users u ON c.user_id = u.id
        LEFT JOIN village_workers vw ON c.worker_id = vw.id
        WHERE {' AND '.join(complaint_clauses)}
        ORDER BY COALESCE(c.updated_at, c.resolved_at, c.assigned_at, c.created_at) DESC
    """, tuple(complaint_params))
    rows = cursor.fetchall() or []
    for row in rows:
        row['complaint_status'] = normalize_complaint_status(row.get('complaint_status'))
        row['created_at'] = _format_timestamp(row.get('created_at'), include_time=True)
        row['last_action_at'] = _format_timestamp(row.get('last_action_at'), include_time=True)
        row['resolved_at'] = _format_timestamp(row.get('resolved_at'), include_time=True)

    return {
        'summary_cards': [
            _report_card('Complaints', len(rows), 'navy', 'fa-file-circle-exclamation'),
            _report_card('Pending', sum(1 for row in rows if complaint_status_key(row.get('complaint_status')) == 'pending'), 'rose', 'fa-hourglass-half'),
            _report_card('In Progress', sum(1 for row in rows if complaint_status_key(row.get('complaint_status')) in ('assigned', 'in_progress')), 'amber', 'fa-spinner'),
            _report_card('Completed', sum(1 for row in rows if complaint_status_key(row.get('complaint_status')) == 'completed'), 'mint', 'fa-circle-check'),
        ],
        'preview_sections': [{
            'title': 'Complaint Register',
            'description': 'Reported service issues and resolution progress.',
            'columns': [
                {'key': 'title', 'label': 'Title'},
                {'key': 'complaint_status', 'label': 'Status'},
                {'key': 'priority', 'label': 'Priority'},
                {'key': 'village', 'label': 'Village'},
                {'key': 'taluka', 'label': 'Taluka'},
                {'key': 'created_at', 'label': 'Created'},
            ],
            'rows': rows[:REPORT_PREVIEW_LIMIT],
            'total_rows': len(rows),
        }],
        'export_sections': [{
            'title': 'Complaints',
            'columns': [
                {'key': 'id', 'label': 'Complaint ID'},
                {'key': 'title', 'label': 'Title'},
                {'key': 'description', 'label': 'Description'},
                {'key': 'priority', 'label': 'Priority'},
                {'key': 'complaint_status', 'label': 'Status'},
                {'key': 'citizen_name', 'label': 'Citizen'},
                {'key': 'worker_name', 'label': 'Assigned Worker'},
                {'key': 'district', 'label': 'District'},
                {'key': 'taluka', 'label': 'Taluka'},
                {'key': 'village', 'label': 'Village'},
                {'key': 'created_at', 'label': 'Created At'},
                {'key': 'last_action_at', 'label': 'Last Action'},
                {'key': 'resolved_at', 'label': 'Resolved At'},
            ],
            'rows': rows,
        }],
        'notes': [],
    }


def _build_request_report_dataset(scope, date_filters, cursor):
    request_clauses, request_params = _report_geo_scope_clauses(
        scope,
        "d.state_id",
        "d.id",
        "COALESCE(ut.id, uvt.id, wt.id)",
        "COALESCE(uv.id, wv.id)",
    )
    _apply_report_date_filter(request_clauses, request_params, "r.created_at", date_filters)
    cursor.execute(f"""
        SELECT
            r.id,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(r.garbage_type, 'Garbage Collection') AS service_type,
            COALESCE(r.status, 'pending') AS service_status,
            COALESCE(r.amount, 0) AS amount,
            COALESCE(r.weight, 0) AS weight,
            COALESCE(vw.name, 'Unassigned') AS worker_name,
            COALESCE(d.name, 'Unknown District') AS district_name,
            COALESCE(ut.name, uvt.name, wt.name, 'Unknown Taluka') AS taluka_name,
            COALESCE(uv.name, wv.name, 'Unknown Village') AS village_name,
            r.created_at
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas uvt ON uv.taluka_id = uvt.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        LEFT JOIN villages wv ON vw.village_id = wv.id
        LEFT JOIN talukas wt ON wv.taluka_id = wt.id
        INNER JOIN districts d ON COALESCE(u.district_id, ut.district_id, uvt.district_id, wt.district_id) = d.id
        WHERE {' AND '.join(request_clauses)}
        ORDER BY r.created_at DESC, r.id DESC
    """, tuple(request_params))
    rows = cursor.fetchall() or []
    for row in rows:
        row['service_status'] = (row.get('service_status') or 'pending').replace('_', ' ').title()
        row['created_at'] = _format_timestamp(row.get('created_at'), include_time=True)

    return {
        'summary_cards': [
            _report_card('Pickup Requests', len(rows), 'navy', 'fa-truck-ramp-box'),
            _report_card('Pending', sum(1 for row in rows if _slug_text(row.get('service_status')) == 'pending'), 'rose', 'fa-hourglass-half'),
            _report_card('Completed', sum(1 for row in rows if _slug_text(row.get('service_status')) == 'completed'), 'mint', 'fa-circle-check'),
            _report_card('Total Value', f"Rs. {sum((row.get('amount') or 0) for row in rows)}", 'amber', 'fa-wallet'),
        ],
        'preview_sections': [{
            'title': 'Door Pickup Services',
            'description': 'Door-to-door and bulk collection requests created by citizens.',
            'columns': [
                {'key': 'citizen_name', 'label': 'Citizen'},
                {'key': 'service_type', 'label': 'Service'},
                {'key': 'service_status', 'label': 'Status'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'worker_name', 'label': 'Worker'},
                {'key': 'created_at', 'label': 'Created'},
            ],
            'rows': rows[:REPORT_PREVIEW_LIMIT],
            'total_rows': len(rows),
        }],
        'export_sections': [{
            'title': 'Door Pickup Services',
            'columns': [
                {'key': 'id', 'label': 'Request ID'},
                {'key': 'citizen_name', 'label': 'Citizen'},
                {'key': 'service_type', 'label': 'Service Type'},
                {'key': 'service_status', 'label': 'Status'},
                {'key': 'amount', 'label': 'Amount'},
                {'key': 'weight', 'label': 'Weight'},
                {'key': 'worker_name', 'label': 'Worker'},
                {'key': 'district_name', 'label': 'District'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'created_at', 'label': 'Created At'},
            ],
            'rows': rows,
        }],
        'notes': [],
    }


def _build_payment_report_dataset(scope, date_filters, cursor):
    payment_clauses, payment_params = _report_geo_scope_clauses(
        scope,
        "d.state_id",
        "d.id",
        "COALESCE(ut.id, uvt.id, wt.id)",
        "COALESCE(uv.id, wv.id)",
    )
    _apply_report_date_filter(payment_clauses, payment_params, "p.created_at", date_filters)
    cursor.execute(f"""
        SELECT
            p.id,
            p.request_id,
            COALESCE(u.name, 'Citizen') AS citizen_name,
            COALESCE(r.garbage_type, 'Garbage Collection') AS service_type,
            COALESCE(p.total, 0) AS total,
            COALESCE(p.owner_share, 0) AS owner_share,
            COALESCE(p.admin_share, 0) AS admin_share,
            COALESCE(p.worker_share, 0) AS worker_share,
            COALESCE(d.name, 'Unknown District') AS district_name,
            COALESCE(ut.name, uvt.name, wt.name, 'Unknown Taluka') AS taluka_name,
            COALESCE(uv.name, wv.name, 'Unknown Village') AS village_name,
            p.created_at
        FROM payments p
        LEFT JOIN requests r ON p.request_id = r.id
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas uvt ON uv.taluka_id = uvt.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        LEFT JOIN villages wv ON vw.village_id = wv.id
        LEFT JOIN talukas wt ON wv.taluka_id = wt.id
        INNER JOIN districts d ON COALESCE(u.district_id, ut.district_id, uvt.district_id, wt.district_id) = d.id
        WHERE {' AND '.join(payment_clauses)}
        ORDER BY p.created_at DESC, p.id DESC
    """, tuple(payment_params))
    rows = cursor.fetchall() or []
    for row in rows:
        row['created_at'] = _format_timestamp(row.get('created_at'), include_time=True)

    return {
        'summary_cards': [
            _report_card('Payments', len(rows), 'navy', 'fa-wallet'),
            _report_card('Collection Total', f"Rs. {sum((row.get('total') or 0) for row in rows)}", 'mint', 'fa-indian-rupee-sign'),
            _report_card('Admin Share', f"Rs. {sum((row.get('admin_share') or 0) for row in rows)}", 'amber', 'fa-scale-balanced'),
            _report_card('Worker Share', f"Rs. {sum((row.get('worker_share') or 0) for row in rows)}", 'sky', 'fa-hand-holding-dollar'),
        ],
        'preview_sections': [{
            'title': 'Payment Records',
            'description': 'Collection payments and revenue sharing across the selected scope.',
            'columns': [
                {'key': 'request_id', 'label': 'Request'},
                {'key': 'citizen_name', 'label': 'Citizen'},
                {'key': 'service_type', 'label': 'Service'},
                {'key': 'total', 'label': 'Total'},
                {'key': 'admin_share', 'label': 'Admin Share'},
                {'key': 'created_at', 'label': 'Paid At'},
            ],
            'rows': rows[:REPORT_PREVIEW_LIMIT],
            'total_rows': len(rows),
        }],
        'export_sections': [{
            'title': 'Payments',
            'columns': [
                {'key': 'id', 'label': 'Payment ID'},
                {'key': 'request_id', 'label': 'Request ID'},
                {'key': 'citizen_name', 'label': 'Citizen'},
                {'key': 'service_type', 'label': 'Service Type'},
                {'key': 'total', 'label': 'Total'},
                {'key': 'owner_share', 'label': 'Owner Share'},
                {'key': 'admin_share', 'label': 'Admin Share'},
                {'key': 'worker_share', 'label': 'Worker Share'},
                {'key': 'district_name', 'label': 'District'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'created_at', 'label': 'Paid At'},
            ],
            'rows': rows,
        }],
        'notes': [],
    }


def _build_task_report_dataset(scope, date_filters, cursor):
    task_clauses, task_params = _report_geo_scope_clauses(scope, "d.state_id", "d.id", "t.id", "v.id")
    _apply_report_date_filter(task_clauses, task_params, "ts.assigned_at", date_filters)
    cursor.execute(f"""
        SELECT
            ts.id,
            COALESCE(vw.name, 'Unknown Worker') AS worker_name,
            COALESCE(ts.location_name, v.name, 'Assigned Area') AS location_name,
            COALESCE(ts.description, '') AS description,
            COALESCE(ts.status, 'pending') AS task_status,
            COALESCE(ts.priority, 'medium') AS priority,
            d.name AS district_name,
            t.name AS taluka_name,
            v.name AS village_name,
            ts.assigned_at
        FROM tasks ts
        INNER JOIN village_workers vw ON ts.worker_id = vw.id
        INNER JOIN villages v ON vw.village_id = v.id
        INNER JOIN talukas t ON v.taluka_id = t.id
        INNER JOIN districts d ON t.district_id = d.id
        WHERE {' AND '.join(task_clauses)}
        ORDER BY ts.assigned_at DESC, ts.id DESC
    """, tuple(task_params))
    rows = cursor.fetchall() or []
    for row in rows:
        row['task_status'] = (row.get('task_status') or 'pending').replace('_', ' ').title()
        row['priority'] = (row.get('priority') or 'medium').title()
        row['assigned_at'] = _format_timestamp(row.get('assigned_at'), include_time=True)

    return {
        'summary_cards': [
            _report_card('Tasks', len(rows), 'navy', 'fa-list-check'),
            _report_card('Pending', sum(1 for row in rows if _slug_text(row.get('task_status')) == 'pending'), 'rose', 'fa-hourglass-half'),
            _report_card('Completed', sum(1 for row in rows if _slug_text(row.get('task_status')) == 'completed'), 'mint', 'fa-circle-check'),
            _report_card('High Priority', sum(1 for row in rows if _slug_text(row.get('priority')) == 'high'), 'amber', 'fa-bolt'),
        ],
        'preview_sections': [{
            'title': 'Operations Tasks',
            'description': 'Manual operational tasks assigned to field workers.',
            'columns': [
                {'key': 'worker_name', 'label': 'Worker'},
                {'key': 'location_name', 'label': 'Location'},
                {'key': 'task_status', 'label': 'Status'},
                {'key': 'priority', 'label': 'Priority'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'assigned_at', 'label': 'Assigned'},
            ],
            'rows': rows[:REPORT_PREVIEW_LIMIT],
            'total_rows': len(rows),
        }],
        'export_sections': [{
            'title': 'Operations Tasks',
            'columns': [
                {'key': 'id', 'label': 'Task ID'},
                {'key': 'worker_name', 'label': 'Worker'},
                {'key': 'location_name', 'label': 'Location'},
                {'key': 'description', 'label': 'Description'},
                {'key': 'task_status', 'label': 'Status'},
                {'key': 'priority', 'label': 'Priority'},
                {'key': 'district_name', 'label': 'District'},
                {'key': 'taluka_name', 'label': 'Taluka'},
                {'key': 'village_name', 'label': 'Village'},
                {'key': 'assigned_at', 'label': 'Assigned At'},
            ],
            'rows': rows,
        }],
        'notes': [],
    }


def _build_report_dataset(scope, date_filters, dataset_key, cursor):
    if dataset_key == 'summary':
        return _build_summary_dataset(scope, date_filters, cursor)
    if dataset_key == 'villages':
        return _build_village_report_dataset(scope, date_filters, cursor)
    if dataset_key == 'citizens':
        return _build_citizen_report_dataset(scope, date_filters, cursor)
    if dataset_key == 'workers':
        return _build_worker_report_dataset(scope, date_filters, cursor)
    if dataset_key == 'complaints':
        return _build_complaint_report_dataset(scope, date_filters, cursor)
    if dataset_key == 'door_pickups':
        return _build_request_report_dataset(scope, date_filters, cursor)
    if dataset_key == 'payments':
        return _build_payment_report_dataset(scope, date_filters, cursor)
    if dataset_key == 'tasks':
        return _build_task_report_dataset(scope, date_filters, cursor)
    if dataset_key == 'all_operations':
        summary_data = _build_summary_dataset(scope, date_filters, cursor)
        citizen_data = _build_citizen_report_dataset(scope, date_filters, cursor)
        worker_data = _build_worker_report_dataset(scope, date_filters, cursor)
        complaint_data = _build_complaint_report_dataset(scope, date_filters, cursor)
        request_data = _build_request_report_dataset(scope, date_filters, cursor)
        payment_data = _build_payment_report_dataset(scope, date_filters, cursor)
        task_data = _build_task_report_dataset(scope, date_filters, cursor)

        preview_sections = []
        export_sections = []
        for block in (summary_data, citizen_data, worker_data, complaint_data, request_data, payment_data, task_data):
            export_sections.extend(block.get('export_sections', []))
            for section in block.get('export_sections', []):
                preview_sections.append({
                    'title': section.get('title'),
                    'description': REPORT_DATASETS.get('all_operations', {}).get('description', ''),
                    'columns': section.get('columns', []),
                    'rows': (section.get('rows') or [])[:REPORT_SECTION_PREVIEW_LIMIT],
                    'total_rows': len(section.get('rows') or []),
                })

        return {
            'summary_cards': summary_data.get('summary_cards', []) + [
                _report_card('Preview Sections', len(preview_sections), 'sky', 'fa-layer-group'),
            ],
            'preview_sections': preview_sections,
            'export_sections': export_sections,
            'notes': ['The all-operations bundle exports multiple sheets in Excel and multiple sections in PDF.'],
        }

    raise ValueError("Unknown report dataset selected.")


def _report_security_notes(role, scope):
    role_label = {
        'admin': 'Taluka admins',
        'district_admin': 'District admins',
        'state_admin': 'State admins',
    }.get(role, 'Admins')

    notes = [
        f"{role_label} can export data only inside {_report_scope_label(scope)}.",
        "Every report request is re-validated on the server before preview or download.",
        "Tampered district, taluka, or village filters are blocked before any data is returned.",
    ]

    if role == 'admin':
        notes.append("Taluka admins cannot open district-wide or state-wide records.")
    elif role == 'district_admin':
        notes.append("District admins can drill into taluka and village data, but not outside their district.")
    else:
        notes.append("State admins can narrow down by district, taluka, or village across their own state.")

    return notes


def _report_filename(role, scope, dataset_key, extension):
    role_prefix = {
        'admin': 'taluka-report',
        'district_admin': 'district-report',
        'state_admin': 'state-report',
    }.get(role, 'report')

    scope_bits = [
        scope.get('district_name'),
        scope.get('taluka_name'),
        scope.get('village_name'),
    ]
    raw_scope = '-'.join([bit for bit in scope_bits if bit]) or (scope.get('state_name') or 'scope')
    clean_scope = re.sub(r'[^a-z0-9]+', '-', raw_scope.strip().lower()).strip('-') or 'scope'
    dataset_slug = re.sub(r'[^a-z0-9]+', '-', dataset_key.strip().lower()).strip('-')
    timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
    return f"{role_prefix}-{dataset_slug}-{clean_scope}-{timestamp}.{extension}"


def _autosize_report_sheet(sheet):
    for column_cells in sheet.columns:
        cell_lengths = [len(str(cell.value)) for cell in column_cells if cell.value not in (None, '')]
        if not cell_lengths:
            continue
        column_letter = column_cells[0].column_letter
        sheet.column_dimensions[column_letter].width = min(max(cell_lengths) + 2, 36)


def _write_report_sheet(workbook, section):
    sheet_name = (section.get('title') or 'Report')[:31]
    worksheet = workbook.create_sheet(title=sheet_name)
    title_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="14213D")

    columns = section.get('columns', [])
    rows = section.get('rows', [])

    for column_index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column.get('label'))
        cell.font = title_font
        cell.fill = header_fill

    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            worksheet.cell(row=row_index, column=column_index, value=row.get(column.get('key')))

    worksheet.freeze_panes = "A2"
    _autosize_report_sheet(worksheet)


def _build_excel_report_bytes(report_payload):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for section in report_payload.get('export_sections', []):
        _write_report_sheet(workbook, section)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _escape_pdf_text(value):
    return str(value).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _section_to_pdf_lines(section):
    lines = [section.get('title', 'Report Section')]
    for row in section.get('rows', []):
        row_text = " | ".join(
            f"{column.get('label')}: {row.get(column.get('key'), '')}"
            for column in section.get('columns', [])
        )
        for wrapped_line in textwrap.wrap(row_text, width=108) or ['']:
            lines.append(wrapped_line)
        lines.append("")
    if len(lines) == 1:
        lines.append("No records found.")
    lines.append("")
    return lines


def _build_pdf_report_bytes(title, subtitle, report_payload):
    page_width = 842
    page_height = 595
    line_height = 13
    start_x = 36
    start_y = 560
    max_lines_per_page = 38

    lines = [title, subtitle, ""]
    for section in report_payload.get('export_sections', []):
        lines.extend(_section_to_pdf_lines(section))

    pages = []
    for index in range(0, len(lines), max_lines_per_page):
        pages.append(lines[index:index + max_lines_per_page])

    objects = []

    def add_object(payload):
        objects.append(payload)
        return len(objects)

    font_object_id = add_object("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids = []
    content_ids = []

    for page_lines in pages or [["No report data available."]]:
        content_stream = [
            "BT",
            f"/F1 10 Tf",
            f"{line_height} TL",
            f"{start_x} {start_y} Td",
        ]
        first_line = True
        for line in page_lines:
            escaped = _escape_pdf_text(line)
            if first_line:
                content_stream.append(f"({escaped}) Tj")
                first_line = False
            else:
                content_stream.append("T*")
                content_stream.append(f"({escaped}) Tj")
        content_stream.append("ET")
        stream_text = "\n".join(content_stream)
        content_object_id = add_object(f"<< /Length {len(stream_text.encode('latin-1', errors='replace'))} >>\nstream\n{stream_text}\nendstream")
        content_ids.append(content_object_id)

        page_object_id = add_object(
            f"<< /Type /Page /Parent 0 0 R /MediaBox [0 0 {page_width} {page_height}] "
            f"/Resources << /Font << /F1 {font_object_id} 0 R >> >> /Contents {content_object_id} 0 R >>"
        )
        page_ids.append(page_object_id)

    pages_object_id = add_object(
        f"<< /Type /Pages /Kids [{' '.join(f'{page_id} 0 R' for page_id in page_ids)}] /Count {len(page_ids)} >>"
    )

    for page_id in page_ids:
        objects[page_id - 1] = objects[page_id - 1].replace("/Parent 0 0 R", f"/Parent {pages_object_id} 0 R")

    catalog_object_id = add_object(f"<< /Type /Catalog /Pages {pages_object_id} 0 R >>")

    pdf_chunks = [b"%PDF-1.4\n"]
    offsets = [0]
    running_length = len(pdf_chunks[0])

    for object_index, object_payload in enumerate(objects, start=1):
        offsets.append(running_length)
        object_bytes = f"{object_index} 0 obj\n{object_payload}\nendobj\n".encode('latin-1', errors='replace')
        pdf_chunks.append(object_bytes)
        running_length += len(object_bytes)

    xref_start = running_length
    xref_lines = [f"xref\n0 {len(objects) + 1}\n", "0000000000 65535 f \n"]
    for offset in offsets[1:]:
        xref_lines.append(f"{offset:010d} 00000 n \n")
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_object_id} 0 R >>\nstartxref\n{xref_start}\n%%EOF"
    pdf_chunks.append("".join(xref_lines).encode('latin-1'))
    pdf_chunks.append(trailer.encode('latin-1'))

    buffer = BytesIO(b"".join(pdf_chunks))
    buffer.seek(0)
    return buffer


def _render_report_center(role):
    dataset_key = (request.args.get('dataset') or 'summary').strip().lower()
    if dataset_key not in REPORT_DATASETS:
        dataset_key = 'summary'

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        if role == 'state_admin':
            profile = _get_state_admin_profile(session.get('user_id'), cursor)
            template_name = 'dashboard/state_reports.html'
            page_context = {'state_page': 'reports'}
            export_endpoint = 'auth_bp.state_report_export'
        elif role == 'district_admin':
            profile = _get_district_admin_profile(session.get('user_id'), cursor)
            template_name = 'dashboard/district_reports.html'
            page_context = {'district_page': 'reports'}
            export_endpoint = 'auth_bp.district_report_export'
        else:
            profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
            template_name = 'dashboard/taluka_reports.html'
            page_context = {'worker_page': 'reports'}
            export_endpoint = 'auth_bp.taluka_report_export'

        if not profile:
            flash("Unable to load your reporting profile right now.", "warning")
            return redirect(url_for('auth_bp.login'))

        date_filters = _resolve_report_date_filters()
        scope = _resolve_report_scope(role, profile, cursor)
        filter_options = _build_report_filter_options(role, profile, scope, cursor)
        report_payload = _build_report_dataset(scope, date_filters, dataset_key, cursor)
        export_query = _report_export_query(scope, date_filters, dataset_key)

        dataset_options = []
        for key, item in REPORT_DATASETS.items():
            dataset_options.append({
                'key': key,
                'label': item.get('label'),
                'description': item.get('description'),
                'icon': item.get('icon'),
                'selected': key == dataset_key,
            })

        context = {
            'report_dataset_key': dataset_key,
            'report_dataset_label': REPORT_DATASETS[dataset_key]['label'],
            'report_dataset_description': REPORT_DATASETS[dataset_key]['description'],
            'report_scope_label': _report_scope_label(scope),
            'report_scope': scope,
            'report_date_filters': date_filters,
            'report_filter_options': filter_options,
            'report_dataset_options': dataset_options,
            'report_payload': report_payload,
            'report_security_notes': _report_security_notes(role, scope),
            'report_export_excel_url': url_for(export_endpoint, format_name='xlsx', **export_query),
            'report_export_pdf_url': url_for(export_endpoint, format_name='pdf', **export_query),
        }
        context.update(page_context)
        if role == 'state_admin':
            context['state_profile'] = profile
        elif role == 'district_admin':
            context['district_profile'] = profile
        else:
            context['admin_profile'] = profile

        return render_template(template_name, **context)
    except PermissionError as exc:
        flash(str(exc), "warning")
        if role == 'state_admin':
            return redirect(url_for('auth_bp.state_dashboard'))
        if role == 'district_admin':
            return redirect(url_for('auth_bp.district_dashboard'))
        return redirect(url_for('auth_bp.taluka_dashboard'))
    except ValueError as exc:
        flash(str(exc), "warning")
        if role == 'state_admin':
            return redirect(url_for('auth_bp.state_reports'))
        if role == 'district_admin':
            return redirect(url_for('auth_bp.district_reports'))
        return redirect(url_for('auth_bp.taluka_reports'))
    finally:
        conn.close()


def _export_report_center(role, format_name):
    format_name = (format_name or '').strip().lower()
    if format_name not in {'xlsx', 'pdf'}:
        flash("Unsupported report format requested.", "warning")
        if role == 'state_admin':
            return redirect(url_for('auth_bp.state_reports'))
        if role == 'district_admin':
            return redirect(url_for('auth_bp.district_reports'))
        return redirect(url_for('auth_bp.taluka_reports'))

    dataset_key = (request.args.get('dataset') or 'summary').strip().lower()
    if dataset_key not in REPORT_DATASETS:
        dataset_key = 'summary'

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        if role == 'state_admin':
            profile = _get_state_admin_profile(session.get('user_id'), cursor)
            fallback_endpoint = 'auth_bp.state_reports'
        elif role == 'district_admin':
            profile = _get_district_admin_profile(session.get('user_id'), cursor)
            fallback_endpoint = 'auth_bp.district_reports'
        else:
            profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
            fallback_endpoint = 'auth_bp.taluka_reports'

        if not profile:
            flash("Unable to load your reporting profile right now.", "warning")
            return redirect(url_for('auth_bp.login'))

        date_filters = _resolve_report_date_filters()
        scope = _resolve_report_scope(role, profile, cursor)
        report_payload = _build_report_dataset(scope, date_filters, dataset_key, cursor)
        filename = _report_filename(role, scope, dataset_key, format_name)

        if format_name == 'xlsx':
            return send_file(
                _build_excel_report_bytes(report_payload),
                as_attachment=True,
                download_name=filename,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )

        pdf_title = f"{REPORT_DATASETS[dataset_key]['label']} Report"
        pdf_subtitle = f"{_report_scope_label(scope)} | {date_filters.get('label')}"
        return send_file(
            _build_pdf_report_bytes(pdf_title, pdf_subtitle, report_payload),
            as_attachment=True,
            download_name=filename,
            mimetype='application/pdf',
        )
    except PermissionError as exc:
        flash(str(exc), "warning")
        return redirect(url_for(fallback_endpoint))
    except ValueError as exc:
        flash(str(exc), "warning")
        return redirect(url_for(fallback_endpoint))
    finally:
        conn.close()


def _get_district_admin_task_items(district_admin_id, cursor, limit=None):
    ensure_district_admin_task_table(cursor)
    limit_clause = ""
    if limit is not None:
        limit_clause = f"\n        LIMIT {max(int(limit), 0)}"

    cursor.execute(f"""
        SELECT
            dat.id,
            dat.district_admin_id,
            dat.taluka_admin_id,
            COALESCE(dat.location_name, '') AS location_name,
            dat.description,
            COALESCE(dat.priority, 'medium') AS priority,
            COALESCE(dat.status, 'pending') AS status,
            dat.assigned_at,
            dat.completed_at,
            COALESCE(ta.name, 'Taluka Admin') AS taluka_admin_name,
            COALESCE(t.name, 'Taluka') AS taluka_name
        FROM district_admin_tasks dat
        LEFT JOIN taluka_admins ta ON dat.taluka_admin_id = ta.id
        LEFT JOIN talukas t ON ta.taluka_id = t.id
        WHERE dat.district_admin_id = %s
        ORDER BY dat.assigned_at DESC
        {limit_clause}
    """, (district_admin_id,))
    tasks = cursor.fetchall() or []

    for task in tasks:
        task['status'] = (task.get('status') or 'pending').lower()
        task['status_label'] = task['status'].replace('_', ' ').title()
        task['priority_key'] = _slug_text(task.get('priority') or 'medium')
        task['assigned_at_text'] = _format_timestamp(task.get('assigned_at'), include_time=True)
        task['completed_at_text'] = _format_timestamp(task.get('completed_at'), include_time=True)

    return tasks


def _get_taluka_admin_district_tasks(taluka_admin_id, cursor, limit=None):
    ensure_district_admin_task_table(cursor)
    limit_clause = ""
    if limit is not None:
        limit_clause = f"\n        LIMIT {max(int(limit), 0)}"

    cursor.execute(f"""
        SELECT
            dat.id,
            dat.taluka_admin_id,
            COALESCE(dat.location_name, '') AS location_name,
            dat.description,
            COALESCE(dat.priority, 'medium') AS priority,
            COALESCE(dat.status, 'pending') AS status,
            dat.assigned_at,
            dat.completed_at,
            COALESCE(da.name, 'District Office') AS district_admin_name
        FROM district_admin_tasks dat
        LEFT JOIN district_admins da ON dat.district_admin_id = da.id
        WHERE dat.taluka_admin_id = %s
        ORDER BY dat.assigned_at DESC
        {limit_clause}
    """, (taluka_admin_id,))
    tasks = cursor.fetchall() or []

    for task in tasks:
        task['status'] = (task.get('status') or 'pending').lower()
        task['status_label'] = task['status'].replace('_', ' ').title()
        task['priority_key'] = _slug_text(task.get('priority') or 'medium')
        task['assigned_at_text'] = _format_timestamp(task.get('assigned_at'), include_time=True)
        task['completed_at_text'] = _format_timestamp(task.get('completed_at'), include_time=True)

    return tasks


def _build_district_map_payload(admin_profile, complaints, requests, assigned_tasks):
    district_name = admin_profile.get('district_name')
    markers = []

    for complaint in complaints or []:
        status_key = complaint.get('status_key') or complaint_status_key(complaint.get('status'))
        created_at = complaint.get('created_at')
        updated_at = complaint.get('updated_at') or complaint.get('resolved_at') or complaint.get('assigned_at') or created_at
        map_query = _build_map_query(
            complaint.get('village'),
            complaint.get('taluka'),
            district_name,
            'Maharashtra',
            'India',
        )
        location_payload = build_location_payload(
            latitude=complaint.get('latitude'),
            longitude=complaint.get('longitude'),
            accuracy=complaint.get('location_accuracy_meters'),
            query=map_query,
            directions=True,
        )
        markers.append({
            'id': f"CMP-{complaint.get('id')}",
            'category': 'complaint',
            'category_label': 'Complaint',
            'title': complaint.get('title') or 'Complaint',
            'description': complaint.get('description') or 'Citizen complaint',
            'status': status_key,
            'status_label': complaint.get('status') or 'Pending',
            'location_label': f"{complaint.get('village') or 'Unknown Village'} • {complaint.get('taluka') or district_name or 'District'}",
            'map_query': map_query,
            'secondary_label': complaint.get('citizen_name') or 'Citizen',
            'assigned_to': complaint.get('worker_name') or 'Unassigned',
            'time_label': _format_timestamp(updated_at, include_time=True),
            'sort_at': updated_at.isoformat() if updated_at else '',
            'is_today': _is_today(created_at),
            'is_open': status_key != 'completed',
            **location_payload,
        })

    for pickup_request in requests or []:
        status_key = (pickup_request.get('status') or 'pending').lower()
        created_at = pickup_request.get('created_at')
        map_query = _build_map_query(
            pickup_request.get('village_name'),
            pickup_request.get('taluka_name'),
            district_name,
            'Maharashtra',
            'India',
        )
        location_payload = build_location_payload(
            latitude=pickup_request.get('latitude'),
            longitude=pickup_request.get('longitude'),
            accuracy=pickup_request.get('location_accuracy_meters'),
            query=map_query,
            directions=True,
        )
        markers.append({
            'id': f"REQ-{pickup_request.get('id')}",
            'category': 'request',
            'category_label': 'Pickup Request',
            'title': pickup_request.get('garbage_type') or 'Garbage Collection',
            'description': f"Pickup request from {pickup_request.get('citizen_name') or 'Citizen'}",
            'status': status_key,
            'status_label': pickup_request.get('status_label') or status_key.replace('_', ' ').title(),
            'location_label': f"{pickup_request.get('village_name') or 'Unknown Village'} • {pickup_request.get('taluka_name') or district_name or 'District'}",
            'map_query': map_query,
            'secondary_label': pickup_request.get('citizen_name') or 'Citizen',
            'assigned_to': pickup_request.get('worker_name') or 'Unassigned',
            'time_label': _format_timestamp(created_at, include_time=True),
            'sort_at': created_at.isoformat() if created_at else '',
            'is_today': _is_today(created_at),
            'is_open': status_key != 'completed',
            **location_payload,
        })

    for assigned_task in assigned_tasks or []:
        status_key = (assigned_task.get('status') or 'pending').lower()
        assigned_at = assigned_task.get('assigned_at')
        markers.append({
            'id': f"TASK-{assigned_task.get('id')}",
            'category': 'assigned_task',
            'category_label': 'Assigned Task',
            'title': assigned_task.get('description') or assigned_task.get('location_name') or 'Assigned task',
            'description': assigned_task.get('description') or 'Manual task assigned by district office',
            'status': status_key,
            'status_label': assigned_task.get('status_label') or status_key.replace('_', ' ').title(),
            'location_label': f"{assigned_task.get('location_name') or assigned_task.get('village_name') or 'Assigned Area'} • {assigned_task.get('taluka_name') or district_name or 'District'}",
            'map_query': _build_map_query(
                assigned_task.get('location_name') or assigned_task.get('village_name'),
                assigned_task.get('taluka_name'),
                district_name,
                'Maharashtra',
                'India',
            ),
            'secondary_label': assigned_task.get('worker_name') or 'Worker',
            'assigned_to': assigned_task.get('worker_name') or 'Worker',
            'time_label': _format_timestamp(assigned_at, include_time=True),
            'sort_at': assigned_at.isoformat() if assigned_at else '',
            'is_today': _is_today(assigned_at),
            'is_open': status_key != 'completed',
        })

    markers.sort(key=lambda item: item.get('sort_at') or '', reverse=True)

    today_counts = {
        'complaints': sum(1 for item in markers if item.get('category') == 'complaint' and item.get('is_today')),
        'requests': sum(1 for item in markers if item.get('category') == 'request' and item.get('is_today')),
        'assigned_tasks': sum(1 for item in markers if item.get('category') == 'assigned_task' and item.get('is_today')),
    }
    today_counts['total'] = sum(today_counts.values())

    return {
        'scope_label': f"{district_name or 'District'} Scope",
        'center': None,
        'center_query': _build_map_query(district_name, 'Maharashtra', 'India'),
        'markers': markers,
        'today_counts': today_counts,
        'open_total': sum(1 for item in markers if item.get('is_open')),
    }


def _empty_district_overview_data(admin_profile):
    return {
        'dashboard_stats': {
            'talukas': 0,
            'villages': 0,
            'taluka_admins': 0,
            'citizens': 0,
            'workers': 0,
            'active_workers': 0,
            'complaints_total': 0,
            'complaints_pending': 0,
            'complaints_in_progress': 0,
            'complaints_completed': 0,
            'requests_total': 0,
            'requests_pending': 0,
            'requests_in_progress': 0,
            'requests_completed': 0,
            'tasks_total': 0,
            'tasks_pending': 0,
            'tasks_completed': 0,
            'completed_value': 0,
            'open_work': 0,
            'escalated_open': 0,
            'service_completion_rate': 0,
            'admin_coverage_rate': 0,
        },
        'district_insights': [],
        'taluka_performance': [],
        'hotspots': [],
        'recent_complaints': [],
        'recent_requests': [],
        'worker_summary': [],
        'escalations': [],
        'revenue_rows': [],
        'recent_assigned_tasks': [],
        'district_talukas': [],
        'district_taluka_admins': [],
        'chart_data': {
            'service_mix': {
                'labels': ['Pending', 'In Progress', 'Completed'],
                'complaints': [0, 0, 0],
                'requests': [0, 0, 0],
                'tasks': [0, 0, 0],
            },
            'taluka_performance': {
                'labels': [],
                'completion': [],
                'open_work': [],
            },
        },
        'district_map': _build_district_map_payload(admin_profile or {}, [], [], []),
    }


def _get_district_overview_data(admin_profile, cursor):
    district_id = admin_profile.get('district_id')
    district_name = admin_profile.get('district_name')

    cursor.execute("SELECT COUNT(*) AS total FROM talukas WHERE district_id = %s", (district_id,))
    talukas_count = (cursor.fetchone() or {}).get('total', 0)

    cursor.execute("""
        SELECT COUNT(*) AS total
        FROM villages
        WHERE taluka_id IN (SELECT id FROM talukas WHERE district_id = %s)
    """, (district_id,))
    villages_count = (cursor.fetchone() or {}).get('total', 0)

    cursor.execute("""
        SELECT COUNT(*) AS total
        FROM taluka_admins
        WHERE taluka_id IN (SELECT id FROM talukas WHERE district_id = %s)
    """, (district_id,))
    taluka_admin_count = (cursor.fetchone() or {}).get('total', 0)

    cursor.execute("""
        SELECT COUNT(*) AS total
        FROM users u
        LEFT JOIN villages v ON u.village_id = v.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas vt ON v.taluka_id = vt.id
        WHERE COALESCE(u.district_id, ut.district_id, vt.district_id) = %s
    """, (district_id,))
    citizens_count = (cursor.fetchone() or {}).get('total', 0)

    cursor.execute("""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN LOWER(COALESCE(vw.status, 'Active')) = 'active' THEN 1 ELSE 0 END) AS active_total
        FROM village_workers vw
        LEFT JOIN villages v ON vw.village_id = v.id
        LEFT JOIN talukas t ON v.taluka_id = t.id
        WHERE t.district_id = %s
    """, (district_id,))
    worker_counts = cursor.fetchone() or {}
    workers_count = worker_counts.get('total') or 0
    active_workers = worker_counts.get('active_total') or 0

    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN LOWER(COALESCE(status, 'Pending')) = 'pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN LOWER(COALESCE(status, 'Pending')) IN ('assigned', 'in progress', 'in_progress') THEN 1 ELSE 0 END) AS in_progress,
            SUM(CASE WHEN LOWER(COALESCE(status, 'Pending')) = 'completed' THEN 1 ELSE 0 END) AS completed
        FROM complaints
        WHERE district = %s
    """, (district_name,))
    complaints_stats = cursor.fetchone() or {}

    cursor.execute("""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'in_progress' THEN 1 ELSE 0 END) AS in_progress,
            SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'completed' THEN 1 ELSE 0 END) AS completed,
            COALESCE(SUM(CASE WHEN LOWER(COALESCE(r.status, 'pending')) = 'completed' THEN r.amount ELSE 0 END), 0) AS completed_value
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages uv ON u.village_id = uv.id
        LEFT JOIN talukas ut ON u.taluka_id = ut.id
        LEFT JOIN talukas uvt ON uv.taluka_id = uvt.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        LEFT JOIN villages wv ON vw.village_id = wv.id
        LEFT JOIN talukas wt ON wv.taluka_id = wt.id
        WHERE COALESCE(u.district_id, ut.district_id, uvt.district_id, wt.district_id) = %s
    """, (district_id,))
    requests_stats = cursor.fetchone() or {}

    cursor.execute("""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN t.status = 'pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN t.status = 'completed' THEN 1 ELSE 0 END) AS completed
        FROM tasks t
        INNER JOIN village_workers vw ON vw.id = t.worker_id
        INNER JOIN villages v ON vw.village_id = v.id
        INNER JOIN talukas tl ON v.taluka_id = tl.id
        WHERE tl.district_id = %s
    """, (district_id,))
    task_stats = cursor.fetchone() or {}

    all_complaints = _get_district_complaints(admin_profile, cursor)
    all_requests = _get_district_request_items(district_id, cursor)
    recent_assigned_tasks = _get_district_recent_manual_tasks(district_id, cursor, limit=8)
    map_assigned_tasks = _get_district_recent_manual_tasks(district_id, cursor, limit=None)
    worker_summary = _get_district_worker_summary(district_id, cursor, limit=8)
    taluka_performance = _get_district_taluka_performance(district_id, district_name, cursor)
    district_talukas = _get_district_taluka_rows(district_id, cursor)
    district_taluka_admins = _get_district_taluka_admin_options(district_id, cursor)
    hotspots = _get_district_hotspots(district_name, cursor, limit=6)

    escalations = sorted(
        [complaint for complaint in all_complaints if complaint.get('is_escalated')],
        key=lambda item: (
            -(item.get('hours_waiting') or 0),
            str(item.get('created_at') or ''),
        )
    )[:6]

    total_service_items = (
        (complaints_stats.get('total') or 0)
        + (requests_stats.get('total') or 0)
        + (task_stats.get('total') or 0)
    )
    total_completed_items = (
        (complaints_stats.get('completed') or 0)
        + (requests_stats.get('completed') or 0)
        + (task_stats.get('completed') or 0)
    )

    dashboard_stats = {
        'talukas': talukas_count or 0,
        'villages': villages_count or 0,
        'taluka_admins': taluka_admin_count or 0,
        'citizens': citizens_count or 0,
        'workers': workers_count or 0,
        'active_workers': active_workers or 0,
        'complaints_total': complaints_stats.get('total') or 0,
        'complaints_pending': complaints_stats.get('pending') or 0,
        'complaints_in_progress': complaints_stats.get('in_progress') or 0,
        'complaints_completed': complaints_stats.get('completed') or 0,
        'requests_total': requests_stats.get('total') or 0,
        'requests_pending': requests_stats.get('pending') or 0,
        'requests_in_progress': requests_stats.get('in_progress') or 0,
        'requests_completed': requests_stats.get('completed') or 0,
        'tasks_total': task_stats.get('total') or 0,
        'tasks_pending': task_stats.get('pending') or 0,
        'tasks_completed': task_stats.get('completed') or 0,
        'completed_value': requests_stats.get('completed_value') or 0,
        'open_work': (
            (complaints_stats.get('pending') or 0)
            + (complaints_stats.get('in_progress') or 0)
            + (requests_stats.get('pending') or 0)
            + (requests_stats.get('in_progress') or 0)
            + (task_stats.get('pending') or 0)
        ),
        'escalated_open': len(escalations),
        'service_completion_rate': round((total_completed_items / total_service_items) * 100) if total_service_items else 0,
        'admin_coverage_rate': round((taluka_admin_count / talukas_count) * 100) if talukas_count else 0,
    }

    revenue_rows = sorted(
        taluka_performance,
        key=lambda item: (
            -(item.get('completed_value') or 0),
            (item.get('name') or '').lower(),
        )
    )[:8]

    active_talukas = [item for item in taluka_performance if item.get('total_services')]
    best_taluka = active_talukas[0] if active_talukas else None
    watch_taluka = max(
        taluka_performance or [{}],
        key=lambda item: (
            item.get('open_work') or 0,
            item.get('total_services') or 0,
            100 - (item.get('completion_rate') or 0),
        )
    ) if taluka_performance else None
    top_hotspot = hotspots[0] if hotspots else None
    lead_worker = worker_summary[0] if worker_summary else None

    district_insights = [
        {
            'title': 'Best Taluka',
            'value': best_taluka.get('name') if best_taluka else 'Awaiting activity',
            'description': (
                f"{best_taluka.get('completion_rate', 0)}% completion with {best_taluka.get('open_work', 0)} open items."
                if best_taluka else
                'The highest-performing taluka will appear here after service activity is recorded.'
            ),
            'tone': 'mint',
            'icon': 'fa-circle-check',
        },
        {
            'title': 'Needs Review',
            'value': watch_taluka.get('name') if watch_taluka else 'All clear',
            'description': (
                f"{watch_taluka.get('open_work', 0)} open items need follow-up."
                if watch_taluka else
                'No taluka has an elevated backlog right now.'
            ),
            'tone': 'rose',
            'icon': 'fa-triangle-exclamation',
        },
        {
            'title': 'Top Hotspot',
            'value': top_hotspot.get('location_label') if top_hotspot else 'No complaint hotspots',
            'description': (
                f"{top_hotspot.get('open_complaints', 0)} open complaints across this area."
                if top_hotspot else
                'Complaint pressure by village will appear here as soon as reports come in.'
            ),
            'tone': 'amber',
            'icon': 'fa-fire-flame-curved',
        },
        {
            'title': 'Top Worker',
            'value': lead_worker.get('name') if lead_worker else 'No workers yet',
            'description': (
                f"{lead_worker.get('completed_tasks', 0)} completed assignments in {lead_worker.get('taluka_name', 'district scope')}."
                if lead_worker else
                'Worker performance rankings will update automatically once assignments are created.'
            ),
            'tone': 'navy',
            'icon': 'fa-user-shield',
        },
    ]

    chart_data = {
        'service_mix': {
            'labels': ['Pending', 'In Progress', 'Completed'],
            'complaints': [
                dashboard_stats['complaints_pending'],
                dashboard_stats['complaints_in_progress'],
                dashboard_stats['complaints_completed'],
            ],
            'requests': [
                dashboard_stats['requests_pending'],
                dashboard_stats['requests_in_progress'],
                dashboard_stats['requests_completed'],
            ],
            'tasks': [
                dashboard_stats['tasks_pending'],
                0,
                dashboard_stats['tasks_completed'],
            ],
        },
        'taluka_performance': {
            'labels': [item.get('name') for item in taluka_performance[:8]],
            'completion': [item.get('completion_rate') or 0 for item in taluka_performance[:8]],
            'open_work': [item.get('open_work') or 0 for item in taluka_performance[:8]],
        },
    }

    return {
        'dashboard_stats': dashboard_stats,
        'district_insights': district_insights,
        'taluka_performance': taluka_performance,
        'hotspots': hotspots,
        'recent_complaints': all_complaints[:6],
        'recent_requests': all_requests[:8],
        'worker_summary': worker_summary,
        'escalations': escalations,
        'revenue_rows': revenue_rows,
        'recent_assigned_tasks': recent_assigned_tasks,
        'district_talukas': district_talukas,
        'district_taluka_admins': district_taluka_admins,
        'chart_data': chart_data,
        'district_map': _build_district_map_payload(admin_profile, all_complaints, all_requests, map_assigned_tasks),
    }


def _get_citizen_dashboard_data(user_id, cursor):
    cursor.execute("""
        SELECT
            id,
            name,
            phone,
            email,
            district_id,
            taluka_id,
            village_id,
            created_at
        FROM users
        WHERE id = %s
    """, (user_id,))
    user_profile = cursor.fetchone() or {}

    cursor.execute("""
        SELECT id, name
        FROM districts
        ORDER BY name ASC
    """)
    district_options = cursor.fetchall() or []

    ensure_complaint_workflow_columns(cursor)
    cursor.execute("""
        SELECT
            COUNT(*) AS total,
            SUM(CASE WHEN LOWER(status) = 'pending' THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN LOWER(status) IN ('assigned', 'in progress', 'in_progress') THEN 1 ELSE 0 END) AS in_progress,
            SUM(CASE WHEN LOWER(status) = 'completed' THEN 1 ELSE 0 END) AS completed
        FROM complaints
        WHERE user_id = %s
    """, (user_id,))
    stats_row = cursor.fetchone() or {}
    stats = {
        'total': stats_row.get('total') or 0,
        'pending': stats_row.get('pending') or 0,
        'in_progress': stats_row.get('in_progress') or 0,
        'completed': stats_row.get('completed') or 0,
    }

    recent_complaints = _get_user_complaints(user_id, cursor, limit=5)

    ensure_request_location_columns(cursor)
    cursor.execute("""
        SELECT
            r.id,
            COALESCE(r.garbage_type, 'Garbage Collection') AS garbage_type,
            COALESCE(r.status, 'pending') AS status,
            COALESCE(r.amount, 0) AS amount,
            r.latitude,
            r.longitude,
            r.location_accuracy_meters,
            r.created_at,
            COALESCE(v.name, 'Unknown Village') AS village_name,
            COALESCE(t.name, 'Unknown Taluka') AS taluka_name,
            COALESCE(vw.name, 'Pending Assignment') AS worker_name
        FROM requests r
        LEFT JOIN users u ON r.user_id = u.id
        LEFT JOIN villages v ON u.village_id = v.id
        LEFT JOIN talukas t ON COALESCE(u.taluka_id, v.taluka_id) = t.id
        LEFT JOIN village_workers vw ON r.worker_id = vw.id
        WHERE r.user_id = %s
        ORDER BY r.created_at DESC
        LIMIT 10
    """, (user_id,))
    user_requests = cursor.fetchall() or []
    for request_item in user_requests:
        request_item.update(build_location_payload(
            latitude=request_item.get('latitude'),
            longitude=request_item.get('longitude'),
            accuracy=request_item.get('location_accuracy_meters'),
            query=_build_map_query(
                request_item.get('village_name'),
                request_item.get('taluka_name'),
                'Maharashtra',
                'India',
            ),
        ))

    cursor.execute("""
        SELECT
            p.id,
            p.request_id,
            p.total,
            p.owner_share,
            p.admin_share,
            p.worker_share,
            p.created_at,
            COALESCE(r.garbage_type, 'General Payment') AS garbage_type
        FROM payments p
        INNER JOIN requests r ON p.request_id = r.id
        WHERE r.user_id = %s
        ORDER BY p.created_at DESC
        LIMIT 10
    """, (user_id,))
    payment_history = cursor.fetchall() or []

    return {
        'stats': stats,
        'recent_complaints': recent_complaints,
        'user_requests': user_requests,
        'payment_history': payment_history,
        'user_profile': user_profile,
        'district_options': district_options,
    }

# --- 🟢 REGISTRATION ROUTE ---
@auth_bp.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        name = (request.form.get('name') or '').strip()
        email = (request.form.get('email') or '').strip()
        phone = (request.form.get('phone') or '').strip()
        password = request.form.get('password')
        role = request.form.get('role', 'user')
        aadhaar_no = (request.form.get('aadhaar_no') or '').strip()
        state_id = _safe_int(request.form.get('state_id'))
        dist_id = _safe_int(request.form.get('district_id'))
        tal_id = _safe_int(request.form.get('taluka_id'))
        vil_id = _safe_int(request.form.get('village_id'))

        if not name or not email or not password:
            flash("Name, Email, and Password are required!", "danger")
            return redirect(url_for('auth_bp.register'))
        if not phone:
            flash("Phone number is required!", "danger")
            return redirect(url_for('auth_bp.register'))
        is_password_valid, missing_password_rules = _validate_registration_password(password)
        if not is_password_valid:
            flash(
                "Password must include " + ", ".join(missing_password_rules) + ".",
                "danger"
            )
            return redirect(url_for('auth_bp.register'))
        if not state_id:
            flash("Please select a state before registering.", "danger")
            return redirect(url_for('auth_bp.register'))
        if role == 'district_admin' and not dist_id:
            flash("District selection is required for district admin registration.", "danger")
            return redirect(url_for('auth_bp.register'))
        if role == 'admin' and not tal_id:
            flash("Taluka selection is required for taluka admin registration.", "danger")
            return redirect(url_for('auth_bp.register'))
        if role in ('worker', 'user') and not vil_id:
            flash("Village selection is required for this registration.", "danger")
            return redirect(url_for('auth_bp.register'))
        if staff_role_requires_verification(role) and not aadhaar_no:
            flash("Aadhaar card number is required for staff background verification.", "danger")
            return redirect(url_for('auth_bp.register'))

        table_name = ROLE_MAP.get(role, 'users')
        hashed_pw = generate_password_hash(password)
        
        conn = get_db()
        cursor = conn.cursor(dictionary=True)

        try:
            ensure_state_admin_table(cursor)
            if staff_role_requires_verification(role):
                ensure_staff_verification_columns(cursor, table_name)

            check_query = f"SELECT * FROM {table_name} WHERE email=%s OR phone=%s"
            cursor.execute(check_query, (email, phone))
            existing_user = cursor.fetchone()

            if role == 'state_admin':
                cursor.execute(
                    """
                    SELECT id
                    FROM state_admins
                    WHERE state_id = %s
                    LIMIT 1
                    """,
                    (state_id,)
                )
                existing_scope_admin = cursor.fetchone()
                if existing_scope_admin and (
                    not existing_user or existing_scope_admin.get('id') != existing_user.get('id')
                ):
                    flash("A state admin already exists for the selected state.", "danger")
                    return redirect(url_for('auth_bp.register'))

            if role == 'district_admin':
                cursor.execute(
                    """
                    SELECT id
                    FROM district_admins
                    WHERE district_id = %s
                    LIMIT 1
                    """,
                    (dist_id,)
                )
                existing_scope_admin = cursor.fetchone()
                if existing_scope_admin and (
                    not existing_user or existing_scope_admin.get('id') != existing_user.get('id')
                ):
                    flash("A district admin already exists for the selected district.", "danger")
                    return redirect(url_for('auth_bp.register'))

            if role == 'admin':
                cursor.execute(
                    """
                    SELECT id
                    FROM taluka_admins
                    WHERE taluka_id = %s
                    LIMIT 1
                    """,
                    (tal_id,)
                )
                existing_scope_admin = cursor.fetchone()
                if existing_scope_admin and (
                    not existing_user or existing_scope_admin.get('id') != existing_user.get('id')
                ):
                    flash("A taluka admin already exists for the selected taluka.", "danger")
                    return redirect(url_for('auth_bp.register'))

            if existing_user and not (
                staff_role_requires_verification(role)
                and normalize_verification_status(existing_user.get('verification_status')) == 'rejected'
            ):
                flash(f"User with this Email or Mobile already exists in {role_display_name(role)}.", "danger")
                return redirect(url_for('auth_bp.register'))

            verification_files = {'aadhaar_photo_path': None, 'selfie_photo_path': None}
            if staff_role_requires_verification(role):
                verification_files, upload_error = _save_staff_verification_files(role)
                if upload_error:
                    flash(upload_error, "danger")
                    return redirect(url_for('auth_bp.register'))

            if existing_user:
                if role == 'district_admin':
                    query = """
                        UPDATE district_admins
                        SET name=%s, email=%s, phone=%s, password=%s, district_id=%s,
                            aadhaar_no=%s, aadhaar_photo_path=%s, selfie_photo_path=%s,
                            verification_status='pending', verification_notes=NULL,
                            verified_by_role=NULL, verified_by_id=NULL, verified_at=NULL
                        WHERE id=%s
                    """
                    cursor.execute(query, (
                        name, email, phone, hashed_pw, dist_id,
                        aadhaar_no, verification_files['aadhaar_photo_path'], verification_files['selfie_photo_path'],
                        existing_user['id'],
                    ))
                elif role == 'admin':
                    query = """
                        UPDATE taluka_admins
                        SET name=%s, email=%s, phone=%s, password=%s, taluka_id=%s,
                            aadhaar_no=%s, aadhaar_photo_path=%s, selfie_photo_path=%s,
                            verification_status='pending', verification_notes=NULL,
                            verified_by_role=NULL, verified_by_id=NULL, verified_at=NULL
                        WHERE id=%s
                    """
                    cursor.execute(query, (
                        name, email, phone, hashed_pw, tal_id,
                        aadhaar_no, verification_files['aadhaar_photo_path'], verification_files['selfie_photo_path'],
                        existing_user['id'],
                    ))
                else:
                    query = """
                        UPDATE village_workers
                        SET name=%s, email=%s, phone=%s, password=%s, village_id=%s,
                            aadhaar_no=%s, aadhaar_photo_path=%s, selfie_photo_path=%s,
                            verification_status='pending', verification_notes=NULL,
                            verified_by_role=NULL, verified_by_id=NULL, verified_at=NULL
                        WHERE id=%s
                    """
                    cursor.execute(query, (
                        name, email, phone, hashed_pw, vil_id,
                        aadhaar_no, verification_files['aadhaar_photo_path'], verification_files['selfie_photo_path'],
                        existing_user['id'],
                    ))
                conn.commit()
                flash("Registration submitted again. Verification is pending.", "success")
                return redirect(url_for('auth_bp.login', role=role))

            if role == 'state_admin':
                query = "INSERT INTO state_admins (name, email, phone, password, state_id) VALUES (%s, %s, %s, %s, %s)"
                cursor.execute(query, (name, email, phone, hashed_pw, state_id))
            elif role == 'district_admin':
                query = """
                    INSERT INTO district_admins (
                        name, email, phone, password, district_id,
                        aadhaar_no, aadhaar_photo_path, selfie_photo_path, verification_status
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'pending')
                """
                cursor.execute(query, (
                    name, email, phone, hashed_pw, dist_id,
                    aadhaar_no, verification_files['aadhaar_photo_path'], verification_files['selfie_photo_path'],
                ))
            elif role == 'admin':
                query = """
                    INSERT INTO taluka_admins (
                        name, email, phone, password, taluka_id,
                        aadhaar_no, aadhaar_photo_path, selfie_photo_path, verification_status
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'pending')
                """
                cursor.execute(query, (
                    name, email, phone, hashed_pw, tal_id,
                    aadhaar_no, verification_files['aadhaar_photo_path'], verification_files['selfie_photo_path'],
                ))
            elif role == 'worker':
                query = """
                    INSERT INTO village_workers (
                        name, email, phone, password, village_id,
                        aadhaar_no, aadhaar_photo_path, selfie_photo_path, verification_status
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'pending')
                """
                cursor.execute(query, (
                    name, email, phone, hashed_pw, vil_id,
                    aadhaar_no, verification_files['aadhaar_photo_path'], verification_files['selfie_photo_path'],
                ))
            else:
                query = """
                    INSERT INTO users (
                        name, email, phone, password, role, state_id, district_id, taluka_id, village_id
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """
                cursor.execute(query, (name, email, phone, hashed_pw, 'user', state_id, dist_id, tal_id, vil_id))

            conn.commit()
            if staff_role_requires_verification(role):
                flash("Registration successful. Your verification is pending approval.", "success")
            else:
                flash("Registration Successful! Please login.", "success")
            return redirect(url_for('auth_bp.login', role=role))

        except Exception as e:
            conn.rollback()
            print(f"Database Error: {e}")
            flash("Registration failed due to a server error.", "danger")
        finally:
            conn.close()

    # GET Request: Fetch states for the dropdown
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("SELECT id, name FROM states ORDER BY name ASC")
    states = cursor.fetchall()
    conn.close()
    return render_template('auth/register.html', states=states)


@auth_bp.route('/find-village')
def find_village():
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT id, name FROM states ORDER BY name ASC")
        states = cursor.fetchall() or []

        cursor.execute("""
            SELECT
                (SELECT COUNT(*) FROM states) AS states_count,
                (SELECT COUNT(*) FROM districts) AS districts_count,
                (SELECT COUNT(*) FROM talukas) AS talukas_count,
                (SELECT COUNT(*) FROM villages) AS villages_count
        """)
        directory_stats = cursor.fetchone() or {}
    finally:
        conn.close()

    return render_template(
        'auth/find_village.html',
        states=states,
        directory_stats=directory_stats,
    )


# --- 🔵 LOGIN ROUTE ---
@auth_bp.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        identifier = (request.form.get('identifier') or '').strip()
        password = request.form.get('password')
        form_role = request.form.get('role') 

        table_name = ROLE_MAP.get(form_role, 'users')
        
        conn = get_db()
        cursor = conn.cursor(dictionary=True)
        try:
            ensure_state_admin_table(cursor)
            if staff_role_requires_verification(form_role):
                ensure_staff_verification_columns(cursor, table_name)

            query = f"SELECT * FROM {table_name} WHERE email = %s OR phone = %s"
            cursor.execute(query, (identifier, identifier))
            user = cursor.fetchone()

            if user and check_password_hash(user['password'], password):
                if staff_role_requires_verification(form_role):
                    verification_status = normalize_verification_status(user.get('verification_status'))
                    if verification_status == 'rejected':
                        flash(STAFF_REJECTED_LOGIN_MESSAGE, "danger")
                        return redirect(url_for('auth_bp.login', role=form_role))
                    if verification_status != 'approved':
                        flash(STAFF_PENDING_LOGIN_MESSAGE, "warning")
                        return redirect(url_for('auth_bp.login', role=form_role))

                session.clear()
                session['user_id'] = user['id']
                session['user_name'] = user['name']
                session['email'] = user.get('email')
                session['role'] = form_role
                _set_session_scope_data(form_role, user, cursor)
                
                flash(f"Welcome back, {user['name']}!", "success")
                
                redirect_map = {
                    'state_admin': 'auth_bp.state_dashboard',
                    'district_admin': 'auth_bp.district_dashboard',
                    'admin': 'auth_bp.taluka_dashboard',
                    'worker': 'auth_bp.worker_dashboard',
                    'user': 'auth_bp.citizen_dashboard'
                }
                return redirect(url_for(redirect_map.get(form_role, 'auth_bp.citizen_dashboard')))
            
            flash("Invalid credentials for the selected role.", "danger")
        except Exception as e:
            print(f"Login Error: {e}")
            flash("An error occurred during login.", "danger")
        finally:
            conn.close()
            
    return render_template('auth/login.html')

# --- 🔴 LOGOUT ROUTE ---
@auth_bp.route('/logout')
def logout():
    reason = (request.args.get('reason') or '').strip().lower()
    session.clear()
    if reason == 'expired':
        flash("Your session expired after 60 minutes. Please log in again.", "warning")
    else:
        flash("Successfully logged out.", "success")
    return redirect(url_for('auth_bp.login'))

# --- 🟡 DASHBOARD ROUTES (Unified Check) ---

@auth_bp.route('/citizen-dashboard')
def citizen_dashboard():
    if session.get('role') != 'user':
        return redirect(url_for('auth_bp.login'))
    
    user_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    
    try:
        data = _get_citizen_dashboard_data(user_id, cursor)
    except Exception as e:
        print(f"Error fetching stats: {e}")
        data = {
            'stats': {'total': 0, 'pending': 0, 'in_progress': 0, 'completed': 0},
            'recent_complaints': [],
            'user_requests': [],
            'payment_history': [],
            'user_profile': {},
        }
    finally:
        conn.close()

    return render_template('dashboard/citizen_dash.html', **data, citizen_page='dashboard')

@auth_bp.route('/district-dashboard')
def district_dashboard():
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    admin_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(admin_id, cursor)
        if not district_profile:
            flash("Unable to find your district profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        try:
            overview = _get_district_overview_data(district_profile, cursor)
        except Exception as e:
            print(f"Error loading district dashboard overview: {e}")
            flash("Some district dashboard data is unavailable right now. Showing the basic district view instead.", "warning")
            overview = _empty_district_overview_data(district_profile)
        overview['announcement_feed'] = fetch_visible_announcements(
            cursor,
            'district_admin',
            state_id=district_profile.get('state_id'),
            district_id=district_profile.get('district_id'),
            limit=6,
        )

        return render_template(
            'dashboard/district_admin.html',
            district_profile=district_profile,
            district_page='dashboard',
            **overview
        )
    except Exception as e:
        print(f"Error loading district dashboard: {e}")
        flash("Unable to load the district dashboard right now.", "danger")
        return redirect(url_for('auth_bp.login'))
    finally:
        conn.close()


@auth_bp.route('/district-reports')
def district_reports():
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))
    return _render_report_center('district_admin')


@auth_bp.route('/district-reports/export/<format_name>')
def district_report_export(format_name):
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))
    return _export_report_center('district_admin', format_name)


@auth_bp.route('/district-equipment')
def district_equipment():
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        if not district_profile:
            flash("Unable to find your district profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        return render_template(
            'dashboard/district_equipment.html',
            district_profile=district_profile,
            district_page='equipment',
            **_build_district_equipment_context(district_profile, session.get('user_id'), cursor),
        )
    except Exception as e:
        print(f"Error loading district equipment dashboard: {e}")
        flash("Unable to load district equipment data right now.", "danger")
        return redirect(url_for('auth_bp.district_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/district-equipment-issue', methods=['POST'])
def district_issue_equipment():
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.district_equipment')
    request_id = _safe_int(request.form.get('request_id'))
    taluka_admin_id = _safe_int(request.form.get('taluka_admin_id'))
    item_key = normalize_equipment_item(request.form.get('item_key'))
    notes = (request.form.get('notes') or '').strip()

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        if not district_profile:
            flash("District profile not found.", "warning")
            return redirect(redirect_target)

        taluka_admin = None
        related_request_id = None
        requested_quantity = None

        if request_id:
            request_record = _get_equipment_request_record(request_id, cursor)
            pending_quantity = max(
                (request_record or {}).get('quantity_requested', 0) - (request_record or {}).get('quantity_fulfilled', 0),
                0,
            )
            if not request_record or request_record.get('status') != 'pending':
                flash("That equipment request is no longer pending.", "warning")
                return redirect(redirect_target)
            if request_record.get('target_role') != 'district_admin' or request_record.get('target_scope_id') != district_profile.get('district_id'):
                flash("You can only fulfill requests assigned to your district office.", "warning")
                return redirect(redirect_target)

            taluka_admin = _get_district_taluka_admin_record(
                district_profile.get('district_id'),
                request_record.get('requester_id'),
                cursor,
            )
            if not taluka_admin:
                flash("This taluka request is outside your district scope.", "warning")
                return redirect(redirect_target)

            taluka_admin_id = taluka_admin.get('id')
            item_key = request_record.get('item_key')
            related_request_id = request_record.get('id')
            requested_quantity = pending_quantity or request_record.get('quantity_requested') or 1

        if not taluka_admin and taluka_admin_id:
            taluka_admin = _get_district_taluka_admin_record(
                district_profile.get('district_id'),
                taluka_admin_id,
                cursor,
            )

        quantity = _positive_int(request.form.get('quantity'), default=requested_quantity)

        if not taluka_admin or not item_key or quantity is None:
            flash("Choose a taluka admin, equipment item, and quantity before assigning stock.", "warning")
            return redirect(redirect_target)

        transfer_note = _merge_equipment_note(
            f"Issued to {taluka_admin.get('name') or 'Taluka Admin'} for {taluka_admin.get('taluka_name') or 'assigned taluka'}.",
            notes,
        )
        issue_inventory(
            cursor,
            actor_role='district_admin',
            actor_id=session.get('user_id'),
            from_scope_role='district',
            from_scope_id=district_profile.get('district_id'),
            to_scope_role='taluka',
            to_scope_id=taluka_admin.get('taluka_id'),
            item_key=item_key,
            quantity=quantity,
            notes=transfer_note,
            related_request_id=related_request_id,
            transaction_type='issue',
        )
        if related_request_id:
            mark_request_fulfilled(cursor, related_request_id, quantity, transfer_note)

        conn.commit()
        flash("Equipment stock assigned to the taluka successfully.", "success")
    except ValueError as e:
        conn.rollback()
        flash(str(e), "warning")
    except Exception as e:
        conn.rollback()
        print(f"Error issuing district equipment: {e}")
        flash("Unable to assign equipment stock right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/district-equipment-restock', methods=['POST'])
def district_restock_equipment():
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.district_equipment')
    item_key = normalize_equipment_item(request.form.get('item_key'))
    quantity = _positive_int(request.form.get('quantity'))
    notes = (request.form.get('notes') or '').strip()

    if not item_key or quantity is None:
        flash("Choose equipment and quantity before updating district reserve stock.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        if not district_profile:
            flash("District profile not found.", "warning")
            return redirect(redirect_target)

        restock_inventory(
            cursor,
            actor_role='district_admin',
            actor_id=session.get('user_id'),
            scope_role='district',
            scope_id=district_profile.get('district_id'),
            item_key=item_key,
            quantity=quantity,
            notes=_merge_equipment_note("District reserve stock updated.", notes),
        )
        conn.commit()
        flash("District reserve stock updated successfully.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error restocking district equipment: {e}")
        flash("Unable to update district reserve stock right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/district-equipment-request-stock', methods=['POST'])
def district_request_equipment_stock():
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.district_equipment')
    item_key = normalize_equipment_item(request.form.get('item_key'))
    quantity = _positive_int(request.form.get('quantity'))
    reason = (request.form.get('reason') or '').strip()
    notes = (request.form.get('notes') or '').strip()
    priority = (request.form.get('priority') or 'normal').strip().lower()
    if priority not in {'low', 'normal', 'high', 'urgent'}:
        priority = 'normal'

    if not item_key or quantity is None or not reason:
        flash("Choose equipment, quantity, and reason before requesting stock from the state office.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        if not district_profile:
            flash("District profile not found.", "warning")
            return redirect(redirect_target)

        create_equipment_request(
            cursor,
            requester_role='district_admin',
            requester_id=session.get('user_id'),
            requester_scope_role='district',
            requester_scope_id=district_profile.get('district_id'),
            target_role='state_admin',
            target_scope_role='state',
            target_scope_id=district_profile.get('state_id'),
            item_key=item_key,
            quantity_requested=quantity,
            reason=reason,
            priority=priority,
            notes=notes or None,
        )
        conn.commit()
        flash("Bulk stock request sent to the state office.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error creating district equipment request: {e}")
        flash("Unable to send the state stock request right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/district-equipment-request/<int:request_id>/reject', methods=['POST'])
def district_reject_equipment_request(request_id):
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.district_equipment')
    notes = (request.form.get('notes') or '').strip()

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        request_record = _get_equipment_request_record(request_id, cursor)
        if not district_profile or not request_record or request_record.get('status') != 'pending':
            flash("That equipment request cannot be rejected now.", "warning")
            return redirect(redirect_target)
        if request_record.get('target_role') != 'district_admin' or request_record.get('target_scope_id') != district_profile.get('district_id'):
            flash("You can only reject requests assigned to your district office.", "warning")
            return redirect(redirect_target)

        mark_request_rejected(cursor, request_id, _merge_equipment_note("Rejected by district office.", notes))
        conn.commit()
        flash("Equipment request rejected.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error rejecting district equipment request: {e}")
        flash("Unable to reject the equipment request right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/state-dashboard')
def state_dashboard():
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        state_profile = _get_state_admin_profile(session.get('user_id'), cursor)
        if not state_profile:
            flash("Unable to find your state admin profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        overview = _get_state_overview_data(state_profile, cursor)
        return render_template(
            'dashboard/state_admin.html',
            state_profile=state_profile,
            state_page='dashboard',
            **overview
        )
    except Exception as e:
        print(f"Error loading state dashboard: {e}")
        flash("Unable to load the state dashboard right now.", "danger")
        return redirect(url_for('auth_bp.login'))
    finally:
        conn.close()


@auth_bp.route('/state-reports')
def state_reports():
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))
    return _render_report_center('state_admin')


@auth_bp.route('/state-reports/export/<format_name>')
def state_report_export(format_name):
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))
    return _export_report_center('state_admin', format_name)


@auth_bp.route('/state-equipment')
def state_equipment():
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        state_profile = _get_state_admin_profile(session.get('user_id'), cursor)
        if not state_profile:
            flash("Unable to find your state admin profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        return render_template(
            'dashboard/state_equipment.html',
            state_profile=state_profile,
            state_page='equipment',
            **_build_state_equipment_context(state_profile, cursor),
        )
    except Exception as e:
        print(f"Error loading state equipment dashboard: {e}")
        flash("Unable to load state equipment data right now.", "danger")
        return redirect(url_for('auth_bp.state_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/state-inquiries')
def state_inquiries():
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        state_profile = _get_state_admin_profile(session.get('user_id'), cursor)
        if not state_profile:
            flash("Unable to find your state admin profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        return render_template(
            'dashboard/state_inquiries.html',
            state_profile=state_profile,
            state_page='inquiries',
            inquiries=fetch_inquiries(cursor),
        )
    except Exception as e:
        print(f"Error loading state inquiries: {e}")
        flash("Unable to load inquiries right now.", "danger")
        return redirect(url_for('auth_bp.state_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/state-equipment-issue', methods=['POST'])
def state_issue_equipment():
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.state_equipment')
    request_id = _safe_int(request.form.get('request_id'))
    district_admin_id = _safe_int(request.form.get('district_admin_id'))
    item_key = normalize_equipment_item(request.form.get('item_key'))
    notes = (request.form.get('notes') or '').strip()

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        state_profile = _get_state_admin_profile(session.get('user_id'), cursor)
        if not state_profile:
            flash("State profile not found.", "warning")
            return redirect(redirect_target)

        district_admin = None
        related_request_id = None
        requested_quantity = None

        if request_id:
            request_record = _get_equipment_request_record(request_id, cursor)
            pending_quantity = max(
                (request_record or {}).get('quantity_requested', 0) - (request_record or {}).get('quantity_fulfilled', 0),
                0,
            )
            if not request_record or request_record.get('status') != 'pending':
                flash("That equipment request is no longer pending.", "warning")
                return redirect(redirect_target)
            if request_record.get('target_role') != 'state_admin' or request_record.get('target_scope_id') != state_profile.get('state_id'):
                flash("You can only fulfill requests assigned to your state office.", "warning")
                return redirect(redirect_target)

            district_admins = _get_state_district_admin_options(state_profile.get('state_id'), cursor)
            district_admin = next((item for item in district_admins if item.get('id') == request_record.get('requester_id')), None)
            if not district_admin:
                flash("This district request is outside your state scope.", "warning")
                return redirect(redirect_target)

            district_admin_id = district_admin.get('id')
            item_key = request_record.get('item_key')
            related_request_id = request_record.get('id')
            requested_quantity = pending_quantity or request_record.get('quantity_requested') or 1

        if not district_admin and district_admin_id:
            district_admins = _get_state_district_admin_options(state_profile.get('state_id'), cursor)
            district_admin = next((item for item in district_admins if item.get('id') == district_admin_id), None)

        quantity = _positive_int(request.form.get('quantity'), default=requested_quantity)

        if not district_admin or not item_key or quantity is None:
            flash("Choose a district admin, equipment item, and quantity before assigning stock.", "warning")
            return redirect(redirect_target)

        transfer_note = _merge_equipment_note(
            f"Issued to {district_admin.get('name') or 'District Admin'} for {district_admin.get('district_name') or 'assigned district'}.",
            notes,
        )
        issue_inventory(
            cursor,
            actor_role='state_admin',
            actor_id=session.get('user_id'),
            from_scope_role='state',
            from_scope_id=state_profile.get('state_id'),
            to_scope_role='district',
            to_scope_id=district_admin.get('district_id'),
            item_key=item_key,
            quantity=quantity,
            notes=transfer_note,
            related_request_id=related_request_id,
            transaction_type='issue',
        )
        if related_request_id:
            mark_request_fulfilled(cursor, related_request_id, quantity, transfer_note)

        conn.commit()
        flash("Equipment stock assigned to the district successfully.", "success")
    except ValueError as e:
        conn.rollback()
        flash(str(e), "warning")
    except Exception as e:
        conn.rollback()
        print(f"Error issuing state equipment: {e}")
        flash("Unable to assign equipment stock right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/state-equipment-restock', methods=['POST'])
def state_restock_equipment():
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.state_equipment')
    item_key = normalize_equipment_item(request.form.get('item_key'))
    quantity = _positive_int(request.form.get('quantity'))
    notes = (request.form.get('notes') or '').strip()

    if not item_key or quantity is None:
        flash("Choose equipment and quantity before updating state reserve stock.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        state_profile = _get_state_admin_profile(session.get('user_id'), cursor)
        if not state_profile:
            flash("State profile not found.", "warning")
            return redirect(redirect_target)

        restock_inventory(
            cursor,
            actor_role='state_admin',
            actor_id=session.get('user_id'),
            scope_role='state',
            scope_id=state_profile.get('state_id'),
            item_key=item_key,
            quantity=quantity,
            notes=_merge_equipment_note("State reserve stock updated.", notes),
        )
        conn.commit()
        flash("State reserve stock updated successfully.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error restocking state equipment: {e}")
        flash("Unable to update state reserve stock right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/state-equipment-request/<int:request_id>/reject', methods=['POST'])
def state_reject_equipment_request(request_id):
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.state_equipment')
    notes = (request.form.get('notes') or '').strip()

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        state_profile = _get_state_admin_profile(session.get('user_id'), cursor)
        request_record = _get_equipment_request_record(request_id, cursor)
        if not state_profile or not request_record or request_record.get('status') != 'pending':
            flash("That equipment request cannot be rejected now.", "warning")
            return redirect(redirect_target)
        if request_record.get('target_role') != 'state_admin' or request_record.get('target_scope_id') != state_profile.get('state_id'):
            flash("You can only reject requests assigned to your state office.", "warning")
            return redirect(redirect_target)

        mark_request_rejected(cursor, request_id, _merge_equipment_note("Rejected by state office.", notes))
        conn.commit()
        flash("Equipment request rejected.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error rejecting state equipment request: {e}")
        flash("Unable to reject the equipment request right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/post-announcement', methods=['POST'])
def post_announcement():
    actor_role = session.get('role')
    if actor_role not in ('district_admin', 'state_admin'):
        return redirect(url_for('auth_bp.login'))

    title = (request.form.get('title') or '').strip()
    message = (request.form.get('message') or '').strip()
    redirect_target = _post_redirect_target(
        'auth_bp.state_dashboard' if actor_role == 'state_admin' else 'auth_bp.district_dashboard'
    )

    if not title or not message:
        flash("Announcement title and message are required.", "danger")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        ensure_announcements_table(cursor)

        state_id = session.get('scope_state_id')
        district_id = session.get('scope_district_id')
        author_name = session.get('user_name')

        if actor_role == 'state_admin':
            if not state_id:
                state_profile = _get_state_admin_profile(session.get('user_id'), cursor)
                state_id = state_profile.get('state_id') if state_profile else None
        else:
            if not state_id or not district_id:
                district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
                state_id = district_profile.get('state_id') if district_profile else None
                district_id = district_profile.get('district_id') if district_profile else None

        cursor.execute("""
            INSERT INTO announcements (
                title,
                message,
                author_role,
                author_id,
                author_name,
                target_roles,
                state_id,
                district_id
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """, (
            title,
            message,
            actor_role,
            session.get('user_id'),
            author_name,
            encode_target_roles(ANNOUNCEMENT_TARGET_ROLES),
            state_id,
            None if actor_role == 'state_admin' else district_id,
        ))
        conn.commit()
        flash("Announcement sent successfully.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error posting announcement: {e}")
        flash("Unable to send the announcement right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/mark-announcements-seen', methods=['POST'])
def mark_announcements_seen_route():
    actor_role = session.get('role')
    if actor_role not in ('worker', 'admin', 'district_admin'):
        return jsonify({'ok': False, 'message': 'Unsupported role.'}), 403

    payload = request.get_json(silent=True) or {}
    announcement_ids = payload.get('announcement_ids') or []

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        ensure_announcements_table(cursor)
        ensure_announcement_views_table(cursor)
        marked_count = mark_announcements_seen(
            cursor,
            actor_role,
            session.get('user_id'),
            announcement_ids,
        )
        conn.commit()
        return jsonify({'ok': True, 'marked_count': marked_count})
    except Exception as e:
        conn.rollback()
        print(f"Error marking announcements seen: {e}")
        return jsonify({'ok': False, 'message': 'Unable to update announcement status.'}), 500
    finally:
        conn.close()


@auth_bp.route('/taluka-verify-worker/<int:worker_id>', methods=['POST'])
def taluka_verify_worker(worker_id):
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    decision = normalize_verification_status(request.form.get('decision'))
    notes = (request.form.get('verification_notes') or '').strip()
    redirect_target = _post_redirect_target('auth_bp.taluka_workers')

    if decision not in ('approved', 'rejected'):
        flash("Choose whether to approve or reject the worker verification.", "danger")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        worker_record = _get_taluka_worker_record(admin_profile.get('taluka_id'), worker_id, cursor) if admin_profile else None
        if not worker_record:
            flash("That worker is outside your taluka scope.", "warning")
            return redirect(redirect_target)

        ensure_staff_verification_columns(cursor, 'village_workers')
        cursor.execute("""
            UPDATE village_workers
            SET verification_status=%s,
                verification_notes=%s,
                verified_by_role='admin',
                verified_by_id=%s,
                verified_at=NOW()
            WHERE id=%s
        """, (decision, notes or None, session.get('user_id'), worker_id))
        conn.commit()
        flash(f"Worker verification marked as {decision}.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error verifying worker: {e}")
        flash("Unable to update worker verification right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/district-verify-taluka-admin/<int:taluka_admin_id>', methods=['POST'])
def district_verify_taluka_admin(taluka_admin_id):
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    decision = normalize_verification_status(request.form.get('decision'))
    notes = (request.form.get('verification_notes') or '').strip()
    redirect_target = _post_redirect_target('auth_bp.district_taluka_admins_page')

    if decision not in ('approved', 'rejected'):
        flash("Choose whether to approve or reject the taluka admin verification.", "danger")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        taluka_admin_record = _get_district_taluka_admin_record(district_profile.get('district_id'), taluka_admin_id, cursor) if district_profile else None
        if not taluka_admin_record:
            flash("That taluka admin is outside your district scope.", "warning")
            return redirect(redirect_target)

        ensure_staff_verification_columns(cursor, 'taluka_admins')
        cursor.execute("""
            UPDATE taluka_admins
            SET verification_status=%s,
                verification_notes=%s,
                verified_by_role='district_admin',
                verified_by_id=%s,
                verified_at=NOW()
            WHERE id=%s
        """, (decision, notes or None, session.get('user_id'), taluka_admin_id))
        conn.commit()
        flash(f"Taluka admin verification marked as {decision}.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error verifying taluka admin: {e}")
        flash("Unable to update taluka admin verification right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/state-verify-district-admin/<int:district_admin_id>', methods=['POST'])
def state_verify_district_admin(district_admin_id):
    if session.get('role') != 'state_admin':
        return redirect(url_for('auth_bp.login'))

    decision = normalize_verification_status(request.form.get('decision'))
    notes = (request.form.get('verification_notes') or '').strip()
    redirect_target = _post_redirect_target('auth_bp.state_dashboard')

    if decision not in ('approved', 'rejected'):
        flash("Choose whether to approve or reject the district admin verification.", "danger")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        state_profile = _get_state_admin_profile(session.get('user_id'), cursor)
        district_admins = _get_state_district_admin_options(state_profile.get('state_id'), cursor) if state_profile else []
        if district_admin_id not in {item.get('id') for item in district_admins}:
            flash("That district admin is outside your state scope.", "warning")
            return redirect(redirect_target)

        ensure_staff_verification_columns(cursor, 'district_admins')
        cursor.execute("""
            UPDATE district_admins
            SET verification_status=%s,
                verification_notes=%s,
                verified_by_role='state_admin',
                verified_by_id=%s,
                verified_at=NOW()
            WHERE id=%s
        """, (decision, notes or None, session.get('user_id'), district_admin_id))
        conn.commit()
        flash(f"District admin verification marked as {decision}.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error verifying district admin: {e}")
        flash("Unable to update district admin verification right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/district-talukas')
def district_talukas_page():
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        if not district_profile:
            flash("Unable to find your district profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        talukas = _get_district_taluka_rows(district_profile.get('district_id'), cursor)
        return render_template(
            'dashboard/district_talukas.html',
            district_profile=district_profile,
            district_talukas=talukas,
            district_page='talukas'
        )
    except Exception as e:
        print(f"Error loading district talukas: {e}")
        flash("Unable to load district talukas right now.", "danger")
        return redirect(url_for('auth_bp.district_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/district-taluka/<int:taluka_id>')
def district_taluka_detail(taluka_id):
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        if not district_profile:
            flash("Unable to find your district profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        taluka = _get_district_taluka_record(district_profile.get('district_id'), taluka_id, cursor)
        if not taluka:
            flash("That taluka is outside your district scope.", "warning")
            return redirect(url_for('auth_bp.district_talukas_page'))

        taluka_scope = _get_taluka_overview_data({
            'taluka_id': taluka.get('id'),
            'taluka_name': taluka.get('name'),
            'district_name': district_profile.get('district_name'),
        }, cursor)

        return render_template(
            'dashboard/district_taluka_detail.html',
            district_profile=district_profile,
            taluka=taluka,
            villages=_get_taluka_village_rows(taluka_id, cursor),
            workers=_get_taluka_worker_options(taluka_id, cursor),
            assigned_tasks=_get_taluka_recent_manual_tasks(taluka_id, cursor, limit=10),
            recent_requests=_get_taluka_request_items(taluka_id, cursor, limit=10),
            dashboard_stats=taluka_scope['dashboard_stats'],
            district_page='talukas'
        )
    except Exception as e:
        print(f"Error loading district taluka detail: {e}")
        flash("Unable to load that taluka right now.", "danger")
        return redirect(url_for('auth_bp.district_talukas_page'))
    finally:
        conn.close()


@auth_bp.route('/district-taluka-admins')
def district_taluka_admins_page():
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        if not district_profile:
            flash("Unable to find your district profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        taluka_admins = _get_district_taluka_admin_options(district_profile.get('district_id'), cursor)
        return render_template(
            'dashboard/district_taluka_admins.html',
            district_profile=district_profile,
            district_taluka_admins=taluka_admins,
            district_page='taluka_admins'
        )
    except Exception as e:
        print(f"Error loading district taluka admins: {e}")
        flash("Unable to load taluka admins right now.", "danger")
        return redirect(url_for('auth_bp.district_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/district-taluka-admin/<int:taluka_admin_id>')
def district_taluka_admin_detail(taluka_admin_id):
    if session.get('role') != 'district_admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        district_profile = _get_district_admin_profile(session.get('user_id'), cursor)
        if not district_profile:
            flash("Unable to find your district profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        taluka_admin = _get_district_taluka_admin_record(district_profile.get('district_id'), taluka_admin_id, cursor)
        if not taluka_admin:
            flash("That taluka admin is outside your district scope.", "warning")
            return redirect(url_for('auth_bp.district_taluka_admins_page'))

        taluka_scope = _get_taluka_overview_data(taluka_admin, cursor)

        return render_template(
            'dashboard/district_taluka_admin_detail.html',
            district_profile=district_profile,
            taluka_admin=taluka_admin,
            villages=_get_taluka_village_rows(taluka_admin.get('taluka_id'), cursor),
            workers=_get_taluka_worker_options(taluka_admin.get('taluka_id'), cursor),
            district_tasks=_get_taluka_admin_district_tasks(taluka_admin_id, cursor, limit=10),
            dashboard_stats=taluka_scope['dashboard_stats'],
            district_page='taluka_admins'
        )
    except Exception as e:
        print(f"Error loading district taluka admin detail: {e}")
        flash("Unable to load that taluka admin right now.", "danger")
        return redirect(url_for('auth_bp.district_taluka_admins_page'))
    finally:
        conn.close()


@auth_bp.route('/taluka-dashboard')
def taluka_dashboard():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    admin_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(admin_id, cursor)
        overview = _get_taluka_overview_data(admin_profile, cursor)
        return render_template(
            'dashboard/taluka_admin.html',
            admin_profile=admin_profile,
            worker_page='dashboard',
            **overview
        )
    except Exception as e:
        print(f"Error loading taluka dashboard: {e}")
        flash("Unable to load taluka dashboard right now.", "danger")
        return redirect(url_for('auth_bp.login'))
    finally:
        conn.close()


@auth_bp.route('/taluka-reports')
def taluka_reports():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))
    return _render_report_center('admin')


@auth_bp.route('/taluka-reports/export/<format_name>')
def taluka_report_export(format_name):
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))
    return _export_report_center('admin', format_name)


@auth_bp.route('/taluka-equipment')
def taluka_equipment():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        if not admin_profile:
            flash("Unable to find your taluka admin profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        return render_template(
            'dashboard/taluka_equipment.html',
            admin_profile=admin_profile,
            worker_page='equipment',
            **_build_taluka_equipment_context(admin_profile, session.get('user_id'), cursor),
        )
    except Exception as e:
        print(f"Error loading taluka equipment dashboard: {e}")
        flash("Unable to load taluka equipment data right now.", "danger")
        return redirect(url_for('auth_bp.taluka_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/taluka-equipment-issue', methods=['POST'])
def taluka_issue_equipment():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.taluka_equipment')
    request_id = _safe_int(request.form.get('request_id'))
    worker_id = _safe_int(request.form.get('worker_id'))
    item_key = normalize_equipment_item(request.form.get('item_key'))
    notes = (request.form.get('notes') or '').strip()

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        if not admin_profile:
            flash("Taluka profile not found.", "warning")
            return redirect(redirect_target)

        worker = None
        related_request_id = None
        requested_quantity = None

        if request_id:
            request_record = _get_equipment_request_record(request_id, cursor)
            pending_quantity = max(
                (request_record or {}).get('quantity_requested', 0) - (request_record or {}).get('quantity_fulfilled', 0),
                0,
            )
            if not request_record or request_record.get('status') != 'pending':
                flash("That equipment request is no longer pending.", "warning")
                return redirect(redirect_target)
            if request_record.get('target_role') != 'admin' or request_record.get('target_scope_id') != admin_profile.get('taluka_id'):
                flash("You can only fulfill requests assigned to your taluka office.", "warning")
                return redirect(redirect_target)

            worker = _get_taluka_worker_record(admin_profile.get('taluka_id'), request_record.get('requester_id'), cursor)
            if not worker:
                flash("This worker request is outside your taluka scope.", "warning")
                return redirect(redirect_target)

            worker_id = worker.get('id')
            item_key = request_record.get('item_key')
            related_request_id = request_record.get('id')
            requested_quantity = pending_quantity or request_record.get('quantity_requested') or 1

        if not worker and worker_id:
            worker = _get_taluka_worker_record(admin_profile.get('taluka_id'), worker_id, cursor)

        quantity = _positive_int(request.form.get('quantity'), default=requested_quantity)

        if not worker or not item_key or quantity is None:
            flash("Choose a worker, equipment item, and quantity before assigning stock.", "warning")
            return redirect(redirect_target)

        transfer_note = _merge_equipment_note(
            f"Issued to {worker.get('name') or 'Worker'} for {worker.get('village_name') or 'assigned area'}.",
            notes,
        )
        issue_inventory(
            cursor,
            actor_role='admin',
            actor_id=session.get('user_id'),
            from_scope_role='taluka',
            from_scope_id=admin_profile.get('taluka_id'),
            to_scope_role='worker',
            to_scope_id=worker.get('id'),
            item_key=item_key,
            quantity=quantity,
            notes=transfer_note,
            related_request_id=related_request_id,
            transaction_type='issue',
        )
        if related_request_id:
            mark_request_fulfilled(cursor, related_request_id, quantity, transfer_note)

        conn.commit()
        flash("Equipment assigned to the worker successfully.", "success")
    except ValueError as e:
        conn.rollback()
        flash(str(e), "warning")
    except Exception as e:
        conn.rollback()
        print(f"Error issuing taluka equipment: {e}")
        flash("Unable to assign equipment right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/taluka-equipment-restock', methods=['POST'])
def taluka_restock_equipment():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.taluka_equipment')
    item_key = normalize_equipment_item(request.form.get('item_key'))
    quantity = _positive_int(request.form.get('quantity'))
    notes = (request.form.get('notes') or '').strip()

    if not item_key or quantity is None:
        flash("Choose equipment and quantity before updating taluka reserve stock.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        if not admin_profile:
            flash("Taluka profile not found.", "warning")
            return redirect(redirect_target)

        restock_inventory(
            cursor,
            actor_role='admin',
            actor_id=session.get('user_id'),
            scope_role='taluka',
            scope_id=admin_profile.get('taluka_id'),
            item_key=item_key,
            quantity=quantity,
            notes=_merge_equipment_note("Taluka reserve stock updated.", notes),
        )
        conn.commit()
        flash("Taluka reserve stock updated successfully.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error restocking taluka equipment: {e}")
        flash("Unable to update taluka reserve stock right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/taluka-equipment-request-stock', methods=['POST'])
def taluka_request_equipment_stock():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.taluka_equipment')
    item_key = normalize_equipment_item(request.form.get('item_key'))
    quantity = _positive_int(request.form.get('quantity'))
    reason = (request.form.get('reason') or '').strip()
    notes = (request.form.get('notes') or '').strip()
    priority = (request.form.get('priority') or 'normal').strip().lower()
    if priority not in {'low', 'normal', 'high', 'urgent'}:
        priority = 'normal'

    if not item_key or quantity is None or not reason:
        flash("Choose equipment, quantity, and reason before requesting stock from the district office.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        if not admin_profile:
            flash("Taluka profile not found.", "warning")
            return redirect(redirect_target)

        create_equipment_request(
            cursor,
            requester_role='admin',
            requester_id=session.get('user_id'),
            requester_scope_role='taluka',
            requester_scope_id=admin_profile.get('taluka_id'),
            target_role='district_admin',
            target_scope_role='district',
            target_scope_id=admin_profile.get('district_id'),
            item_key=item_key,
            quantity_requested=quantity,
            reason=reason,
            priority=priority,
            notes=notes or None,
        )
        conn.commit()
        flash("Bulk stock request sent to the district office.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error creating taluka equipment request: {e}")
        flash("Unable to send the district stock request right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/taluka-equipment-request/<int:request_id>/reject', methods=['POST'])
def taluka_reject_equipment_request(request_id):
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.taluka_equipment')
    notes = (request.form.get('notes') or '').strip()

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        request_record = _get_equipment_request_record(request_id, cursor)
        if not admin_profile or not request_record or request_record.get('status') != 'pending':
            flash("That equipment request cannot be rejected now.", "warning")
            return redirect(redirect_target)
        if request_record.get('target_role') != 'admin' or request_record.get('target_scope_id') != admin_profile.get('taluka_id'):
            flash("You can only reject requests assigned to your taluka office.", "warning")
            return redirect(redirect_target)

        mark_request_rejected(cursor, request_id, _merge_equipment_note("Rejected by taluka office.", notes))
        conn.commit()
        flash("Equipment request rejected.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error rejecting taluka equipment request: {e}")
        flash("Unable to reject the equipment request right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/taluka-workers')
def taluka_workers():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    admin_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(admin_id, cursor)
        workers = _get_taluka_worker_options(admin_profile.get('taluka_id'), cursor)
        assigned_tasks = _get_taluka_recent_manual_tasks(admin_profile.get('taluka_id'), cursor, limit=10)
        return render_template(
            'dashboard/taluka_workers.html',
            admin_profile=admin_profile,
            workers=workers,
            assigned_tasks=assigned_tasks,
            worker_page='workers'
        )
    except Exception as e:
        print(f"Error loading taluka workers: {e}")
        flash("Unable to load worker management right now.", "danger")
        return redirect(url_for('auth_bp.taluka_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/taluka-complaints')
def taluka_complaints():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    admin_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(admin_id, cursor)
        complaints = _get_taluka_complaints(admin_profile, cursor)
        available_workers = _get_taluka_worker_options(admin_profile.get('taluka_id'), cursor)
        map_requests = _get_taluka_request_items(admin_profile.get('taluka_id'), cursor)
        map_assigned_tasks = _get_taluka_recent_manual_tasks(admin_profile.get('taluka_id'), cursor, limit=None)
        return render_template(
            'dashboard/taluka_complaints.html',
            admin_profile=admin_profile,
            complaints=complaints,
            available_workers=available_workers,
            taluka_map=_build_taluka_map_payload(admin_profile, complaints, map_requests, map_assigned_tasks),
            worker_page='complaints'
        )
    except Exception as e:
        print(f"Error loading taluka complaints: {e}")
        flash("Unable to load complaints right now.", "danger")
        return redirect(url_for('auth_bp.taluka_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/taluka-manage-complaint/<int:complaint_id>', methods=['POST'])
def taluka_manage_complaint(complaint_id):
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.taluka_complaints')
    requested_worker_id = request.form.get('worker_id', type=int)
    requested_status = normalize_complaint_status(request.form.get('status') or 'Pending')

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        ensure_complaint_workflow_columns(cursor)
        cursor.execute("""
            SELECT
                id,
                worker_id,
                COALESCE(status, 'Pending') AS status
            FROM complaints
            WHERE id = %s AND taluka = %s
        """, (complaint_id, admin_profile.get('taluka_name')))
        complaint = cursor.fetchone()
        if not complaint:
            flash("That complaint is outside your taluka scope.", "warning")
            return redirect(redirect_target)

        worker_id = requested_worker_id if requested_worker_id is not None else complaint.get('worker_id')
        if request.form.get('worker_id') == '':
            worker_id = None

        worker = None
        if worker_id:
            worker = _get_taluka_worker_record(admin_profile.get('taluka_id'), worker_id, cursor)
            if not worker:
                flash("Choose a worker from your taluka before saving the complaint.", "warning")
                return redirect(redirect_target)

        status_key = complaint_status_key(requested_status)
        if worker_id and status_key == 'pending':
            status_key = 'assigned'
        if not worker_id and status_key in ('assigned', 'in_progress'):
            status_key = 'pending'
        if status_key == 'completed' and not worker_id:
            flash("Assign the complaint to a worker before marking it completed.", "warning")
            return redirect(redirect_target)

        columns = [
            "worker_id = %s",
            "admin_id = %s",
            "status = %s",
            "updated_at = NOW()",
        ]
        values = [worker_id, session.get('user_id'), normalize_complaint_status(status_key)]

        if worker_id and not complaint.get('worker_id'):
            columns.append("assigned_at = NOW()")
        elif not worker_id:
            columns.append("assigned_at = NULL")

        if status_key == 'completed':
            columns.append("resolved_at = NOW()")
        else:
            columns.append("resolved_at = NULL")

        cursor.execute(f"""
            UPDATE complaints
            SET {', '.join(columns)}
            WHERE id = %s
        """, tuple(values + [complaint_id]))
        conn.commit()

        if worker and status_key == 'assigned':
            flash(f"Complaint assigned to {worker.get('name')}.", "success")
        else:
            flash(f"Complaint status updated to {normalize_complaint_status(status_key)}.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error updating taluka complaint: {e}")
        flash("Unable to update the complaint right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/taluka-villages')
def taluka_villages():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    admin_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(admin_id, cursor)
        return render_template(
            'dashboard/taluka_villages.html',
            admin_profile=admin_profile,
            villages=_get_taluka_village_rows(admin_profile.get('taluka_id'), cursor),
            worker_page='villages'
        )
    except Exception as e:
        print(f"Error loading taluka villages: {e}")
        flash("Unable to load villages right now.", "danger")
        return redirect(url_for('auth_bp.taluka_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/taluka-profile')
def taluka_profile():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    admin_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(admin_id, cursor)
        overview = _get_taluka_overview_data(admin_profile, cursor)
        return render_template(
            'dashboard/taluka_profile.html',
            admin_profile=admin_profile,
            dashboard_stats=overview['dashboard_stats'],
            worker_page='profile'
        )
    except Exception as e:
        print(f"Error loading taluka profile: {e}")
        flash("Unable to load profile right now.", "danger")
        return redirect(url_for('auth_bp.taluka_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/taluka-assign-request/<int:request_id>', methods=['POST'])
def taluka_assign_request(request_id):
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.taluka_dashboard')
    worker_id = request.form.get('worker_id', type=int)
    if not worker_id:
        flash("Please choose a worker before assigning a request.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        worker = _get_taluka_worker_record(admin_profile.get('taluka_id'), worker_id, cursor)
        if not worker:
            flash("Select a worker from your taluka before assigning the request.", "warning")
            return redirect(redirect_target)

        cursor.execute("""
            SELECT
                r.id
            FROM requests r
            LEFT JOIN users u ON r.user_id = u.id
            LEFT JOIN villages v ON u.village_id = v.id
            WHERE r.id = %s AND COALESCE(u.taluka_id, v.taluka_id) = %s
        """, (request_id, admin_profile.get('taluka_id')))
        request_row = cursor.fetchone()
        if not request_row:
            flash("That request is outside your taluka scope.", "warning")
            return redirect(redirect_target)

        cursor.execute("""
            UPDATE requests
            SET worker_id = %s,
                admin_id = %s,
                status = 'pending'
            WHERE id = %s
        """, (worker_id, session.get('user_id'), request_id))
        conn.commit()
        flash(f"Request assigned to {worker.get('name')} successfully.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error assigning taluka request: {e}")
        flash("Unable to assign the request right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/taluka-assign-task', methods=['POST'])
def taluka_assign_task():
    if session.get('role') != 'admin':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.taluka_workers')
    worker_id = request.form.get('worker_id', type=int)
    location_name = (request.form.get('location_name') or '').strip()
    description = (request.form.get('description') or '').strip()
    priority = (request.form.get('priority') or 'medium').strip().lower()

    if priority not in {'low', 'medium', 'high'}:
        priority = 'medium'

    if not worker_id or not location_name or not description:
        flash("Choose a worker, area, and task description before assigning work.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        admin_profile = _get_taluka_admin_profile(session.get('user_id'), cursor)
        worker = _get_taluka_worker_record(admin_profile.get('taluka_id'), worker_id, cursor)
        if not worker:
            flash("You can only assign tasks to workers inside your taluka.", "warning")
            return redirect(redirect_target)

        cursor.execute("""
            INSERT INTO tasks (worker_id, location_name, description, status, priority, assigned_at)
            VALUES (%s, %s, %s, 'pending', %s, NOW())
        """, (worker_id, location_name, description, priority))
        conn.commit()
        flash(f"Task assigned to {worker.get('name')} for {worker.get('village_name')}.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error assigning taluka task: {e}")
        flash("Unable to assign the task right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)

@auth_bp.route('/worker-dashboard')
def worker_dashboard():
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))
    
    worker_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    
    try:
        worker_profile = _get_worker_profile(worker_id, cursor)
        worker_stats = _get_worker_assignment_stats(worker_id, cursor)
        completed = worker_stats.get('completed') or 0
        assigned = worker_stats.get('assigned') or 0
        worker_stats['rating'] = round(min(5, max(1, completed / 5)), 1) if completed else 0
        worker_stats['efficiency'] = round((completed / assigned) * 100) if assigned else 0
        recent_active_tasks = _get_worker_work_items(worker_id, cursor, worker_profile=worker_profile, status_filter='active', limit=5)
        
        return render_template('dashboard/worker_dash.html', 
                             worker_profile=worker_profile,
                             worker_stats=worker_stats,
                             recent_active_tasks=recent_active_tasks if recent_active_tasks else [],
                             worker_map=_build_worker_map_config(worker_profile),
                             worker_page='dashboard')
    
    except Exception as e:
        print(f"Error fetching worker data: {e}")
        flash("Error loading dashboard data.", "danger")
        return redirect(url_for('auth_bp.login'))
    finally:
        conn.close()


@auth_bp.route('/worker-stats-api')
def worker_stats_api():
    if session.get('role') != 'worker':
        return jsonify({'error': 'Unauthorized'}), 403

    worker_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        stats = _get_worker_assignment_stats(worker_id, cursor)
        assigned = stats.get('assigned') or 0
        completed = stats.get('completed') or 0
        efficiency = round((completed / assigned) * 100) if assigned else 0
        rating = round(min(5, max(1, completed / 5)), 1) if completed else 0

        return jsonify({
            'assigned': assigned,
            'pending': stats.get('pending') or 0,
            'in_progress': stats.get('in_progress') or 0,
            'completed': completed,
            'completed_today': stats.get('completed_today') or 0,
            'monthly_earnings': stats.get('monthly_earnings') or 0,
            'today_earnings': stats.get('today_earnings') or 0,
            'efficiency': efficiency,
            'rating': rating,
        })
    except Exception as e:
        print(f"Error fetching worker stats api data: {e}")
        return jsonify({'assigned': 0, 'pending': 0, 'in_progress': 0, 'completed': 0, 'completed_today': 0, 'monthly_earnings': 0, 'today_earnings': 0, 'efficiency': 0, 'rating': 0})
    finally:
        conn.close()


@auth_bp.route('/worker-tasks-api')
def worker_tasks_api():
    if session.get('role') != 'worker':
        return jsonify({'error': 'Unauthorized'}), 403

    worker_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        worker_profile = _get_worker_profile(worker_id, cursor)
        tasks = _get_worker_work_items(worker_id, cursor, worker_profile=worker_profile, status_filter='active', limit=10)
        return jsonify({'active_tasks': tasks})
    except Exception as e:
        print(f"Error fetching worker tasks api data: {e}")
        return jsonify({'active_tasks': []})
    finally:
        conn.close()


@auth_bp.route('/worker-start-task/<int:task_id>', methods=['POST'])
def worker_start_task(task_id):
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.worker_dashboard')
    conn = get_db()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            UPDATE requests
            SET status = 'in_progress'
            WHERE id = %s AND worker_id = %s AND status = 'pending'
        """, (task_id, session.get('user_id')))
        conn.commit()
        if cursor.rowcount:
            flash("Task started successfully.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error starting worker task: {e}")
        flash("Unable to start the task right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/worker-start-complaint/<int:complaint_id>', methods=['POST'])
def worker_start_complaint(complaint_id):
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.worker_dashboard')
    conn = get_db()
    cursor = conn.cursor()

    try:
        ensure_complaint_workflow_columns(cursor)
        cursor.execute("""
            UPDATE complaints
            SET status = 'In Progress',
                updated_at = NOW()
            WHERE id = %s
              AND worker_id = %s
              AND LOWER(COALESCE(status, 'Pending')) IN ('pending', 'assigned', 'in progress', 'in_progress')
        """, (complaint_id, session.get('user_id')))
        conn.commit()
        if cursor.rowcount:
            flash("Complaint work started successfully.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error starting worker complaint: {e}")
        flash("Unable to start the complaint right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/worker-complete-task/<int:task_id>', methods=['POST'])
def worker_complete_task(task_id):
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.worker_dashboard')
    conn = get_db()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            UPDATE requests
            SET status = 'completed'
            WHERE id = %s AND worker_id = %s AND status IN ('pending', 'in_progress')
        """, (task_id, session.get('user_id')))
        conn.commit()
        if cursor.rowcount:
            flash("Task marked as completed.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error completing worker task: {e}")
        flash("Unable to complete the task right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/worker-complete-complaint/<int:complaint_id>', methods=['POST'])
def worker_complete_complaint(complaint_id):
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.worker_dashboard')
    completion_photo = request.files.get('completion_photo')
    if not completion_photo or not completion_photo.filename:
        flash("Upload a completion photo before resolving the complaint.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor()

    try:
        complaint_columns = ensure_complaint_workflow_columns(cursor)
        completion_photo_name = save_complaint_upload(completion_photo, current_app.root_path, prefix='worker')
        cursor.execute("""
            UPDATE complaints
            SET status = 'Completed',
                resolution_photo_path = %s,
                updated_at = NOW(),
                resolved_at = NOW()
              WHERE id = %s
                AND worker_id = %s
                AND LOWER(COALESCE(status, 'Pending')) IN ('pending', 'assigned', 'in progress', 'in_progress')
        """, (completion_photo_name, complaint_id, session.get('user_id')))
        conn.commit()
        if cursor.rowcount:
            flash("Complaint marked as completed.", "success")
        else:
            flash("Complaint could not be completed from the current state.", "warning")
    except Exception as e:
        conn.rollback()
        print(f"Error completing worker complaint: {e}")
        flash("Unable to complete the complaint right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/worker-complete-assigned-task/<int:task_id>', methods=['POST'])
def worker_complete_assigned_task(task_id):
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.worker_dashboard')
    conn = get_db()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            UPDATE tasks
            SET status = 'completed'
            WHERE id = %s AND worker_id = %s AND status = 'pending'
        """, (task_id, session.get('user_id')))
        conn.commit()
        if cursor.rowcount:
            flash("Assigned task marked as completed.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error completing assigned task: {e}")
        flash("Unable to complete the assigned task right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)


@auth_bp.route('/worker-requests')
def worker_requests():
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    worker_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        worker_profile = _get_worker_profile(worker_id, cursor)
        tasks = _get_worker_work_items(worker_id, cursor, worker_profile=worker_profile, status_filter='all')

        return render_template(
            'dashboard/worker_requests.html',
            worker_profile=worker_profile,
            tasks=tasks,
            worker_page='tasks'
        )
    except Exception as e:
        print(f"Error loading worker requests: {e}")
        flash("Unable to load worker tasks right now.", "danger")
        return redirect(url_for('auth_bp.worker_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/worker-assigned-areas')
def worker_assigned_areas():
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    worker_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        worker_profile = _get_worker_profile(worker_id, cursor)
        area_stats = _get_worker_assignment_stats(worker_id, cursor)

        assigned_areas = []
        if worker_profile:
            assigned_areas.append({
                'village_name': worker_profile.get('village_name') or 'Village not assigned',
                'vehicle_no': worker_profile.get('vehicle_no') or 'Not assigned',
                'status': worker_profile.get('status') or 'Active',
                'total_tasks': area_stats.get('total_tasks') or 0,
                'pending_tasks': area_stats.get('pending_tasks') or 0,
                'in_progress_tasks': area_stats.get('in_progress_tasks') or 0,
                'completed_tasks': area_stats.get('completed_tasks') or 0,
            })

        return render_template(
            'dashboard/worker_assigned_areas.html',
            worker_profile=worker_profile,
            assigned_areas=assigned_areas,
            worker_page='areas'
        )
    except Exception as e:
        print(f"Error loading assigned areas: {e}")
        flash("Unable to load assigned area details right now.", "danger")
        return redirect(url_for('auth_bp.worker_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/worker-history')
def worker_history():
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    worker_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        worker_profile = _get_worker_profile(worker_id, cursor)
        history_items = _get_worker_work_items(worker_id, cursor, worker_profile=worker_profile, status_filter='completed')

        return render_template(
            'dashboard/worker_history.html',
            worker_profile=worker_profile,
            history_items=history_items,
            worker_page='history'
        )
    except Exception as e:
        print(f"Error loading worker history: {e}")
        flash("Unable to load worker history right now.", "danger")
        return redirect(url_for('auth_bp.worker_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/worker-earnings')
def worker_earnings():
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    worker_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        worker_profile = _get_worker_profile(worker_id, cursor)
        cursor.execute("""
            SELECT
                id,
                COALESCE(garbage_type, 'Garbage Collection') AS garbage_type,
                COALESCE(amount, 0) AS amount,
                COALESCE(status, 'pending') AS status,
                created_at
            FROM requests
            WHERE worker_id = %s
            ORDER BY created_at DESC
        """, (worker_id,))
        earnings_rows = cursor.fetchall() or []

        total_earned = sum((row.get('amount') or 0) for row in earnings_rows if row.get('status') == 'completed')
        pending_amount = sum((row.get('amount') or 0) for row in earnings_rows if row.get('status') != 'completed')

        return render_template(
            'dashboard/worker_earnings.html',
            worker_profile=worker_profile,
            earnings_rows=earnings_rows,
            total_earned=total_earned,
            pending_amount=pending_amount,
            worker_page='earnings'
        )
    except Exception as e:
        print(f"Error loading worker earnings: {e}")
        flash("Unable to load worker earnings right now.", "danger")
        return redirect(url_for('auth_bp.worker_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/worker-profile')
def worker_profile():
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    worker_id = session.get('user_id')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        profile = _get_worker_profile(worker_id, cursor)
        profile_stats = _get_worker_assignment_stats(worker_id, cursor)

        return render_template(
            'dashboard/worker_profile.html',
            worker_profile=profile,
            profile_stats=profile_stats,
            worker_page='profile'
        )
    except Exception as e:
        print(f"Error loading worker profile: {e}")
        flash("Unable to load worker profile right now.", "danger")
        return redirect(url_for('auth_bp.worker_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/worker-equipment')
def worker_equipment():
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        worker_profile = _get_worker_profile(session.get('user_id'), cursor)
        if not worker_profile:
            flash("Unable to find your worker profile.", "warning")
            return redirect(url_for('auth_bp.login'))

        return render_template(
            'dashboard/worker_equipment.html',
            worker_profile=worker_profile,
            worker_page='equipment',
            **_build_worker_equipment_context(worker_profile, cursor),
        )
    except Exception as e:
        print(f"Error loading worker equipment dashboard: {e}")
        flash("Unable to load your equipment page right now.", "danger")
        return redirect(url_for('auth_bp.worker_dashboard'))
    finally:
        conn.close()


@auth_bp.route('/worker-equipment-request', methods=['POST'])
def worker_equipment_request():
    if session.get('role') != 'worker':
        return redirect(url_for('auth_bp.login'))

    redirect_target = _post_redirect_target('auth_bp.worker_equipment')
    item_key = normalize_equipment_item(request.form.get('item_key'))
    quantity = _positive_int(request.form.get('quantity'))
    reason = (request.form.get('reason') or '').strip()
    notes = (request.form.get('notes') or '').strip()
    priority = (request.form.get('priority') or 'normal').strip().lower()
    if priority not in {'low', 'normal', 'high', 'urgent'}:
        priority = 'normal'

    if not item_key or quantity is None or not reason:
        flash("Choose equipment, quantity, and reason before sending a request.", "warning")
        return redirect(redirect_target)

    conn = get_db()
    cursor = conn.cursor(dictionary=True)

    try:
        worker_profile = _get_worker_profile(session.get('user_id'), cursor)
        if not worker_profile or not worker_profile.get('taluka_id'):
            flash("Your taluka scope is not available right now.", "warning")
            return redirect(redirect_target)

        create_equipment_request(
            cursor,
            requester_role='worker',
            requester_id=session.get('user_id'),
            requester_scope_role='worker',
            requester_scope_id=session.get('user_id'),
            target_role='admin',
            target_scope_role='taluka',
            target_scope_id=worker_profile.get('taluka_id'),
            item_key=item_key,
            quantity_requested=quantity,
            reason=reason,
            priority=priority,
            notes=notes or None,
        )
        conn.commit()
        flash("Equipment request sent to your taluka admin.", "success")
    except Exception as e:
        conn.rollback()
        print(f"Error creating worker equipment request: {e}")
        flash("Unable to send the equipment request right now.", "danger")
    finally:
        conn.close()

    return redirect(redirect_target)

# --- 🟡 API ROUTES FOR DROPDOWNS ---

@auth_bp.route('/get_districts/<int:state_id>')
def get_districts(state_id):
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("SELECT id, name FROM districts WHERE state_id = %s ORDER BY name ASC", (state_id,))
    data = cursor.fetchall()
    conn.close()
    return jsonify(data)

@auth_bp.route('/get_talukas/<int:district_id>')
def get_talukas(district_id):
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("SELECT id, name FROM talukas WHERE district_id = %s ORDER BY name ASC", (district_id,))
    data = cursor.fetchall()
    conn.close()
    return jsonify(data)

@auth_bp.route('/get_villages/<int:taluka_id>')
def get_villages(taluka_id):
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("SELECT id, name FROM villages WHERE taluka_id = %s ORDER BY name ASC", (taluka_id,))
    data = cursor.fetchall()
    conn.close()
    return jsonify(data)


@auth_bp.route('/search-villages')
def search_villages():
    state_id = _safe_int(request.args.get('state_id'))
    district_id = _safe_int(request.args.get('district_id'))
    taluka_id = _safe_int(request.args.get('taluka_id'))
    village_id = _safe_int(request.args.get('village_id'))
    search_term = (request.args.get('search') or '').strip()

    clauses = []
    params = []

    if state_id:
        clauses.append("s.id = %s")
        params.append(state_id)
    if district_id:
        clauses.append("d.id = %s")
        params.append(district_id)
    if taluka_id:
        clauses.append("t.id = %s")
        params.append(taluka_id)
    if village_id:
        clauses.append("v.id = %s")
        params.append(village_id)
    if search_term:
        clauses.append("LOWER(v.name) LIKE %s")
        params.append(f"%{search_term.lower()}%")

    where_sql = f"WHERE {' AND '.join(clauses)}" if clauses else ""

    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    try:
        count_query = f"""
            SELECT COUNT(*) AS total
            FROM villages v
            INNER JOIN talukas t ON v.taluka_id = t.id
            INNER JOIN districts d ON t.district_id = d.id
            INNER JOIN states s ON d.state_id = s.id
            {where_sql}
        """
        cursor.execute(count_query, tuple(params))
        total = (cursor.fetchone() or {}).get('total', 0)

        data_query = f"""
            SELECT
                s.id AS state_id,
                s.name AS state_name,
                d.id AS district_id,
                d.name AS district_name,
                t.id AS taluka_id,
                t.name AS taluka_name,
                v.id AS village_id,
                v.name AS village_name
            FROM villages v
            INNER JOIN talukas t ON v.taluka_id = t.id
            INNER JOIN districts d ON t.district_id = d.id
            INNER JOIN states s ON d.state_id = s.id
            {where_sql}
            ORDER BY s.name ASC, d.name ASC, t.name ASC, v.name ASC
            LIMIT 250
        """
        cursor.execute(data_query, tuple(params))
        rows = cursor.fetchall() or []
    finally:
        conn.close()

    return jsonify({
        'rows': rows,
        'total': total,
        'returned': len(rows),
        'limited': total > len(rows),
    })
